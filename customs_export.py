"""Create the five-sheet customs package and its immutable data snapshot."""

from collections import OrderedDict
from copy import copy
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins


NAVY = '1F4E78'


def _text(value):
    if value is None:
        return ''
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return str(value)


def _unique(values):
    return '；'.join(dict.fromkeys(
        _text(value).strip() for value in values if _text(value).strip()
    ))


def build_customs_snapshot(pi, form_data):
    """Capture every value used by the export so confirmed files stay stable."""
    groups = OrderedDict()
    for item in pi.items:
        product = item.product
        key = (
            product.customs_hs_code,
            product.customs_name_cn,
            product.customs_name_en,
            product.customs_unit,
        )
        group = groups.setdefault(key, {
            'hs_code': product.customs_hs_code,
            'name_cn': product.customs_name_cn,
            'name_en': product.customs_name_en,
            'unit': product.customs_unit,
            'quantity': 0,
            'amount': 0.0,
            'product_details': [],
            'brand_type': [],
            'preferential': [],
            'purpose': [],
            'brand': [],
            'origin_country': [],
            'domestic_source': [],
            'tax_exemption': [],
            'elements': [],
        })
        group['quantity'] += int(item.quantity or 0)
        group['amount'] += float(item.amount or 0)
        detail = ' / '.join(
            value for value in (
                product.name, product.product_code, product.specification,
            ) if value
        )
        if detail:
            group['product_details'].append(detail)
        for key_name, field_name in (
            ('brand_type', 'customs_brand_type'),
            ('preferential', 'customs_preferential'),
            ('purpose', 'customs_purpose'),
            ('brand', 'customs_brand'),
            ('origin_country', 'customs_origin_country'),
            ('domestic_source', 'customs_domestic_source'),
            ('tax_exemption', 'customs_tax_exemption'),
            ('elements', 'customs_elements'),
        ):
            group[key_name].append(getattr(product, field_name, '') or '')

    normalized_groups = []
    for group in groups.values():
        for key_name in (
            'product_details', 'brand_type', 'preferential', 'purpose',
            'brand', 'origin_country', 'domestic_source', 'tax_exemption',
            'elements',
        ):
            group[key_name] = _unique(group[key_name])
        group['amount'] = round(group['amount'], 2)
        group['unit_price'] = round(
            group['amount'] / group['quantity'], 6
        ) if group['quantity'] else 0
        normalized_groups.append(group)

    packing_list = pi.packing_list
    boxes = []
    for box in packing_list.boxes if packing_list else []:
        boxes.append({
            'box_no': box.box_no,
            'net_weight': float(box.net_weight or 0),
            'gross_weight': float(box.gross_weight or 0),
            'length_cm': float(box.length_cm or 0),
            'width_cm': float(box.width_cm or 0),
            'height_cm': float(box.height_cm or 0),
            'volume_cbm': float(box.volume_cbm or 0),
            'shipping_mark': box.shipping_mark or '',
            'note': box.note or '',
            'items': [
                {
                    'product_name': item.product_name,
                    'product_code': item.product_code or '',
                    'specification': item.specification or '',
                    'quantity': int(item.quantity or 0),
                    'note': item.note or '',
                }
                for item in box.items
            ],
        })

    company_name = (
        'Changzhou Qisuo Welding And Cutting Equipment Co., Ltd.'
        if (pi.company or '').lower() == 'qisuo'
        else 'CHANGZHOU KLISTA INTERNATIONAL TRADE CO., LTD.'
    )
    return {
        'pi_number': pi.pi_number,
        'issue_date': _text(pi.issue_date),
        'currency': (pi.currency or 'USD').upper(),
        'salesperson': pi.salesperson or '',
        'payment_terms': pi.payment_terms or '',
        'price_terms': pi.price_terms or '',
        'delivery_time': pi.delivery_time or '',
        'bank_info': pi.bank_info or '',
        'company_name': company_name,
        'customer': {
            'name': pi.customer.name,
            'contact_person': pi.customer.contact_person or '',
            'country': pi.customer.country or '',
            'address': pi.customer.address or '',
            'email': pi.customer.email or '',
            'phone': pi.customer.phone or '',
        },
        'product_total': round(sum(
            group['amount'] for group in normalized_groups
        ), 2),
        'other_charges': round(float(pi.other_charges or 0), 2),
        'other_charges_treatment': 'exclude',
        'customs': {
            key: _text(value) for key, value in form_data.items()
            if not key.startswith('_')
        },
        'groups': normalized_groups,
        'boxes': boxes,
    }


def _title(ws, title, columns):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = ws.cell(1, 1, title)
    cell.font = Font(name='Microsoft YaHei', size=18, bold=True, color=NAVY)
    cell.alignment = Alignment(horizontal='center')


def _header(ws, row, labels):
    for index, label in enumerate(labels, 1):
        cell = ws.cell(row, index, label)
        cell.font = Font(
            name='Microsoft YaHei', bold=True, color='FFFFFF'
        )
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.alignment = Alignment(horizontal='center', wrap_text=True)


def _finish(ws, widths, landscape=False):
    ws.sheet_view.showGridLines = False
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    thin = Side(style='thin', color='B7C4CE')
    for row in ws.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = 'Microsoft YaHei'
            cell.font = font
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if cell.row > 2:
                cell.border = Border(bottom=thin)
    ws.freeze_panes = 'A3'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = (
        ws.ORIENTATION_LANDSCAPE if landscape else ws.ORIENTATION_PORTRAIT
    )
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(
        left=0.3, right=0.3, top=0.45, bottom=0.45,
        header=0.2, footer=0.2,
    )


def _write_info(ws, rows, start=3, end_column=7):
    for row, (label, value) in enumerate(rows, start):
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.merge_cells(
            start_row=row, start_column=2,
            end_row=row, end_column=end_column,
        )
        ws.cell(row, 2, value)


def generate_customs_workbook(pi, document, data):
    """Generate a customs workbook from a confirmed snapshot or live draft."""
    snapshot = data.get('_confirmed_snapshot') or build_customs_snapshot(pi, data)
    groups = snapshot['groups']
    boxes = snapshot['boxes']
    customs = snapshot['customs']
    customer = snapshot['customer']
    currency = snapshot['currency']
    workbook = Workbook()
    workbook.remove(workbook.active)

    ws = workbook.create_sheet('报关单')
    _title(ws, '中华人民共和国海关出口货物报关单（资料草单）', 14)
    _write_info(ws, [
        ('境内发货人', snapshot['company_name']),
        ('生产销售单位', snapshot['company_name']),
        ('合同协议号', snapshot['pi_number']),
        ('境外收货人', customer['name']),
        ('出境关别', customs.get('export_customs', '')),
        ('运输方式', customs.get('transport_mode', '')),
        ('运输工具及航次', customs.get('vehicle_voyage', '')),
        ('提运单号', customs.get('bill_no', '')),
        ('贸易国（地区）', customs.get('trade_country', '')),
        ('运抵国（地区）', customs.get('destination_country', '')),
        ('指运港', customs.get('destination_port', '')),
        ('离境口岸', customs.get('departure_port', '')),
        ('出口日期', customs.get('export_date', '')),
        ('申报日期', customs.get('declaration_date', '')),
        ('监管方式', customs.get('supervision_mode', '')),
        ('征免性质', customs.get('exemption_nature', '')),
        ('许可证号', customs.get('license_no', '')),
        ('成交方式', snapshot['price_terms']),
        ('运输包装种类', customs.get('packing_type', '')),
        ('随附单证及编号', customs.get('accompanying_documents', '')),
        ('标记唛码及备注', customs.get('marks_notes', '')),
        ('申报人员', customs.get('declaration_agent', '')),
    ], end_column=6)
    for row, label, value in (
        (3, '包装件数', len(boxes)),
        (4, '总毛重', sum(box['gross_weight'] for box in boxes)),
        (5, '总净重', sum(box['net_weight'] for box in boxes)),
    ):
        ws.cell(row, 8, label).font = Font(bold=True)
        ws.cell(row, 9, value)
    start = 25
    _header(ws, start, [
        '项号', '商品编号', '商品名称', '规格/备注', '数量', '分项净重',
        '单位', '单价', '总价', '币制', '原产国（地区）',
        '最终目的国（地区）', '境内货源地', '征免',
    ])
    for number, group in enumerate(groups, 1):
        values = [
            number, group['hs_code'], group['name_cn'],
            group['product_details'], group['quantity'], '', group['unit'],
            group['unit_price'], group['amount'], currency,
            group['origin_country'], customs.get('destination_country', ''),
            group['domestic_source'], group['tax_exemption'],
        ]
        for column, value in enumerate(values, 1):
            ws.cell(start + number, column, value)
        ws.cell(start + number, 5).number_format = '#,##0'
        ws.cell(start + number, 8).number_format = '#,##0.000000'
        ws.cell(start + number, 9).number_format = '#,##0.00'
    _finish(ws, {
        'A': 7, 'B': 14, 'C': 18, 'D': 34, 'E': 10, 'F': 12,
        'G': 9, 'H': 13, 'I': 13, 'J': 8, 'K': 14, 'L': 16,
        'M': 14, 'N': 12,
    }, landscape=True)

    ws = workbook.create_sheet('发票')
    _title(ws, 'INVOICE / 发票', 7)
    _write_info(ws, [
        ('Seller', snapshot['company_name']),
        ('Invoice No.', snapshot['pi_number']),
        ('Date', snapshot['issue_date']),
        ('Buyer', customer['name']),
        ('Contact', customer['contact_person']),
        ('Address', customer['address']),
        ('From / To', (
            f"{customs.get('departure_port', '')} / "
            f"{customs.get('destination_port', '')}"
        )),
        ('Currency', currency),
        ('Price Term', snapshot['price_terms']),
    ])
    table_row = 14
    _header(ws, table_row, [
        'No.', 'HS Code', 'Descriptions', 'Qty', 'Unit',
        'Unit Price', 'Amount',
    ])
    for number, group in enumerate(groups, 1):
        values = [
            number, group['hs_code'],
            f"{group['name_en']}\n{group['product_details']}",
            group['quantity'], group['unit'], group['unit_price'],
            group['amount'],
        ]
        for column, value in enumerate(values, 1):
            ws.cell(table_row + number, column, value)
        ws.cell(table_row + number, 4).number_format = '#,##0'
        ws.cell(table_row + number, 6).number_format = '#,##0.000000'
        ws.cell(table_row + number, 7).number_format = '#,##0.00'
    total_row = table_row + len(groups) + 1
    ws.cell(total_row, 6, 'Product Total').font = Font(bold=True)
    ws.cell(total_row, 7, snapshot['product_total'])
    ws.cell(total_row + 1, 6, 'Other Charges (excluded)').font = Font(italic=True)
    ws.cell(total_row + 1, 7, snapshot['other_charges'])
    ws.cell(total_row, 7).number_format = '#,##0.00'
    ws.cell(total_row + 1, 7).number_format = '#,##0.00'
    ws.cell(total_row + 3, 1, (
        f"Packing: {len(boxes)} CARTONS    "
        f"Gross Weight: {sum(box['gross_weight'] for box in boxes):g} KGS    "
        f"Net Weight: {sum(box['net_weight'] for box in boxes):g} KGS"
    ))
    ws.merge_cells(
        start_row=total_row + 3, start_column=1,
        end_row=total_row + 3, end_column=7,
    )
    _finish(ws, {
        'A': 8, 'B': 16, 'C': 48, 'D': 12,
        'E': 10, 'F': 16, 'G': 16,
    })

    ws = workbook.create_sheet('装箱单')
    _title(ws, 'PACKING LIST / 装箱单', 8)
    _write_info(ws, [
        ('Seller', snapshot['company_name']),
        ('PI No.', snapshot['pi_number']),
        ('Buyer', customer['name']),
        ('Date', snapshot['issue_date']),
        ('Salesperson', snapshot['salesperson']),
        ('Price Term', snapshot['price_terms']),
        ('From / To', (
            f"{customs.get('departure_port', '')} / "
            f"{customs.get('destination_port', '')}"
        )),
        ('Packing', f"CARTONS: {len(boxes)}"),
    ], start=3, end_column=8)
    _header(ws, 12, [
        '箱序', '箱号', '产品', '数量', '净重(kg)', '毛重(kg)',
        '尺寸(cm)', '体积(m³)',
    ])
    row = 13
    for box_index, box in enumerate(boxes, 1):
        items = box['items'] or [{'product_name': '—', 'quantity': 0}]
        for item_index, item in enumerate(items):
            first = item_index == 0
            values = [
                box_index if first else '', box['box_no'] if first else '',
                item.get('product_name', ''), item.get('quantity', 0),
                box['net_weight'] if first else '',
                box['gross_weight'] if first else '',
                (f"{box['length_cm']:g}×{box['width_cm']:g}×"
                 f"{box['height_cm']:g}") if first else '',
                box['volume_cbm'] if first else '',
            ]
            for column, value in enumerate(values, 1):
                ws.cell(row, column, value)
            row += 1
    ws.cell(row, 3, 'TOTAL / 合计').font = Font(bold=True)
    ws.cell(row, 4, sum(
        item.get('quantity', 0) for box in boxes for item in box['items']
    ))
    ws.cell(row, 5, sum(box['net_weight'] for box in boxes))
    ws.cell(row, 6, sum(box['gross_weight'] for box in boxes))
    ws.cell(row, 8, sum(box['volume_cbm'] for box in boxes))
    _finish(ws, {
        'A': 8, 'B': 12, 'C': 45, 'D': 12,
        'E': 13, 'F': 13, 'G': 22, 'H': 14,
    })

    ws = workbook.create_sheet('申报要素')
    _title(ws, '申报要素', 4)
    _header(ws, 3, ['HS编码', '报关品名', '项目', '内容'])
    row = 4
    for group in groups:
        elements = [
            ('品牌类型', group['brand_type']),
            ('出口享惠情况', group['preferential']),
            ('用途', group['purpose']),
            ('品牌', group['brand']),
            ('型号/规格', group['product_details']),
            ('原产国', group['origin_country']),
            ('境内货源地', group['domestic_source']),
            ('征免', group['tax_exemption']),
            ('申报数量', f"{group['quantity']} {group['unit']}"),
            ('申报金额', f"{currency} {group['amount']:,.2f}"),
            ('其他申报要素', group['elements']),
        ]
        for element_index, (label, value) in enumerate(elements):
            values = [
                group['hs_code'] if element_index == 0 else '',
                group['name_cn'] if element_index == 0 else '',
                label, value,
            ]
            for column, cell_value in enumerate(values, 1):
                ws.cell(row, column, cell_value)
            row += 1
    _finish(ws, {'A': 16, 'B': 22, 'C': 18, 'D': 65})

    ws = workbook.create_sheet('合同')
    _title(ws, 'SALES CONTRACT / 销售合同', 7)
    _write_info(ws, [
        ('Contract No. / 合同号', snapshot['pi_number']),
        ('Seller / 卖方', snapshot['company_name']),
        ('Buyer / 买方', customer['name']),
        ('Date / 日期', snapshot['issue_date']),
        ('Currency / 币种', currency),
        ('Payment Terms / 付款条款', snapshot['payment_terms']),
        ('Price Terms / 价格条款', snapshot['price_terms']),
        ('Delivery / 交货期', snapshot['delivery_time']),
        ('Bank / 银行信息', snapshot['bank_info']),
    ], end_column=7)
    table_row = 14
    _header(ws, table_row, [
        'No.', 'Description / 品名', 'HS Code', 'Qty', 'Unit',
        'Unit Price', 'Amount',
    ])
    for number, group in enumerate(groups, 1):
        values = [
            number, f"{group['name_en']} / {group['name_cn']}",
            group['hs_code'], group['quantity'], group['unit'],
            group['unit_price'], group['amount'],
        ]
        for column, value in enumerate(values, 1):
            ws.cell(table_row + number, column, value)
        ws.cell(table_row + number, 4).number_format = '#,##0'
        ws.cell(table_row + number, 6).number_format = '#,##0.000000'
        ws.cell(table_row + number, 7).number_format = '#,##0.00'
    total_row = table_row + len(groups) + 1
    ws.cell(total_row, 6, 'TOTAL').font = Font(bold=True)
    ws.cell(total_row, 7, snapshot['product_total'])
    ws.cell(total_row, 7).number_format = '#,##0.00'
    ws.cell(total_row + 2, 1, 'Seller Signature / 卖方签章')
    ws.cell(total_row + 2, 5, 'Buyer Signature / 买方签章')
    _finish(ws, {
        'A': 12, 'B': 38, 'C': 16, 'D': 12,
        'E': 10, 'F': 15, 'G': 15,
    })

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
