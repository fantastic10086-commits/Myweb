"""Validation and customer matching for opening customer turnover, USD only."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

DEFAULT_CUTOFF = date(2026, 9, 17)
HEADERS = ('客户编号', '客户名称', '历史成交额（USD）', '截止日期', '备注')
FIELDS = ('historical_deal_usd', 'historical_deal_cutoff', 'historical_deal_note')


def history_values(amount, cutoff, note):
    try:
        value = Decimal(str(amount).strip().replace(',', ''))
    except (InvalidOperation, ValueError):
        raise ValueError('历史成交额必须填写有效美元金额。')
    if not value.is_finite() or value < 0 or value > Decimal('999999999999.99'):
        raise ValueError('历史成交额必须是非负有效金额，且不超过 999999999999.99。')
    if value != value.quantize(Decimal('0.01')):
        raise ValueError('历史成交额最多保留两位小数。')
    if isinstance(cutoff, datetime):
        cutoff = cutoff.date()
    if not isinstance(cutoff, date):
        try:
            cutoff = date.fromisoformat(str(cutoff or DEFAULT_CUTOFF.isoformat()).strip())
        except ValueError:
            raise ValueError('截止日期请使用 YYYY-MM-DD 格式。')
    if cutoff > date.today():
        raise ValueError('截止日期不能晚于今天。')
    note = str(note or '').strip()
    if len(note) > 1000:
        raise ValueError('历史成交额备注不能超过 1000 个字符。')
    return {'historical_deal_usd': float(value), 'historical_deal_cutoff': cutoff,
            'historical_deal_note': note}


def preview_workbook(source, customers):
    wb = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = wb['历史成交额'] if '历史成交额' in wb.sheetnames else wb.worksheets[0]
        iterator = sheet.iter_rows(max_col=5, values_only=True)
        if tuple(next(iterator, ())[:5]) != HEADERS:
            raise ValueError('文件表头不正确，请下载模板填写。')
        by_id = {customer.id: customer for customer in customers}
        by_name = {}
        for customer in customers:
            by_name.setdefault(customer.name.strip().casefold(), []).append(customer)
        entries, errors, seen = [], [], set()
        for number, raw in enumerate(iterator, 2):
            if number > 3001:
                raise ValueError('每次最多导入 3000 行，请分批上传。')
            row = list(raw[:5]) + [None] * max(0, 5 - len(raw))
            if all(value is None or value == '' for value in row):
                continue
            try:
                if any(isinstance(value, str) and value.startswith('=') for value in row):
                    raise ValueError('请填写实际数值，不能使用公式。')
                customer_id, name, amount, cutoff, note = row
                name = str(name or '').strip()
                if customer_id is not None and str(customer_id).strip():
                    numeric_id = Decimal(str(customer_id).strip())
                    if not numeric_id.is_finite() or numeric_id != numeric_id.to_integral_value():
                        raise ValueError('客户编号必须是整数。')
                    customer = by_id.get(int(numeric_id))
                    if customer is None:
                        raise ValueError('客户编号不存在或客户已删除。')
                    if name and customer.name.strip().casefold() != name.casefold():
                        raise ValueError('客户编号与名称不一致。')
                else:
                    matches = by_name.get(name.casefold(), [])
                    if len(matches) != 1:
                        raise ValueError('客户名称未匹配或存在同名客户，请填写客户编号。')
                    customer = matches[0]
                if customer.id in seen:
                    raise ValueError('文件内同一客户重复出现。')
                seen.add(customer.id)
                values = history_values(amount, cutoff, note)
                values['historical_deal_cutoff'] = values['historical_deal_cutoff'].isoformat()
                entries.append(dict(values, id=customer.id, name=customer.name,
                                    version=customer.version, previous=customer.historical_deal_usd or 0,
                                    system=customer.total_deal_usd or 0))
            except (ValueError, InvalidOperation) as exc:
                errors.append(f'第 {number} 行：{exc}')
        if not entries and not errors:
            raise ValueError('文件没有可导入的数据。')
        return entries, errors
    finally:
        wb.close()
