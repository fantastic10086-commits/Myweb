"""Admin monthly reconciliation, independent of operational profit accounting."""
import os
import hashlib
import json
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from sqlalchemy import or_

from flask import abort, flash, redirect, render_template, request, send_file, url_for, current_app
from werkzeug.utils import secure_filename
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload
from sqlalchemy.orm.exc import StaleDataError
from models import db, PI, Payment, Expense, FinanceReview, FinanceAccount, AuditLog

STATUS = {'unconfirmed': '付款待确认', 'unpaid': '未支付', 'paid': '已支付'}


def register_finance_web(app, *, admin_required, current_user, audit):
    @app.after_request
    def protect_finance_cache(response):
        if request.path.startswith('/admin/finance'):
            response.headers['Cache-Control'] = 'private, no-store'
        return response

    def signature(record, kind):
        fields = ['pi_id', 'amount', 'attachment', 'note']
        fields += ['fee', 'order_no', 'receiving_account_name', 'receiving_account_currency'] if kind == 'payment' else ['currency', 'category']
        values = {key: getattr(record, key) for key in fields}
        return hashlib.sha256(json.dumps(values, sort_keys=True, ensure_ascii=False).encode()).hexdigest()

    def get_rows():
        kind = request.args.get('tab', 'payment')
        if kind not in ('payment', 'expense'):
            kind = 'payment'
        month = request.args.get('month', date.today().strftime('%Y-%m'))
        if month:
            try:
                datetime.strptime(month, '%Y-%m')
            except ValueError:
                abort(400, '月份格式错误')
        basis = request.args.get('basis', 'occurred')
        if basis not in ('occurred', 'paid'):
            basis = 'occurred'
        if kind == 'payment':
            basis = 'occurred'
        keyword = request.args.get('q', '').strip()[:200]
        salesperson = request.args.get('salesperson', '').strip()[:100]
        account = request.args.get('account', '').strip()[:200]
        category = request.args.get('category', '').strip()[:200]
        status = request.args.get('status', '')
        checked = request.args.get('checked', '')
        missing = request.args.get('missing') == '1'
        reviews = {(r.kind, r.record_id): r for r in FinanceReview.query.filter_by(kind=kind).all()}
        model = Payment if kind == 'payment' else Expense
        query = model.query.options(joinedload(model.pi).joinedload(PI.customer))
        if kind == 'expense':
            query = query.outerjoin(PI).filter(or_(Expense.pi_id.is_(None), PI.deleted_at.is_(None)))
        else:
            query = query.join(PI).filter(PI.deleted_at.is_(None))
        if kind == 'payment':
            query = query.filter(Payment.deleted_at.is_(None))
        records = query.order_by(model.created_at.desc(), model.id.desc()).all()
        creators = {}
        for log in AuditLog.query.filter_by(entity_type=kind, action='create').order_by(AuditLog.id).all():
            creators.setdefault(log.entity_id, log.username)
        rows = []
        for record in records:
            pi = record.pi or SimpleNamespace(id=None, pi_number="未关联PI", customer=None, salesperson="", freight_forwarder="")
            review = reviews.get((kind, record.id))
            occurred = review.occurred_on if review and review.occurred_on else ((record.created_at + timedelta(hours=8)).date() if record.created_at else None)
            paid = review.paid_on if review else None
            paid_status = review.payment_status if review else 'unconfirmed'
            account_name = (record.receiving_account_name or '未记录账户') if kind == 'payment' else (review.account_name if review else '')
            actual_date = paid if basis == 'paid' else occurred
            if month and (not actual_date or actual_date.strftime('%Y-%m') != month):
                continue
            if keyword and keyword.casefold() not in (' '.join([pi.pi_number, pi.customer.name if pi.customer else '', record.order_no if kind == 'payment' else record.note or ''])).casefold():
                continue
            if salesperson and pi.salesperson != salesperson:
                continue
            if account and account != account_name:
                continue
            if kind == 'expense' and category and category != record.category:
                continue
            if kind == 'expense' and status in STATUS and status != paid_status:
                continue
            is_checked = bool(review and review.checked and review.source_signature == signature(record, kind))
            if checked in ('yes', 'no') and is_checked != (checked == 'yes'):
                continue
            if missing and record.attachment:
                continue
            currency = (record.receiving_account_currency or pi.currency or 'USD') if kind == 'payment' else (record.currency or '未标币种')
            row = dict(id=record.id, kind=kind, record=record, pi=pi, review=review, creator=creators.get(record.id, '未记录'),
                       occurred=occurred, paid=paid, status=paid_status, account=account_name,
                       currency=currency, checked=is_checked, estimated=not (review and review.occurred_on),
                       attachment=url_for('finance_attachment', kind=kind, id=record.id) if record.attachment else '',
                       edit=dict(kind=kind, id=record.id, version=review.version if review else 0, source_signature=signature(record, kind),
                                 occurred_on=review.occurred_on.isoformat() if review and review.occurred_on else '',
                                 payment_status=paid_status, paid_on=paid.isoformat() if paid else '',
                                 account_id=review.account_id if review else '', checked=is_checked,
                                 note=review.note if review else '', label=pi.pi_number))
            rows.append(row)
        return kind, rows, dict(month=month, basis=basis, q=keyword, salesperson=salesperson,
                                account=account, category=category, status=status, checked=checked, missing=missing)

    @app.get('/admin/finance')
    @admin_required
    def finance_reconciliation():
        kind, rows, filters = get_rows()
        totals = defaultdict(float)
        fees = defaultdict(float)
        account_totals = defaultdict(float)
        paid_totals = defaultdict(float)
        unpaid_totals = defaultdict(float)
        for row in rows:
            currency = row['currency']
            totals[currency] += row['record'].amount or 0
            if kind == 'payment' or row['status'] == 'paid':
                account_totals[(row['account'] or '未记录账户', currency)] += row['record'].amount or 0
            if kind == 'payment':
                fees[currency] += row['record'].fee or 0
            elif row['status'] == 'paid':
                paid_totals[currency] += row['record'].amount or 0
            elif row['status'] == 'unpaid':
                unpaid_totals[currency] += row['record'].amount or 0
        page = max(1, request.args.get('page', 1, type=int) or 1)
        pages = max(1, (len(rows) + 49) // 50)
        page = min(page, pages)
        def link(**changes):
            args = request.args.to_dict()
            args.update(changes)
            return url_for('finance_reconciliation', **args)
        return render_template('finance.html', kind=kind, rows=rows[(page-1)*50:page*50],
                               filters=filters, totals=totals, fees=fees, account_totals=account_totals, paid_totals=paid_totals, unpaid_totals=unpaid_totals,
                               count=len(rows), missing_count=sum(not row['attachment'] for row in rows),
                               unchecked_count=sum(not row['checked'] for row in rows),
                               unconfirmed_count=sum(row['status'] == 'unconfirmed' for row in rows) if kind == 'expense' else 0,
                               page=page, pages=pages, link=link, statuses=STATUS,
                               local_time=lambda value: (value + timedelta(hours=8)).strftime('%Y-%m-%d %H:%M') if value else '',
                               accounts=FinanceAccount.query.order_by(FinanceAccount.name).all(),
                               receipt_accounts=sorted({x[0] for x in db.session.query(Payment.receiving_account_name).filter(Payment.deleted_at.is_(None)).distinct() if x[0]}),
                               salespeople=sorted({x[0] for x in db.session.query(PI.salesperson).filter(PI.deleted_at.is_(None)).distinct() if x[0]}),
                               categories=sorted({x[0] for x in db.session.query(Expense.category).distinct() if x[0]}),
                               export_url=url_for('finance_export', **request.args.to_dict()))

    def snapshot(review):
        return {key: getattr(review, key) for key in ('kind', 'record_id', 'occurred_on', 'payment_status', 'paid_on', 'account_id', 'account_name', 'checked', 'checked_by', 'checked_at', 'note', 'source_signature', 'version')}

    @app.post('/admin/finance/review')
    @admin_required
    def finance_review_save():
        kind = request.form.get('kind')
        if kind not in ('expense', 'payment'):
            abort(400)
        model = Expense if kind == 'expense' else Payment
        record = db.session.get(model, request.form.get('id', type=int))
        if not record or (record.pi and record.pi.deleted_at) or (kind == 'payment' and (not record.pi or record.deleted_at)):
            abort(404)
        def go(message, category='danger'):
            flash(message, category)
            # Keep only known, local query parameters supplied by the form.
            return redirect(url_for('finance_reconciliation', tab=kind,
                                    **{k: request.form.get('filter_'+k, '') for k in ('month', 'basis', 'q', 'salesperson', 'account', 'category', 'status', 'checked', 'missing', 'page')}))
        try:
            def parse_date(field):
                value = request.form.get(field, '').strip()
                return date.fromisoformat(value) if value else None
            occurred = parse_date('occurred_on')
            paid = parse_date('paid_on')
        except ValueError:
            return go('日期格式不正确。')
        status = request.form.get('payment_status', 'unconfirmed')
        account = db.session.get(FinanceAccount, request.form.get('account_id', type=int)) if request.form.get('account_id', type=int) else None
        if kind == 'expense':
            if status not in STATUS:
                return go('请选择有效的付款状态。')
            if status == 'paid' and (not paid or not account):
                return go('已支付费用必须填写支付日期和付款账户。')
            if status != 'paid':
                paid, account = None, None
        else:
            status, paid, account = 'unconfirmed', None, None
        review = FinanceReview.query.filter_by(kind=kind, record_id=record.id).first()
        expected = request.form.get('version', type=int)
        if expected is None or expected != (review.version if review else 0):
            return go('记录已被其他管理员修改，请刷新后重新核对。')
        if request.form.get('source_signature') != signature(record, kind):
            return go('原始金额或凭证已变化，请刷新后重新核对。')
        before = snapshot(review) if review else {}
        if not review:
            review = FinanceReview(kind=kind, record_id=record.id)
            db.session.add(review)
        review.occurred_on = occurred
        review.payment_status, review.paid_on = status, paid
        review.account_id = account.id if account else None
        review.account_name = account.name if account else ''
        review.note = request.form.get('note', '').strip()[:1000]
        review.checked = request.form.get('checked') == '1'
        review.source_signature = signature(record, kind)
        review.checked_by = current_user().username if review.checked else ''
        review.checked_at = datetime.utcnow() if review.checked else None
        try:
            db.session.flush()
            audit('update', 'finance_review', review.id, f'财务核对 {record.pi.pi_number if record.pi else "未关联PI"} {kind} #{record.id}', before=before, after=snapshot(review))
            db.session.commit()
        except (StaleDataError, IntegrityError):
            db.session.rollback()
            return go('记录已被其他管理员修改，请刷新后重新核对。')
        return go('财务信息已保存。', 'success')

    @app.post('/admin/finance/accounts')
    @admin_required
    def finance_account_add():
        name = request.form.get('name', '').strip()
        if not name or len(name) > 200:
            abort(400, '请填写不超过200字的账户名称')
        account = FinanceAccount(name=name)
        try:
            db.session.add(account)
            db.session.flush()
            audit('create', 'finance_account', account.id, '新增付款账户', after={'name': name})
            db.session.commit()
            flash('付款账户已新增。', 'success')
        except IntegrityError:
            db.session.rollback()
            flash('该付款账户已存在。', 'warning')
        return redirect(url_for('finance_reconciliation', tab='expense', month=request.form.get('month', '')))

    @app.get('/admin/finance/<kind>/<int:id>/attachment')
    @admin_required
    def finance_attachment(kind, id):
        if kind not in ('payment', 'expense'):
            abort(404)
        record = db.session.get(Payment if kind == 'payment' else Expense, id)
        if not record or (record.pi and record.pi.deleted_at) or (kind == 'payment' and record.deleted_at):
            abort(404)
        filename = secure_filename(record.attachment or '')
        if not filename or filename != record.attachment:
            abort(404)
        path = os.path.join(current_app.config['UPLOAD_DIR'], filename)
        if not os.path.isfile(path):
            abort(404)
        return send_file(path)

    @app.get('/admin/finance/export')
    @admin_required
    def finance_export():
        from openpyxl import Workbook
        kind, rows, filters = get_rows()
        wb = Workbook()
        ws = wb.active
        ws.title = '回款明细' if kind == 'payment' else '费用明细'
        headers = ['PI编号', '客户', '业务员', '到账日期' if kind == 'payment' else '费用发生日期', '日期来源', '币种', '金额', '手续费' if kind == 'payment' else '费用类别', '流水号' if kind == 'payment' else '付款状态', '到账账户' if kind == 'payment' else '付款账户', '支付日期', '凭证链接', '核对状态', '核对人', '核对时间(UTC)', '核对备注', '原始备注', '登记人']
        ws.append(headers)
        for row in rows:
            rec, review, pi = row['record'], row['review'], row['pi']
            ws.append([pi.pi_number, pi.customer.name if pi.customer else '', pi.salesperson,
                       row['occurred'].isoformat() if row['occurred'] else '', '登记日期代替' if row['estimated'] else '管理员确认',
                       row['currency'], rec.amount, rec.fee if kind == 'payment' else rec.category,
                       rec.order_no if kind == 'payment' else STATUS[row['status']], row['account'], row['paid'].isoformat() if row['paid'] else '',
                       request.url_root.rstrip('/') + row['attachment'] if row['attachment'] else '',
                       '已核对' if row['checked'] else '待核对', review.checked_by if review else '',
                       review.checked_at.isoformat() if review and review.checked_at else '', review.note if review else '', rec.note, row['creator']])
        # Treat user-controlled strings as literal text, never spreadsheet formulas.
        for line in ws:
            for cell in line:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        from openpyxl.utils import get_column_letter
        for i in range(1, len(headers)+1):
            ws.column_dimensions[get_column_letter(i)].width = 24
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f'财务核对_{kind}_{filters["month"] or "全部"}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
