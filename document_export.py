"""Safe Excel-template rendering and PDF conversion for PI exports."""

import copy
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels
from openpyxl.worksheet.page import PageMargins


class TemplateError(ValueError):
    pass


PLACEHOLDER_GROUPS = {
    'PI 信息': [
        ('pi_number', 'PI 编号'),
        ('issue_date', '开单日期'),
        ('salesperson', '业务员'),
        ('salesperson_phone', '业务员电话'),
        ('salesperson_wechat', '业务员微信（同电话）'),
        ('salesperson_whatsapp', '业务员 WhatsApp（同电话）'),
        ('salesperson_email', '业务员邮箱'),
        ('currency', '币种代码（USD/RMB）'),
        ('currency_symbol', '币种符号（$/¥）'),
        ('payment_terms', '付款条款'),
        ('price_terms', '价格条款'),
        ('delivery_time', '交货期'),
        ('bank_info', '银行信息'),
        ('bank_beneficiary_name', '收款人名称'),
        ('bank_account_no', '收款账号'),
        ('bank_country_region', '收款人国家/地区'),
        ('bank_beneficiary_address', '收款人地址'),
        ('bank_name', '银行名称'),
        ('bank_address', '银行地址'),
        ('bank_swift_code', 'SWIFT 代码'),
        ('bank_code', 'Bank Code'),
        ('bank_branch_code', 'Branch Code'),
        ('bank_currency', '收款账户币种'),
        ('shipping_address', '收货地址'),
        ('notes', 'PI 备注'),
    ],
    '公司与客户': [
        ('company_name', '公司名称'),
        ('company_address', '公司地址'),
        ('customer_name', '客户名称'),
        ('customer_contact', '联系人'),
        ('customer_country', '国家'),
        ('customer_email', '客户邮箱'),
        ('customer_phone', '客户电话'),
        ('customer_address', '客户地址'),
        ('customer_notes', '客户备注'),
    ],
    '金额': [
        ('product_subtotal', '产品总金额（数值）'),
        ('other_charges', '其他费用/折扣（数值）'),
        ('shipping_note', '客户费用/折扣英文名称（兼容字段）'),
        ('shipping_note_zh', '客户费用/折扣中文名称'),
        ('shipping_note_en', '客户费用/折扣英文名称'),
        ('grand_total', '订单总金额（数值）'),
        ('total_quantity', '产品总数量（数值）'),
        ('hs_code', '海关编码（固定）'),
        ('packing', '包装方式（固定）'),
        ('place_of_loading', '装货地（固定）'),
        ('origin', '原产地（固定）'),
    ],
    '产品明细行': [
        ('item.no', '序号'),
        ('item.image', '产品图片'),
        ('item.code', '产品编码'),
        ('item.name', '产品名称'),
        ('item.specification', '规格'),
        ('item.quantity', '数量（数值）'),
        ('item.unit_price', '单价（数值）'),
        ('item.amount', '金额（数值）'),
    ],
}

SUPPORTED_PLACEHOLDERS = {
    key for pairs in PLACEHOLDER_GROUPS.values() for key, _label in pairs
}
_TOKEN_RE = re.compile(r'\{\{\s*([a-zA-Z0-9_.]+)\s*\}\}')
_EXACT_TOKEN_RE = re.compile(r'^\s*\{\{\s*([a-zA-Z0-9_.]+)\s*\}\}\s*$')


def _company_name(pi):
    override = getattr(pi, '_company_name_override', '')
    if override:
        return override
    if (getattr(pi, 'company', '') or '').lower() == 'qisuo':
        return 'Changzhou Qisuo welding and cutting Equipment Co., LTD'
    return 'CHANGZHOU KLISTA INTERNATIONAL TRADE CO., LTD.'


def _company_address(pi):
    override = getattr(pi, '_company_addr_override', '')
    if override:
        return override
    if (getattr(pi, 'company', '') or '').lower() == 'qisuo':
        return (
            'No. 158, Jingchuang Road, Yaoguan Town, Changzhou District, '
            'Jiangsu Province'
        )
    return 'No. 158 Jinchuang Road, Yaoguan Town, Changzhou City, China'


def _bank_snapshot_value(pi, attribute, labels=()):
    """Read a structured bank snapshot with a legacy bank-info fallback."""
    value = (getattr(pi, attribute, '') or '').strip()
    if value:
        return value
    bank_info = getattr(pi, 'bank_info', '') or ''
    for line in bank_info.splitlines():
        left, separator, right = line.partition(':')
        if separator and left.strip().casefold() in {
            label.casefold() for label in labels
        }:
            return right.strip()
    return ''


def _set_safe_excel_text(cell, value):
    """Write user text as an explicit string, never as an Excel formula."""
    text = '' if value is None else str(value)
    cell.value = text
    if text.startswith(('=', '+', '-', '@')):
        # openpyxl interprets a leading equals sign as a formula by default.
        # Marking the OOXML cell as a string prevents execution without adding
        # a visible apostrophe to the PDF produced by LibreOffice.
        cell.data_type = 's'


def _fixed_values(pi, salesperson_info):
    customer = getattr(pi, 'customer', None)
    currency = (getattr(pi, 'currency', 'USD') or 'USD').upper()
    subtotal = round(float(getattr(pi, 'total_amount', 0) or 0), 2)
    adjustment = round(float(getattr(pi, 'shipping_cost', 0) or 0), 2)
    salesperson_phone = (salesperson_info or {}).get('phone', '')
    items = list(getattr(pi, 'items', []) or [])
    return {
        'pi_number': getattr(pi, 'pi_number', '') or '',
        'issue_date': getattr(pi, 'issue_date', None) or '',
        'salesperson': getattr(pi, 'salesperson', '') or '',
        'salesperson_phone': salesperson_phone,
        'salesperson_wechat': salesperson_phone,
        'salesperson_whatsapp': salesperson_phone,
        'salesperson_email': (salesperson_info or {}).get('email', ''),
        'currency': currency,
        'currency_symbol': '¥' if currency == 'RMB' else '$',
        'company_name': _company_name(pi),
        'company_address': _company_address(pi),
        'customer_name': getattr(customer, 'name', '') or '',
        'customer_contact': getattr(customer, 'contact_person', '') or '',
        'customer_country': getattr(customer, 'country', '') or '',
        'customer_email': getattr(customer, 'email', '') or '',
        'customer_phone': getattr(customer, 'phone', '') or '',
        'customer_address': getattr(customer, 'address', '') or '',
        'customer_notes': getattr(customer, 'notes', '') or '',
        'payment_terms': getattr(pi, 'payment_terms', '') or '',
        'price_terms': getattr(pi, 'price_terms', '') or '',
        'delivery_time': getattr(pi, 'delivery_time', '') or '',
        'bank_info': getattr(pi, 'bank_info', '') or '',
        'bank_beneficiary_name': _bank_snapshot_value(
            pi, 'bank_beneficiary_name', ('Beneficiary', 'Company')
        ),
        'bank_account_no': _bank_snapshot_value(
            pi, 'bank_account_no', ('A/C', 'Account', 'Account No')
        ),
        'bank_country_region': _bank_snapshot_value(
            pi, 'bank_country_region', ('Country/Region', 'Country')
        ),
        'bank_beneficiary_address': _bank_snapshot_value(
            pi, 'bank_beneficiary_address', ('Beneficiary Address',)
        ),
        'bank_name': _bank_snapshot_value(pi, 'bank_name', ('Bank Name', 'Bank')),
        'bank_address': _bank_snapshot_value(pi, 'bank_address', ('Bank Address',)),
        'bank_swift_code': _bank_snapshot_value(
            pi, 'bank_swift_code', ('SWIFT', 'SWIFT Code')
        ),
        'bank_code': _bank_snapshot_value(pi, 'bank_code', ('Bank Code',)),
        'bank_branch_code': _bank_snapshot_value(
            pi, 'bank_branch_code', ('Branch Code',)
        ),
        'bank_currency': _bank_snapshot_value(pi, 'bank_currency', ('Currency',)),
        'shipping_address': getattr(pi, 'shipping_address', '') or '',
        'notes': getattr(pi, 'notes', '') or '',
        'product_subtotal': subtotal,
        'other_charges': adjustment,
        'shipping_note': (
            getattr(pi, 'shipping_note_en', '')
            or getattr(pi, 'shipping_note', '') or ''
        ),
        'shipping_note_zh': getattr(pi, 'shipping_note', '') or '',
        'shipping_note_en': (
            getattr(pi, 'shipping_note_en', '')
            or getattr(pi, 'shipping_note', '') or ''
        ),
        'grand_total': round(subtotal + adjustment, 2),
        'total_quantity': sum(int(getattr(item, 'quantity', 0) or 0) for item in items),
        'hs_code': '8518900090',
        'packing': 'CARTON',
        'place_of_loading': 'Changzhou',
        'origin': 'China',
    }


def _item_values(item, index):
    product = getattr(item, 'product', None)
    return {
        'item.no': index,
        'item.image': getattr(product, 'image', '') or '',
        'item.code': getattr(item, 'display_code', getattr(product, 'product_code', '')) or '',
        'item.name': getattr(item, 'display_name', getattr(product, 'name', '')) or '',
        'item.specification': getattr(item, 'display_specification', getattr(product, 'specification', '')) or '',
        'item.quantity': int(getattr(item, 'quantity', 0) or 0),
        'item.unit_price': round(float(getattr(item, 'unit_price', 0) or 0), 2),
        'item.amount': round(float(getattr(item, 'amount', 0) or 0), 2),
    }


def scan_template(path):
    """Return placeholders and the sheets/rows containing item placeholders."""
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            expanded_size = sum(member.file_size for member in members)
            if len(members) > 2000 or expanded_size > 100 * 1024 * 1024:
                raise TemplateError('模板解压后的内容过大，已为安全起见拒绝上传。')
    except TemplateError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise TemplateError('文件不是有效的 .xlsx 模板。') from exc
    try:
        workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    except Exception as exc:
        raise TemplateError('Excel 模板无法读取，请确认文件未损坏且为 .xlsx 格式。') from exc

    placeholders = set()
    item_rows = []
    try:
        if len(workbook.worksheets) > 20:
            raise TemplateError('模板工作表不能超过 20 个。')
        for sheet in workbook.worksheets:
            if sheet.max_row * sheet.max_column > 100000:
                raise TemplateError(f'工作表“{sheet.title}”范围过大，请删除多余空白行列。')
            sheet_item_rows = set()
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    found = set(_TOKEN_RE.findall(value))
                    placeholders.update(found)
                    if any(token.startswith('item.') for token in found):
                        sheet_item_rows.add(cell.row)
            for row_number in sorted(sheet_item_rows):
                item_rows.append((sheet.title, row_number))
    finally:
        workbook.close()
    return placeholders, item_rows


def validate_template(path):
    placeholders, item_rows = scan_template(path)
    unknown = sorted(placeholders - SUPPORTED_PLACEHOLDERS)
    if unknown:
        raise TemplateError('模板含有不支持的占位符：' + '、'.join('{{' + key + '}}' for key in unknown))
    if 'pi_number' not in placeholders:
        raise TemplateError('模板至少需要包含 {{pi_number}}。')
    if not any(token.startswith('item.') for token in placeholders):
        raise TemplateError('模板至少需要一行产品明细占位符，例如 {{item.name}}。')
    per_sheet = {}
    for sheet_name, row_number in item_rows:
        per_sheet.setdefault(sheet_name, []).append(row_number)
    duplicated = [name for name, rows in per_sheet.items() if len(rows) > 1]
    if duplicated:
        raise TemplateError('每个工作表只能设置一行产品明细模板行：' + '、'.join(duplicated))
    return placeholders


def create_placeholder_template(output_path):
    """Create an administrator-editable starter workbook with placeholders."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'PI Template'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)

    dark_fill = PatternFill('solid', fgColor='1A3A5C')
    light_fill = PatternFill('solid', fgColor='EAF0F6')
    white_bold = Font(name='Arial', bold=True, color='FFFFFF', size=9)
    bold = Font(name='Arial', bold=True, size=9)
    normal = Font(name='Arial', size=9)
    centered = Alignment(horizontal='center', vertical='center', wrap_text=True)
    wrapped = Alignment(vertical='top', wrap_text=True)
    thin = Side(style='thin', color='B9C3CC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    widths = [12, 13, 18, 22, 14, 9, 14, 15]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.merge_cells('A1:H1')
    sheet['A1'] = '{{company_name}}'
    sheet['A1'].font = Font(name='Arial', bold=True, color='FFFFFF', size=16)
    sheet['A1'].fill = dark_fill
    sheet['A1'].alignment = centered
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells('A2:H2')
    sheet['A2'] = '{{company_address}}'
    sheet['A2'].font = Font(name='Arial', color='FFFFFF', size=8)
    sheet['A2'].fill = dark_fill
    sheet['A2'].alignment = centered
    sheet.merge_cells('A4:H4')
    sheet['A4'] = 'PROFORMA INVOICE'
    sheet['A4'].font = Font(name='Arial', bold=True, color='1A3A5C', size=16)
    sheet['A4'].alignment = centered

    info_rows = [
        ('A6', 'PI Number:', 'B6', '{{pi_number}}', 'E6', 'To / Buyer:', 'F6', '{{customer_name}}'),
        ('A7', 'Date:', 'B7', '{{issue_date}}', 'E7', 'Contact:', 'F7', '{{customer_contact}}'),
        ('A8', 'Salesperson:', 'B8', '{{salesperson}}', 'E8', 'Country:', 'F8', '{{customer_country}}'),
        ('A9', 'Email:', 'B9', '{{salesperson_email}}', 'E9', 'Email:', 'F9', '{{customer_email}}'),
    ]
    for left_label, left_value, left_data, left_placeholder, right_label, right_value, right_data, right_placeholder in info_rows:
        sheet[left_label] = left_value
        sheet[left_label].font = bold
        sheet[left_data] = left_placeholder
        sheet[left_data].font = normal
        sheet[left_data].alignment = Alignment(horizontal='left', vertical='center')
        sheet.merge_cells(start_row=sheet[left_data].row, start_column=sheet[left_data].column,
                          end_row=sheet[left_data].row, end_column=4)
        sheet[right_label] = right_value
        sheet[right_label].font = bold
        sheet[right_data] = right_placeholder
        sheet[right_data].font = normal
        sheet[right_data].alignment = Alignment(horizontal='left', vertical='center')
        sheet.merge_cells(start_row=sheet[right_data].row, start_column=sheet[right_data].column,
                          end_row=sheet[right_data].row, end_column=8)

    headers = ['No.', 'Image', 'Product Code', 'Description', 'Specification', 'Qty', 'Unit Price', 'Amount']
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(11, column, label)
        cell.font = white_bold
        cell.fill = dark_fill
        cell.alignment = centered
        cell.border = border
    item_tokens = [
        '{{item.no}}', '{{item.image}}', '{{item.code}}', '{{item.name}}',
        '{{item.specification}}', '{{item.quantity}}', '{{item.unit_price}}', '{{item.amount}}',
    ]
    for column, token in enumerate(item_tokens, 1):
        cell = sheet.cell(12, column, token)
        cell.font = normal
        cell.alignment = centered if column != 4 else wrapped
        cell.border = border
        if column in (7, 8):
            cell.number_format = '#,##0.00'
    sheet.row_dimensions[12].height = 48

    totals = [
        (14, 'PRODUCT SUBTOTAL:', '{{product_subtotal}}'),
        (15, 'OTHER CHARGES:', '{{other_charges}}'),
        (16, 'TOTAL AMOUNT:', '{{grand_total}}'),
    ]
    for row_number, label, token in totals:
        sheet.merge_cells(start_row=row_number, start_column=5, end_row=row_number, end_column=7)
        sheet.cell(row_number, 5, label).font = bold
        sheet.cell(row_number, 5).alignment = Alignment(horizontal='right')
        sheet.cell(row_number, 8, token).font = bold if row_number == 16 else normal
        sheet.cell(row_number, 8).number_format = '#,##0.00'
        if row_number == 16:
            for column in range(5, 9):
                sheet.cell(row_number, column).fill = light_fill

    detail_rows = [
        (18, 'Payment Terms:', '{{payment_terms}}'),
        (19, 'Bank Info:', '{{bank_info}}'),
        (20, 'Shipping Address:', '{{shipping_address}}'),
        (21, 'Customer Notes:', '{{customer_notes}}'),
        (22, 'PI Notes:', '{{notes}}'),
    ]
    for row_number, label, token in detail_rows:
        sheet.cell(row_number, 1, label).font = bold
        sheet.cell(row_number, 1).alignment = wrapped
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=2)
        sheet.merge_cells(start_row=row_number, start_column=3, end_row=row_number, end_column=8)
        sheet.cell(row_number, 3, token).font = normal
        sheet.cell(row_number, 3).alignment = wrapped
    sheet.row_dimensions[19].height = 42
    sheet.row_dimensions[20].height = 30
    sheet.row_dimensions[21].height = 30
    sheet.row_dimensions[22].height = 30
    sheet.print_area = 'A1:H22'
    sheet.freeze_panes = 'A11'
    workbook.save(output_path)
    workbook.close()
    return output_path


def _copy_row(sheet, source_row, destination_row):
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(destination_row, column)
        target.value = source.value
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.alignment = copy.copy(source.alignment)
        target.protection = copy.copy(source.protection)
    source_dimension = sheet.row_dimensions[source_row]
    target_dimension = sheet.row_dimensions[destination_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden


def _expand_item_row(sheet, row_number, item_count):
    count = max(1, item_count)
    insert_at = row_number + 1
    shift = count - 1
    merged_on_row = [
        copy.copy(cell_range) for cell_range in sheet.merged_cells.ranges
        if cell_range.min_row == row_number and cell_range.max_row == row_number
    ]
    if shift > 0:
        # openpyxl moves cell values when inserting rows, but not the related
        # merged ranges, row heights, or print area. Move those structures too.
        shifted_merges = []
        for cell_range in list(sheet.merged_cells.ranges):
            if cell_range.max_row < insert_at:
                continue
            sheet.unmerge_cells(str(cell_range))
            min_row = cell_range.min_row + shift if cell_range.min_row >= insert_at else cell_range.min_row
            max_row = cell_range.max_row + shift
            shifted_merges.append((cell_range.min_col, min_row, cell_range.max_col, max_row))

        shifted_dimensions = {
            index: copy.copy(dimension)
            for index, dimension in list(sheet.row_dimensions.items())
            if index >= insert_at
        }
        for index in shifted_dimensions:
            del sheet.row_dimensions[index]

        original_print_area = sheet.print_area
        sheet.insert_rows(insert_at, shift)

        for index, dimension in shifted_dimensions.items():
            new_index = index + shift
            dimension.index = new_index
            sheet.row_dimensions[new_index] = dimension
        for min_col, min_row, max_col, max_row in shifted_merges:
            start = f'{get_column_letter(min_col)}{min_row}'
            end = f'{get_column_letter(max_col)}{max_row}'
            sheet.merge_cells(f'{start}:{end}')

        if original_print_area:
            sheet.print_area = re.sub(
                r'(\$?[A-Z]{1,3}\$?)(\d+)',
                lambda match: match.group(1) + str(
                    int(match.group(2)) + shift
                    if int(match.group(2)) >= insert_at else int(match.group(2))
                ),
                str(original_print_area),
            )

        for offset in range(1, count):
            _copy_row(sheet, row_number, row_number + offset)
            for cell_range in merged_on_row:
                start = f'{get_column_letter(cell_range.min_col)}{row_number + offset}'
                end = f'{get_column_letter(cell_range.max_col)}{row_number + offset}'
                sheet.merge_cells(f'{start}:{end}')
    return [row_number + offset for offset in range(count)]


def _fit_image_size(width, height, max_width, max_height, allow_upscale=False):
    """Fit an image inside a box without changing its aspect ratio."""
    width = float(width or 0)
    height = float(height or 0)
    max_width = max(float(max_width or 0), 1.0)
    max_height = max(float(max_height or 0), 1.0)
    if width <= 0 or height <= 0:
        return max_width, max_height
    scale = min(max_width / width, max_height / height)
    if not allow_upscale:
        scale = min(scale, 1.0)
    return max(width * scale, 1.0), max(height * scale, 1.0)


def _column_width_to_pixels(width):
    """Approximate Excel column width in pixels using Excel's own scale."""
    width = float(width or 0)
    if width < 1:
        return int(width * 12 + 0.5)
    return int(width * 7 + 5)


def _image_cell_bounds(sheet, cell):
    """Return the anchor and pixel bounds for a cell or its merged range."""
    min_col = max_col = cell.column
    min_row = max_row = cell.row
    for merged_range in sheet.merged_cells.ranges:
        if cell.coordinate in merged_range:
            min_col = merged_range.min_col
            max_col = merged_range.max_col
            min_row = merged_range.min_row
            max_row = merged_range.max_row
            break

    default_column_width = sheet.sheet_format.defaultColWidth or 8.43
    width = 0
    for column in range(min_col, max_col + 1):
        column_letter = get_column_letter(column)
        configured_width = sheet.column_dimensions[column_letter].width
        width += _column_width_to_pixels(configured_width or default_column_width)

    default_row_height = sheet.sheet_format.defaultRowHeight or 15
    height = 0
    for row in range(min_row, max_row + 1):
        configured_height = sheet.row_dimensions[row].height
        height += points_to_pixels(configured_height or default_row_height)

    return min_col, min_row, max(float(width), 1.0), max(float(height), 1.0)


def _put_image(sheet, cell, image_name, upload_dir):
    cell.value = ''
    if not image_name or not upload_dir:
        return
    image_path = os.path.join(upload_dir, os.path.basename(str(image_name)))
    if not os.path.isfile(image_path):
        return
    try:
        image = ExcelImage(image_path)
        anchor_col, anchor_row, cell_width, cell_height = _image_cell_bounds(
            sheet, cell
        )
        padding = 4.0
        max_width = max(cell_width - padding * 2, 1.0)
        max_height = max(cell_height - padding * 2, 1.0)
        image.width, image.height = _fit_image_size(
            image.width,
            image.height,
            max_width,
            max_height,
            allow_upscale=False,
        )
        offset_x = max((cell_width - image.width) / 2, 0)
        offset_y = max((cell_height - image.height) / 2, 0)
        marker = AnchorMarker(
            col=anchor_col - 1,
            colOff=pixels_to_EMU(offset_x),
            row=anchor_row - 1,
            rowOff=pixels_to_EMU(offset_y),
        )
        image.anchor = OneCellAnchor(
            _from=marker,
            ext=XDRPositiveSize2D(
                cx=pixels_to_EMU(image.width),
                cy=pixels_to_EMU(image.height),
            ),
        )
        sheet.add_image(image)
    except Exception:
        return


def _replace_cell(sheet, cell, values, upload_dir):
    if isinstance(cell, MergedCell) or not isinstance(cell.value, str):
        return
    original = cell.value
    matches = list(_TOKEN_RE.finditer(original))
    if not matches:
        return
    if cell.data_type == 'f':
        raise TemplateError(f'{sheet.title}!{cell.coordinate} 的公式中不能使用占位符。')

    exact = _EXACT_TOKEN_RE.match(original)
    if exact:
        key = exact.group(1)
        if key == 'item.image':
            _put_image(sheet, cell, values.get(key, ''), upload_dir)
            return
        value = values.get(key, '')
        if isinstance(value, datetime):
            cell.value = value
        elif isinstance(value, date):
            cell.value = value
            if cell.number_format == 'General':
                cell.number_format = 'yyyy-mm-dd'
        elif isinstance(value, (int, float)):
            cell.value = value
        else:
            _set_safe_excel_text(cell, value)
        return

    result = original
    for match in matches:
        key = match.group(1)
        if key == 'item.image':
            replacement = ''
        else:
            value = values.get(key, '')
            replacement = value.strftime('%Y-%m-%d') if isinstance(value, (date, datetime)) else str(value)
        result = result.replace(match.group(0), replacement)
    _set_safe_excel_text(cell, result)


def render_excel_template(template_path, pi, output_path, salesperson_info=None, upload_dir=''):
    """Fill a trusted .xlsx layout without modifying the source workbook."""
    validate_template(template_path)
    workbook = load_workbook(template_path, data_only=False, read_only=False, keep_links=False)
    fixed_values = _fixed_values(pi, salesperson_info or {})
    items = list(getattr(pi, 'items', []) or [])

    try:
        for sheet in workbook.worksheets:
            item_rows = set()
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and any(
                        key.startswith('item.') for key in _TOKEN_RE.findall(cell.value)
                    ):
                        item_rows.add(cell.row)
            if len(item_rows) > 1:
                raise TemplateError(f'工作表“{sheet.title}”只能有一行产品明细模板行。')

            expanded_rows = []
            if item_rows:
                template_row = next(iter(item_rows))
                expanded_rows = _expand_item_row(sheet, template_row, len(items))

            for item_index, row_number in enumerate(expanded_rows, 1):
                values = dict(fixed_values)
                if items:
                    values.update(_item_values(items[item_index - 1], item_index))
                else:
                    values.update(_item_values(None, item_index))
                for cell in sheet[row_number]:
                    _replace_cell(sheet, cell, values, upload_dir)

            expanded_set = set(expanded_rows)
            for row in sheet.iter_rows():
                if row and row[0].row in expanded_set:
                    continue
                for cell in row:
                    _replace_cell(sheet, cell, fixed_values, upload_dir)

            # Some workbooks created by WPS store print metadata that is lost
            # when openpyxl saves the populated copy.  Re-apply a stable print
            # area here so the generated Excel and its PDF conversion always
            # stay one page wide while allowing long item lists to flow down.
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.orientation = (
                sheet.page_setup.orientation or sheet.ORIENTATION_PORTRAIT
            )
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.print_area = (
                f'$A$1:${get_column_letter(max(1, sheet.max_column))}'
                f'${max(1, sheet.max_row)}'
            )

        if workbook.calculation is not None:
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.calculation.calcMode = 'auto'
        workbook.save(output_path)
    finally:
        workbook.close()
    return output_path


def apply_system_default_page_setup(excel_path):
    """Apply stable A4 print settings to the built-in editable workbook.

    The source template remains an ordinary, unlocked Excel file.  Print
    settings are applied to each generated copy so both the downloaded Excel
    file and the PDF converted from it share the same page geometry.
    """
    workbook = load_workbook(excel_path, data_only=False, read_only=False, keep_links=False)
    try:
        for sheet in workbook.worksheets:
            sheet.sheet_view.showGridLines = False
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.orientation = sheet.ORIENTATION_PORTRAIT
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_margins = PageMargins(
                left=0.28, right=0.28, top=0.30, bottom=0.38,
                header=0.12, footer=0.18,
            )
            sheet.print_options.horizontalCentered = True
            sheet.print_area = (
                f'$A$1:${get_column_letter(max(1, sheet.max_column))}'
                f'${max(1, sheet.max_row)}'
            )
            sheet.oddFooter.center.text = 'Page &P of &N'
            sheet.oddFooter.center.size = 8
            sheet.oddFooter.center.color = '6B7280'
        workbook.save(excel_path)
    finally:
        workbook.close()
    return excel_path


def find_soffice():
    for command in ('libreoffice', 'soffice'):
        found = shutil.which(command)
        if found:
            return found
    mac_path = '/Applications/LibreOffice.app/Contents/MacOS/soffice'
    return mac_path if os.path.isfile(mac_path) else ''


def convert_excel_to_pdf(excel_path, output_path):
    soffice = find_soffice()
    if not soffice:
        raise TemplateError('服务器尚未安装 Excel 转 PDF 组件，请联系管理员。')
    output_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(output_dir, exist_ok=True)
    profile_dir = tempfile.mkdtemp(prefix='pi-lo-profile-', dir=output_dir)
    try:
        result = subprocess.run(
            [
                soffice,
                '--headless',
                f'-env:UserInstallation=file://{profile_dir}',
                '--convert-to', 'pdf',
                '--outdir', output_dir,
                os.path.abspath(excel_path),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=90,
            check=False,
        )
        generated = os.path.join(
            output_dir, os.path.splitext(os.path.basename(excel_path))[0] + '.pdf'
        )
        if result.returncode != 0 or not os.path.isfile(generated):
            raise TemplateError('模板转换为 PDF 失败，请检查模板打印区域和页面设置。')
        if os.path.abspath(generated) != os.path.abspath(output_path):
            os.replace(generated, output_path)
    except subprocess.TimeoutExpired as exc:
        raise TemplateError('模板转换为 PDF 超时，请简化模板后重试。') from exc
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)
    return output_path
