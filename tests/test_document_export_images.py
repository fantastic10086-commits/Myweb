import os
import tempfile
import unittest

from PIL import Image as PillowImage
from openpyxl import Workbook, load_workbook
from openpyxl.utils.units import EMU_to_pixels

from document_export import _fit_image_size, _image_cell_bounds, _put_image
from packing_list_export import (
    COMPACT_PRODUCT_TEXT_UNITS,
    _compact_product_text,
)


class DocumentExportImageTests(unittest.TestCase):
    def test_compact_packing_product_text_is_truncated_only_when_too_long(self):
        self.assertEqual(
            _compact_product_text('P80 Torch', '8M cable'),
            'P80 Torch / 8M cable',
        )

        truncated = _compact_product_text(
            'Black Wolf P80 plasma cutting torch with extended handle',
            'Extra long product specification for carton label',
        )
        self.assertTrue(truncated.endswith('…'))
        self.assertLessEqual(len(truncated), COMPACT_PRODUCT_TEXT_UNITS)
        self.assertNotIn('Extra long product specification', truncated)

    def test_fit_preserves_wide_image_ratio(self):
        width, height = _fit_image_size(400, 100, 52, 48)

        self.assertAlmostEqual(width / height, 4.0)
        self.assertLessEqual(width, 52)
        self.assertLessEqual(height, 48)

    def test_fit_preserves_tall_image_ratio(self):
        width, height = _fit_image_size(100, 400, 52, 48)

        self.assertAlmostEqual(width / height, 0.25)
        self.assertLessEqual(width, 52)
        self.assertLessEqual(height, 48)

    def test_fit_does_not_enlarge_small_image(self):
        self.assertEqual(
            _fit_image_size(20, 10, 52, 48),
            (20.0, 10.0),
        )

    def test_inserted_image_is_centered_and_stays_inside_cell(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = os.path.join(temp_dir, 'wide.png')
            workbook_path = os.path.join(temp_dir, 'result.xlsx')
            PillowImage.new('RGB', (400, 100), 'red').save(image_path)

            workbook = Workbook()
            sheet = workbook.active
            sheet.column_dimensions['B'].width = 8
            sheet.row_dimensions[16].height = 42
            cell = sheet['B16']
            cell.value = '{{item.image}}'

            _put_image(sheet, cell, 'wide.png', temp_dir)

            self.assertEqual(len(sheet._images), 1)
            workbook.save(workbook_path)

            reloaded_sheet = load_workbook(workbook_path).active
            self.assertEqual(len(reloaded_sheet._images), 1)
            image = reloaded_sheet._images[0]
            anchor = image.anchor
            _, _, cell_width, cell_height = _image_cell_bounds(
                reloaded_sheet, reloaded_sheet['B16']
            )
            rendered_width = EMU_to_pixels(anchor.ext.cx)
            rendered_height = EMU_to_pixels(anchor.ext.cy)
            offset_x = EMU_to_pixels(anchor._from.colOff)
            offset_y = EMU_to_pixels(anchor._from.rowOff)

            self.assertEqual(anchor._from.col, 1)
            self.assertEqual(anchor._from.row, 15)
            self.assertAlmostEqual(rendered_width / rendered_height, 4.0, delta=0.1)
            self.assertGreaterEqual(offset_x, 0)
            self.assertGreaterEqual(offset_y, 0)
            self.assertLessEqual(offset_x + rendered_width, cell_width)
            self.assertLessEqual(offset_y + rendered_height, cell_height)


if __name__ == '__main__':
    unittest.main()
