"""Run with the project virtualenv: python tests/security_smoke.py."""

import os
import re
import shutil
import sys
import tempfile
import unittest
import zipfile
from datetime import date, timedelta
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

TEST_ROOT = tempfile.TemporaryDirectory(prefix='pi-manager-tests-')
for name in ('instance', 'uploads', 'pdf', 'backups'):
    os.makedirs(os.path.join(TEST_ROOT.name, name), exist_ok=True)

os.environ.update({
    'FLASK_ENV': 'production',
    'SECRET_KEY': 'test-only-secret-key-never-use-in-production',
    'INITIAL_ADMIN_PASSWORD': 'InitialTestPass123!',
    'DATABASE_DIR': os.path.join(TEST_ROOT.name, 'instance'),
    'UPLOAD_DIR': os.path.join(TEST_ROOT.name, 'uploads'),
    'PDF_DIR': os.path.join(TEST_ROOT.name, 'pdf'),
    'BACKUP_DIR': os.path.join(TEST_ROOT.name, 'backups'),
    'SETTINGS_FILE': os.path.join(TEST_ROOT.name, 'settings.json'),
})
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from werkzeug.security import generate_password_hash
import app as application
from werkzeug.datastructures import MultiDict
from models import (
    Account, AuditLog, Customer, DocumentTemplate, Expense, FieldOption,
    PackingBox, PackingItem, PackingList, Payment, PI, PIItem, Procurement,
    Product, Supplier, User, CustomerFollowUp, CustomsDocument, CustomsRevision, TranslationCache, db,
)
from document_export import _fixed_values, convert_excel_to_pdf, find_soffice
from customs_export import _model_summary


class SecuritySmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with application.app.app_context():
            db.session.add_all([
                User(account='alice-login', username='alice', password_hash=generate_password_hash('StrongPass123!'),
                     role='salesperson', salesperson_name='Alice', must_change_password=False),
                User(account='bob-login', username='bob', password_hash=generate_password_hash('StrongPass123!'),
                     role='salesperson', salesperson_name='Bob', must_change_password=False),
                User(account='admin-test-login', username='admin-test', password_hash=generate_password_hash('AdminPass123!'),
                     role='admin', salesperson_name='', must_change_password=False),
                Customer(name='Alice Customer', salesperson='Alice'),
                Customer(name='Bob Customer', salesperson='Bob'),
                Account(name='Approved', bank_name='SAFE BANK', account_no='123', swift_code='SAFE', brand='klista,qisuo'),
                Product(name='Original Product', product_code='ORIG', specification='Original spec', unit_price=10),
            ])
            db.session.commit()
            cls.alice_customer = Customer.query.filter_by(name='Alice Customer').one().id
            cls.bob_customer = Customer.query.filter_by(name='Bob Customer').one().id
            cls.approved_account = Account.query.filter_by(name='Approved').one().id
            product = Product.query.filter_by(name='Original Product').one()
            pi = PI(pi_number='PI-TEST-001', customer_id=cls.alice_customer,
                    salesperson='Alice', bank_info='SAFE BANK\nA/C: 123\nSWIFT: SAFE',
                    currency='USD', exchange_rate=7.0, total_amount=10)
            db.session.add(pi)
            db.session.flush()
            db.session.add(PIItem(pi_id=pi.id, product_id=product.id,
                                  quantity=1, unit_price=10, amount=10))
            db.session.commit()
            cls.alice_pi = pi.id

    def setUp(self):
        self.client = application.app.test_client()

    def token(self, path='/login'):
        html = self.client.get(path).get_data(as_text=True)
        return re.search(r'<meta name="csrf-token" content="([^"]+)"', html).group(1)

    def login(self, username='alice'):
        password = 'AdminPass123!' if username == 'admin-test' else 'StrongPass123!'
        account = {
            'alice': 'alice-login',
            'bob': 'bob-login',
            'admin-test': 'admin-test-login',
        }.get(username, username)
        return self.client.post('/login', data={
            'account': account,
            'password': password,
            'csrf_token': self.token(),
        })

    @staticmethod
    def image_upload(filename='receipt.png'):
        from PIL import Image
        receipt = BytesIO()
        Image.new('RGB', (4, 4), '#ffffff').save(receipt, format='PNG')
        receipt.seek(0)
        return receipt, filename

    def test_pi_draft_incomplete_save_resume_isolated_and_versioned(self):
        import json
        from models import PIDraft
        self.login('alice')
        token = self.token('/pi/create')
        with application.app.app_context():
            before = PI.query.count()
            product = Product.query.first()
            product_id = product.id
        row = {'id': product_id, 'productId': product_id, 'name': '', 'code': '', 'spec': 'draft spec', 'price': 9.5, 'priceUsd': 9.5, 'qty': 2, 'priceRaw': '', 'qtyRaw': '', 'imageMode': 'keep', 'imageSource': ''}
        result = self.client.post('/pi/drafts/save', data={
            'csrf_token': token, 'notes': '', 'currency': 'USD', 'draft_rows': json.dumps([row]),
            'item_image_file_' + str(product_id): self.image_upload(),
        })
        self.assertEqual(result.status_code, 200)
        saved = result.get_json()
        draft_id = saved['id']
        self.assertIn('draft spec', self.client.get('/pi/create?draft_id=' + str(draft_id)).get_data(as_text=True))
        with application.app.app_context():
            draft = db.session.get(PIDraft, draft_id)
            source = json.loads(draft.payload)['rows'][0]['imageSource']
            self.assertTrue(source)
            self.assertEqual(PI.query.count(), before)
        self.assertEqual(self.client.get('/uploads/' + source).status_code, 200)
        stale = self.client.post('/pi/drafts/save', data={'csrf_token': token, 'draft_id': draft_id, 'draft_version': 0, 'draft_rows': '[]'})
        self.assertEqual(stale.status_code, 409)
        self.login('bob')
        self.assertEqual(self.client.get('/pi/create?draft_id=' + str(draft_id)).status_code, 404)
        self.assertEqual(self.client.get('/uploads/' + source).status_code, 404)
        self.login('alice')
        token = self.token('/pi/create')
        row.update(name='Draft row', imageSource=source)
        updated = self.client.post('/pi/drafts/save', data={'csrf_token': token, 'draft_id': draft_id, 'draft_version': saved['version'], 'draft_rows': json.dumps([row]), 'notes': 'draft saved'}).get_json()
        with patch.object(application, '_generate_default_pi_documents', return_value=('draft.pdf', 'draft.xlsx')):
            response = self.client.post('/pi/create', data={
                'csrf_token': token, 'draft_id': draft_id, 'draft_version': updated['version'],
                'customer_id': self.alice_customer, 'notes': 'converted', 'salesperson': 'Alice',
                'account_id': self.approved_account, 'currency': 'USD', 'exchange_rate': '7',
                'selected_' + str(product_id): 'on', 'qty_' + str(product_id): '2',
                'unit_price_' + str(product_id): '9.5', 'item_image_source_' + str(product_id): source,
            })
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            draft = db.session.get(PIDraft, draft_id)
            self.assertTrue(draft.converted_pi_id)
            self.assertEqual(db.session.get(PI, draft.converted_pi_id).items[0].image_override, source)
            after = PI.query.count()
        self.client.post('/pi/create', data={'csrf_token': token, 'draft_id': draft_id, 'draft_version': updated['version']})
        with application.app.app_context():
            self.assertEqual(PI.query.count(), after)

    def test_login_account_resolves_username_and_authenticates(self):
        token = self.token()
        lookup = self.client.post('/login/account-name', json={
            'account': 'ALICE-LOGIN',
        }, headers={'X-CSRFToken': token})
        self.assertEqual(lookup.status_code, 200)
        self.assertEqual(lookup.get_json(), {'ok': True, 'username': 'alice'})

        missing = self.client.post('/login/account-name', json={
            'account': 'does-not-exist',
        }, headers={'X-CSRFToken': token})
        self.assertEqual(missing.status_code, 200)
        self.assertEqual(missing.get_json(), {'ok': False, 'username': ''})

        response = self.client.post('/login', data={
            'account': 'ALICE-LOGIN',
            'password': 'StrongPass123!',
            'csrf_token': token,
        })
        self.assertEqual(response.status_code, 302)

    def test_anonymous_business_download_redirects_to_login(self):
        self.assertEqual(self.client.get('/pi/1/download').status_code, 302)

    def test_salesperson_can_confirm_and_export_own_customs_package(self):
        with application.app.app_context():
            product = Product.query.filter_by(name='Original Product').one()
            pi = PI(
                pi_number='PI-CUSTOMS-TEST-001',
                customer_id=self.alice_customer,
                salesperson='Alice',
                currency='USD',
                total_amount=32,
                shipping_cost=12,
                customs_required=True,
            )
            db.session.add(pi)
            db.session.flush()
            pi_item = PIItem(
                pi_id=pi.id, product_id=product.id,
                quantity=2, unit_price=10, amount=20,
            )
            packing = PackingList(
                pi_id=pi.id, status='completed', created_by='alice',
            )
            db.session.add_all([pi_item, packing])
            db.session.flush()
            box = PackingBox(
                packing_list_id=packing.id, box_no='1',
                net_weight=1, gross_weight=2,
                length_cm=10, width_cm=20, height_cm=30,
            )
            db.session.add(box)
            db.session.flush()
            db.session.add(PackingItem(
                packing_box_id=box.id, pi_item_id=pi_item.id,
                quantity=2, product_name=product.name,
                product_code=product.product_code,
            ))
            db.session.commit()
            pi_id = pi.id

        values = {
            'trade_country': '美国',
            'destination_country': '美国',
            'export_date': '2026-09-14',
            'declaration_date': '2026-09-14',
            'supervision_mode': '一般贸易',
            'packing_type': '纸箱 / CARTON',
            'customs_freight_amount': '5',
            'customs_insurance_amount': '3',
            'customs_misc_amount': '1',
        }
        try:
            self.login('alice')
            path = f'/pi/{pi_id}/customs-documents'
            form_html = self.client.get(path).get_data(as_text=True)
            for field in (
                'customs_freight_amount', 'customs_insurance_amount',
                'customs_misc_amount',
            ):
                self.assertIn(f'name="{field}"', form_html)
            self.assertNotIn('exclude_charges_confirm', form_html)
            token = self.token(path)
            response = self.client.post(path, data={
                **values, 'csrf_token': token, 'version': '0',
                'action': 'save',
            })
            self.assertEqual(response.status_code, 302)
            with application.app.app_context():
                document = CustomsDocument.query.filter_by(pi_id=pi_id).one()
                version = document.version
                self.assertEqual(document.status, 'draft')

            token = self.token(path)
            response = self.client.post(path, data={
                **values, 'csrf_token': token, 'version': str(version),
                'action': 'confirm',
            })
            self.assertEqual(response.status_code, 302)
            with application.app.app_context():
                document = CustomsDocument.query.filter_by(pi_id=pi_id).one()
                self.assertEqual(document.status, 'confirmed')
                self.assertEqual(len(document.revisions), 1)
                self.assertIn('_confirmed_snapshot', document.data_json)

            response = self.client.get(f'/pi/{pi_id}/customs-documents.xlsx')
            self.assertEqual(response.status_code, 200)
            workbook = load_workbook(BytesIO(response.data), data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ['报关单', '发票', '装箱单', '申报要素', '合同'],
            )
            self.assertEqual(workbook['发票']['E17'].value, 20)
            self.assertEqual(workbook['合同']['E17'].value, 20)
            self.assertEqual(workbook['报关单']['J12'].value, 'USD 5.00')
            self.assertEqual(workbook['报关单']['L12'].value, 'USD 3.00')
            self.assertEqual(workbook['报关单']['N12'].value, 'USD 1.00')
            self.assertEqual(
                workbook['报关单']['A4'].value,
                '常州市克利斯达国际贸易有限公司',
            )
            self.assertTrue(workbook['报关单'].print_area)
            self.assertTrue(workbook['发票'].print_area)
            self.assertTrue(workbook['装箱单'].print_area)
            workbook.close()

            self.client.get('/logout')
            self.login('admin-test')
            token = self.token(path)
            with application.app.app_context():
                version = CustomsDocument.query.filter_by(pi_id=pi_id).one().version
            response = self.client.post(path, data={
                'csrf_token': token,
                'version': str(version),
                'action': 'unlock',
                'unlock_reason': '测试更正报关资料',
            })
            self.assertEqual(response.status_code, 302)
            with application.app.app_context():
                document = CustomsDocument.query.filter_by(pi_id=pi_id).one()
                self.assertEqual(document.status, 'draft')
                self.assertEqual(len(document.revisions), 2)
                self.assertNotIn('_confirmed_snapshot', document.data_json)
        finally:
            with application.app.app_context():
                pi = db.session.get(PI, pi_id)
                if pi:
                    db.session.delete(pi)
                    db.session.commit()

    def test_customs_merged_model_summary_limits_examples(self):
        group = {'product_models': ['A-1', 'B-2', 'C-3', 'D-4', 'A-1']}
        self.assertEqual(
            _model_summary(group),
            'A-1、B-2、C-3\n等，详见装箱单 / SEE PACKING LIST',
        )

    def test_customs_charge_allocation_cannot_exceed_pi_other_charges(self):
        errors = application._validate_customs_charge_allocation(
            SimpleNamespace(other_charges=10),
            {
                'customs_freight_amount': '6',
                'customs_insurance_amount': '3',
                'customs_misc_amount': '2',
            },
        )
        self.assertIn('合计不能超过 PI 其他费用', errors[0])

    def test_customs_draft_exports_before_packing_but_cannot_confirm(self):
        with application.app.app_context():
            product = Product.query.filter_by(name='Original Product').one()
            pi = PI(
                pi_number='PI-CUSTOMS-DRAFT-001',
                customer_id=self.alice_customer,
                salesperson='Alice', currency='USD', total_amount=10,
                customs_required=True,
            )
            db.session.add(pi)
            db.session.flush()
            db.session.add(PIItem(
                pi_id=pi.id, product_id=product.id,
                quantity=1, unit_price=10, amount=10,
            ))
            db.session.commit()
            pi_id = pi.id
        values = {
            'export_customs': '上海海关', 'transport_mode': '海运',
            'vehicle_voyage': 'TEST V002', 'bill_no': 'BL-002',
            'trade_country': '美国', 'destination_country': '美国',
            'destination_port': 'LOS ANGELES', 'departure_port': '上海港',
            'export_date': '2026-09-14', 'declaration_date': '2026-09-14',
            'supervision_mode': '一般贸易', 'packing_type': '纸箱 / CARTON',
        }
        try:
            self.login('alice')
            path = f'/pi/{pi_id}/customs-documents'
            response = self.client.post(path, data={
                **values, 'csrf_token': self.token(path),
                'version': '0', 'action': 'save',
            })
            self.assertEqual(response.status_code, 302)
            export = self.client.get(f'/pi/{pi_id}/customs-documents.xlsx')
            self.assertEqual(export.status_code, 200)
            self.assertIn('DRAFT', export.headers.get('Content-Disposition', ''))
            with application.app.app_context():
                document = CustomsDocument.query.filter_by(pi_id=pi_id).one()
                version = document.version
            response = self.client.post(path, data={
                **values, 'csrf_token': self.token(path),
                'version': str(version), 'action': 'confirm',
            }, follow_redirects=True)
            self.assertIn('装箱单尚未完成'.encode(), response.data)
            with application.app.app_context():
                document = CustomsDocument.query.filter_by(pi_id=pi_id).one()
                self.assertEqual(document.status, 'draft')
        finally:
            with application.app.app_context():
                pi = db.session.get(PI, pi_id)
                if pi:
                    db.session.delete(pi)
                    db.session.commit()

    def test_pi_list_combines_workflow_actions_and_shows_customer_country(self):
        with application.app.app_context():
            customer = db.session.get(Customer, self.alice_customer)
            original_country = customer.country
            customer.country = '测试国家'
            db.session.commit()
        try:
            self.login('alice')
            response = self.client.get('/pi/list')
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            self.assertIn('<th>国家</th>', html)
            self.assertIn('测试国家', html)
            self.assertIn('class="pi-list-workflow"', html)
            for label in ('回款', '采购', '装箱', '报关', '发货'):
                self.assertIn(f'pi-list-workflow-label">{label}</span>', html)
            self.assertNotIn('<th>回款状态</th>', html)
            self.assertNotIn('<th>订单进度</th>', html)
            self.assertNotIn('>业务处理</button>', html)
            self.assertIn(f'id="customsBtn-{self.alice_pi}"', html)
            self.assertIn(f'id="shippingBtn-{self.alice_pi}"', html)
        finally:
            with application.app.app_context():
                customer = db.session.get(Customer, self.alice_customer)
                customer.country = original_country
                db.session.commit()

    def test_customer_deal_counts_partial_payments_fees_and_caps_each_pi(self):
        with application.app.app_context():
            customer = Customer(name='Deal Formula Customer', salesperson='Alice')
            usd_pi = PI(
                pi_number='PI-DEAL-USD', customer=customer, salesperson='Alice',
                currency='USD', exchange_rate=7.0, total_amount=100,
            )
            rmb_pi = PI(
                pi_number='PI-DEAL-RMB', customer=customer, salesperson='Alice',
                currency='RMB', exchange_rate=7.0, total_amount=700,
            )
            db.session.add_all([customer, usd_pi, rmb_pi])
            db.session.flush()
            usd_first = Payment(pi_id=usd_pi.id, amount=40, fee=5, order_no='DEAL-USD-1')
            usd_second = Payment(pi_id=usd_pi.id, amount=60, fee=5, order_no='DEAL-USD-2')
            rmb_partial = Payment(pi_id=rmb_pi.id, amount=350, fee=70, order_no='DEAL-RMB-1')
            db.session.add_all([usd_first, usd_second, rmb_partial])
            db.session.flush()

            application._recalculate_pi_payments(usd_pi)
            application._recalculate_pi_payments(rmb_pi)
            application._recalculate_customer_deal(customer)
            self.assertTrue(usd_pi.paid)
            self.assertFalse(rmb_pi.paid)
            # USD is capped at $100; RMB 420 / 7 contributes another $60.
            self.assertEqual(customer.total_deal_usd, 160.0)

            usd_second.deleted_at = application.datetime.utcnow()
            application._recalculate_pi_payments(usd_pi)
            application._recalculate_customer_deal(customer)
            self.assertFalse(usd_pi.paid)
            self.assertEqual(customer.total_deal_usd, 105.0)

            usd_second.deleted_at = None
            application._recalculate_pi_payments(usd_pi)
            application._recalculate_customer_deal(customer)
            self.assertTrue(usd_pi.paid)
            self.assertEqual(customer.total_deal_usd, 160.0)
            db.session.rollback()

    def test_pi_three_decimal_prices_survive_save_copy_and_all_templates(self):
        from document_export import _item_values, render_excel_template
        self.login('alice')
        with application.app.app_context():
            product_id = Product.query.first().id
        data = {'csrf_token': self.token('/pi/create'), 'customer_id': self.alice_customer,
                'salesperson': 'Alice', 'notes': 'three decimal PI', 'account_id': self.approved_account,
                'currency': 'USD', 'exchange_rate': '7', 'product_order': '1,-1000000000',
                'selected_1': 'on', 'row_product_1': product_id, 'qty_1': '1', 'unit_price_1': '0.005',
                'selected_-1000000000': 'on', 'row_product_-1000000000': product_id,
                'qty_-1000000000': '3', 'unit_price_-1000000000': '1.235'}
        with patch.object(application, '_generate_default_pi_documents', return_value=('precision.pdf', 'precision.xlsx')):
            response = self.client.post('/pi/create', data=data)
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            pi = PI.query.filter_by(notes='three decimal PI').one()
            self.assertEqual([item.unit_price for item in pi.items], [0.005, 1.235])
            self.assertEqual([item.amount for item in pi.items], [0.01, 3.71])
            self.assertEqual(pi.total_amount, 3.72)
            export = application._pi_export_copy(pi)
            self.assertEqual(_item_values(export.items[1], 2)['item.unit_price'], 1.235)
            templates = [application.SYSTEM_DEFAULT_TEMPLATE_SOURCE, application.QISUO_LEGACY_TEMPLATE_SOURCE]
            for index, path in enumerate(templates):
                output = os.path.join(TEST_ROOT.name, 'precision-%s.xlsx' % index)
                render_excel_template(path, export, output)
                workbook = load_workbook(output)
                prices = [cell for sheet in workbook for row in sheet for cell in row if cell.number_format == '#,##0.000' and isinstance(cell.value, (int, float))]
                self.assertEqual([cell.value for cell in prices], [0.005, 1.235])
                workbook.close()
            # An uploaded template's old two-place format cannot truncate prices.
            workbook = Workbook()
            workbook.active['A1'] = '{{item.unit_price}}'
            workbook.active['A1'].number_format = '0.00'
            workbook.active['B1'] = 'USD {{item.unit_price}}'
            workbook.active['C1'] = '{{pi_number}}'
            path = os.path.join(TEST_ROOT.name, 'precision-custom.xlsx')
            workbook.save(path)
            render_excel_template(path, export, output)
            workbook = load_workbook(output)
            self.assertEqual(workbook.active['A2'].value, 1.235)
            self.assertEqual(workbook.active['A2'].number_format, '#,##0.000')
            self.assertEqual(workbook.active['B2'].value, 'USD 1.235')
            workbook.close()
        detail = self.client.get(response.location).get_data(as_text=True)
        self.assertIn('1.235', detail)

    def test_submitted_pi_items_preserve_explicit_selection_order(self):
        with application.app.app_context():
            first = Product(name='Order First', product_code='ORDER-FIRST', unit_price=11)
            second = Product(name='Order Second', product_code='ORDER-SECOND', unit_price=22)
            db.session.add_all([first, second])
            db.session.commit()
            first_id, second_id = first.id, second.id

            # Deliberately put selected_* keys in the opposite order.  The
            # explicit UI order must be authoritative.
            form = MultiDict([
                (f'selected_{first_id}', 'on'),
                (f'qty_{first_id}', '1'),
                (f'unit_price_{first_id}', '11'),
                (f'selected_{second_id}', 'on'),
                (f'qty_{second_id}', '1'),
                (f'unit_price_{second_id}', '22'),
                ('product_order', f'{second_id},{first_id}'),
            ])
            items, total = application._submitted_pi_items(form)
            self.assertEqual(
                [item['product'].id for item in items],
                [second_id, first_id],
            )
            self.assertEqual(total, 33)

            db.session.delete(first)
            db.session.delete(second)
            db.session.commit()

    def test_saved_pi_documents_use_export_workbench_renderer(self):
        calls = []

        def fake_render(export_pi, template, output_format, work_dir, preview=False):
            extension = 'pdf' if output_format == 'pdf' else 'xlsx'
            path = os.path.join(work_dir, f'generated.{extension}')
            with open(path, 'wb') as output:
                output.write(b'%PDF' if extension == 'pdf' else b'xlsx')
            calls.append((template.id, output_format, [item.product.name for item in export_pi.items]))
            return path, 'application/octet-stream'

        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            expected_template_id = application._export_template().id
            with patch.object(application, '_render_pi_export', side_effect=fake_render):
                pdf_name, excel_name = application._generate_default_pi_documents(pi)

        self.assertEqual([call[1] for call in calls], ['pdf', 'xlsx'])
        self.assertEqual([call[0] for call in calls], [expected_template_id, expected_template_id])
        self.assertEqual(calls[0][2], ['Original Product'])
        self.assertEqual(pdf_name, 'PI-TEST-001.pdf')
        self.assertEqual(excel_name, 'PI-TEST-001.xlsx')

    def test_pi_row_image_upload_copy_restore_and_export_are_per_pi(self):
        from PIL import Image
        image = BytesIO()
        Image.new('RGB', (20, 20), 'red').save(image, format='PNG')
        image.seek(0)
        self.login('admin-test')
        with application.app.app_context():
            product = Product.query.filter_by(product_code='ORIG').one()
            product_id = product.id
            catalog_image = product.image
        data = {
            'customer_id': str(self.alice_customer), 'salesperson': 'Alice',
            'notes': 'per-pi-image-test', 'issue_date': '2026-09-17',
            'currency': 'USD', 'exchange_rate': '7', 'company': 'klista',
            'shipping_cost': '0', 'account_id': str(self.approved_account),
            'product_order': str(product_id), f'selected_{product_id}': 'on',
            f'qty_{product_id}': '1', f'unit_price_{product_id}': '10',
            f'item_image_file_{product_id}': (image, 'clipboard.png'),
            'csrf_token': self.token('/pi/create'),
        }
        with patch.object(application, '_generate_default_pi_documents', return_value=('test.pdf', 'test.xlsx')):
            response = self.client.post('/pi/create', data=data, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            pi = PI.query.filter_by(notes='per-pi-image-test').one()
            pi_id = pi.id
            item = pi.items[0]
            item_id = item.id
            filename = item.image_override
            current_version = pi.version
            self.assertTrue(filename.endswith('.png'))
            self.assertEqual(item.display_image, filename)
            self.assertEqual(application._pi_export_copy(pi).items[0].product.image, filename)
            self.assertEqual(db.session.get(Product, product_id).image, catalog_image)
        response = self.client.get('/uploads/' + filename)
        self.assertEqual(response.status_code, 200)
        response.close()
        response = self.client.get('/uploads/thumb/' + filename)
        self.assertEqual(response.status_code, 200)
        response.close()
        self.client.get('/logout')
        self.login('bob')
        self.assertEqual(self.client.get('/uploads/' + filename).status_code, 404)
        self.assertEqual(self.client.get('/uploads/thumb/' + filename).status_code, 404)
        self.client.get('/logout')
        self.login('admin-test')
        url = f'/pi/{pi_id}/edit'
        row = -item_id
        copied_row = -1000000000
        edit = {
            'customer_id': str(self.alice_customer), 'salesperson': 'Alice',
            'notes': 'per-pi-image-test', 'issue_date': '2026-09-17',
            'currency': 'USD', 'exchange_rate': '7', 'company': 'klista',
            'shipping_cost': '0', 'product_order': f'{row},{copied_row}', 'version': str(current_version),
            f'row_item_{row}': str(item_id),
            'csrf_token': self.token(url),
        }
        for key in (row, copied_row):
            edit.update({f'selected_{key}': 'on', f'row_product_{key}': str(product_id),
                         f'qty_{key}': '1', f'unit_price_{key}': '10',
                         f'item_image_source_{key}': filename, f'item_image_mode_{key}': 'keep'})
        with patch.object(application, '_generate_default_pi_documents', return_value=('test.pdf', 'test.xlsx')):
            self.assertEqual(self.client.post(url, data=edit).status_code, 302)
        with application.app.app_context():
            pi = db.session.get(PI, pi_id)
            self.assertEqual([item.display_image for item in pi.items], [filename, filename])
            copied_item_id = pi.items[1].id
            export = application._pi_export_copy(pi)
            template = application._export_template()
            work_dir = tempfile.mkdtemp(dir=TEST_ROOT.name)
            output, _ = application._render_pi_export(export, template, 'xlsx', work_dir)
            workbook = load_workbook(output)
            self.assertEqual(len(workbook.active._images), 2)
            workbook.close()
        with application.app.app_context():
            edit['version'] = str(db.session.get(PI, pi_id).version)
        edit[f'row_item_{copied_row}'] = str(copied_item_id)
        edit[f'item_image_source_{row}'] = ''
        edit[f'item_image_mode_{row}'] = 'catalog'
        edit[f'item_image_source_{copied_row}'] = ''
        edit[f'item_image_mode_{copied_row}'] = 'clear'
        with patch.object(application, '_generate_default_pi_documents', return_value=('test.pdf', 'test.xlsx')):
            self.assertEqual(self.client.post(url, data=edit).status_code, 302)
        with application.app.app_context():
            pi = db.session.get(PI, pi_id)
            self.assertIsNone(pi.items[0].image_override)
            self.assertEqual(pi.items[1].image_override, '')
            self.assertEqual(db.session.get(Product, product_id).image, catalog_image)
        copied = self.client.post(f'/pi/{pi_id}/copy', data={'csrf_token': self.token(url)})
        self.assertEqual(copied.status_code, 302)
        with application.app.app_context():
            copy = PI.query.order_by(PI.id.desc()).first()
            self.assertEqual([item.image_override for item in copy.items], [None, ''])

    def test_pi_row_image_rejects_invalid_upload_and_foreign_sources(self):
        self.login('admin-test')
        with application.app.app_context():
            product_id = Product.query.filter_by(product_code='ORIG').one().id
            before = PI.query.count()
        data = {'customer_id': str(self.alice_customer), 'salesperson': 'Alice',
                'notes': 'invalid-image-test', 'currency': 'USD', 'company': 'klista',
                'exchange_rate': '7', 'shipping_cost': '0',
                f'selected_{product_id}': 'on', f'qty_{product_id}': '1',
                f'unit_price_{product_id}': '10', 'csrf_token': self.token('/pi/create'),
                f'item_image_file_{product_id}': (BytesIO(b'not an image'), 'invalid.png')}
        self.assertEqual(self.client.post('/pi/create', data=data, content_type='multipart/form-data').status_code, 400)
        data.pop(f'item_image_file_{product_id}')
        data[f'item_image_source_{product_id}'] = 'foreign-private-image.png'
        self.assertEqual(self.client.post('/pi/create', data=data).status_code, 400)
        with application.app.app_context():
            self.assertEqual(PI.query.count(), before)

    def test_new_pi_redirects_to_its_preview_after_creation(self):
        self.login('admin-test')
        with application.app.app_context():
            product_id = Product.query.filter_by(product_code='ORIG').one().id

        with patch.object(
            application,
            '_generate_default_pi_documents',
            return_value=('created-preview.pdf', 'created-preview.xlsx'),
        ):
            response = self.client.post('/pi/create', data={
                'customer_id': str(self.alice_customer),
                'salesperson': 'Alice',
                'payment_terms': '100% TT before shipment',
                'price_terms': '',
                'delivery_time': '',
                'bank_info': 'SAFE BANK',
                'notes': 'redirect-to-preview-test',
                'account_id': self.approved_account,
                'issue_date': '2026-09-09',
                'currency': 'USD',
                'exchange_rate': '7',
                'company': 'klista',
                'shipping_address': '',
                'shipping_cost': '0',
                'shipping_note': '',
                f'selected_{product_id}': 'on',
                f'qty_{product_id}': '1',
                f'unit_price_{product_id}': '10',
                'product_order': str(product_id),
                'csrf_token': self.token('/pi/create'),
            })

        self.assertEqual(response.status_code, 302)
        location = response.headers['Location']
        self.assertRegex(location, r'/pi/\d+$')
        self.assertNotEqual(location, '/pi/list')
        created_id = int(location.rsplit('/', 1)[-1])
        preview = self.client.get(location)
        self.assertEqual(preview.status_code, 200)
        self.assertIn('redirect-to-preview-test', preview.get_data(as_text=True))

        with application.app.app_context():
            created = db.session.get(PI, created_id)
            if created:
                db.session.delete(created)
                db.session.commit()

    def test_new_pi_requires_notes_in_browser_and_server(self):
        self.login('admin-test')
        form_html = self.client.get('/pi/create').get_data(as_text=True)
        self.assertRegex(
            form_html,
            r'<label for="notes"[^>]*>PI 备注\s*<span[^>]*>\*</span></label>',
        )
        self.assertRegex(
            form_html,
            r'<textarea[^>]*id="notes"[^>]*name="notes"[^>]*required',
        )

        with application.app.app_context():
            product_id = Product.query.filter_by(product_code='ORIG').one().id
            before_count = PI.query.count()

        response = self.client.post('/pi/create', data={
            'customer_id': str(self.alice_customer),
            'salesperson': 'Alice',
            'payment_terms': '100% TT before shipment',
            'price_terms': '',
            'delivery_time': '',
            'bank_info': 'SAFE BANK',
            'notes': '   ',
            'issue_date': '2026-09-14',
            'currency': 'USD',
            'exchange_rate': '7',
            'company': 'klista',
            'shipping_address': '',
            'shipping_cost': '0',
            'shipping_note': '',
            f'selected_{product_id}': 'on',
            f'qty_{product_id}': '1',
            f'unit_price_{product_id}': '10',
            'product_order': str(product_id),
            'csrf_token': self.token('/pi/create'),
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn('请填写 PI 备注。', response.get_data(as_text=True))
        with application.app.app_context():
            self.assertEqual(PI.query.count(), before_count)

    def test_pi_edit_protects_downstream_records_and_preserves_item_ids(self):
        with application.app.app_context():
            original_product = Product.query.filter_by(product_code='ORIG').one()
            extra_product = Product(
                name='PI Edit Extra Product', product_code='PI-EDIT-EXTRA',
                unit_price=5,
            )
            supplier = Supplier(name='PI Edit Protected Supplier')
            pi = PI(
                pi_number='PI-EDIT-PROTECTED',
                customer_id=self.alice_customer,
                salesperson='Alice', currency='USD', exchange_rate=7,
                issue_date=date(2026, 9, 10), total_amount=25,
            )
            db.session.add_all([extra_product, supplier, pi])
            db.session.flush()
            protected_item = PIItem(
                pi_id=pi.id, product_id=original_product.id,
                quantity=2, unit_price=10, amount=20,
            )
            extra_item = PIItem(
                pi_id=pi.id, product_id=extra_product.id,
                quantity=1, unit_price=5, amount=5,
            )
            db.session.add_all([protected_item, extra_item])
            db.session.flush()
            db.session.add(Procurement(
                pi_id=pi.id, pi_item_id=protected_item.id,
                supplier_id=supplier.id, unit_price=8,
                quantity=2, total=16,
            ))
            packing_list = PackingList(
                pi_id=pi.id, status='draft', created_by='admin-test',
                updated_by='admin-test',
            )
            db.session.add(packing_list)
            db.session.flush()
            box = PackingBox(
                packing_list_id=packing_list.id, box_no='1', sort_order=0,
            )
            db.session.add(box)
            db.session.flush()
            db.session.add(PackingItem(
                packing_box_id=box.id, pi_item_id=protected_item.id,
                quantity=1, product_name=original_product.name,
            ))
            db.session.add(Payment(
                pi_id=pi.id, amount=25, fee=0,
                order_no='PI-EDIT-PROTECTED-PAYMENT',
            ))
            pi.received_amount = 25
            pi.paid = True
            db.session.commit()
            application._recalculate_customer_deal(
                db.session.get(Customer, self.alice_customer)
            )
            db.session.commit()
            pi_id = pi.id
            original_product_id = original_product.id
            protected_item_id = protected_item.id
            extra_item_id = extra_item.id
            extra_product_id = extra_product.id
            supplier_id = supplier.id

        def edit_data(client, *, protected_qty=2, include_protected=True,
                      unlock=False, reason='', note=''):
            with application.app.app_context():
                current = db.session.get(PI, pi_id)
                version = current.version
            data = {
                'version': str(version),
                'customer_id': str(self.alice_customer),
                'salesperson': 'Alice',
                'payment_terms': '100% TT before shipment',
                'price_terms': '',
                'delivery_time': '',
                'bank_info': '',
                'notes': note,
                'issue_date': '2026-09-10',
                'currency': 'USD',
                'exchange_rate': '7',
                'company': 'klista',
                'shipping_address': 'Safe editable address',
                'shipping_cost': '0',
                'shipping_note': '',
                f'selected_{extra_product_id}': 'on',
                f'qty_{extra_product_id}': '1',
                f'unit_price_{extra_product_id}': '5',
                'product_order': (
                    f'{original_product_id},{extra_product_id}'
                    if include_protected else str(extra_product_id)
                ),
                'csrf_token': self.token(f'/pi/{pi_id}/edit'),
            }
            if include_protected:
                data.update({
                    f'selected_{original_product_id}': 'on',
                    f'qty_{original_product_id}': str(protected_qty),
                    f'unit_price_{original_product_id}': '10',
                })
            if unlock:
                data['downstream_unlock'] = '1'
                data['downstream_change_reason'] = reason
            return data

        try:
            self.login('alice')
            page = self.client.get(f'/pi/{pi_id}/edit')
            html = page.get_data(as_text=True)
            self.assertIn('该 PI 已有下游业务记录，结构信息已锁定', html)
            self.assertIn('回款 1 笔', html)
            self.assertIn('采购 1 行 / 1 家供应商', html)
            self.assertIn('装箱单草稿 / 1 箱 / 1 行', html)
            self.assertNotIn('id="downstream_unlock"', html)

            denied = self.client.post(
                f'/pi/{pi_id}/edit',
                data=edit_data(self.client, protected_qty=3, note='denied'),
            )
            self.assertEqual(denied.status_code, 302)
            with application.app.app_context():
                self.assertEqual(db.session.get(PIItem, protected_item_id).quantity, 2)

            with patch.object(
                application, '_generate_default_pi_documents',
                return_value=('protected.pdf', 'protected.xlsx'),
            ):
                notes_only = self.client.post(
                    f'/pi/{pi_id}/edit',
                    data=edit_data(self.client, note='allowed notes-only edit'),
                )
            self.assertEqual(notes_only.status_code, 302)
            with application.app.app_context():
                current = db.session.get(PI, pi_id)
                self.assertEqual(current.notes, 'allowed notes-only edit')
                self.assertEqual(db.session.get(PIItem, protected_item_id).quantity, 2)
                self.assertEqual(db.session.get(PIItem, extra_item_id).quantity, 1)

            self.client.get('/logout')
            self.login('admin-test')
            admin_page = self.client.get(f'/pi/{pi_id}/edit').get_data(as_text=True)
            self.assertIn('id="downstream_unlock"', admin_page)
            self.assertIn('name="downstream_unlock" id="downstream_unlock" form="pi-form"', admin_page)
            self.assertIn('name="downstream_change_reason" id="downstream_change_reason" form="pi-form"', admin_page)
            self.assertIn('原因会写入审计日志', admin_page)

            with patch.object(
                application, '_generate_default_pi_documents',
                return_value=('protected.pdf', 'protected.xlsx'),
            ):
                updated = self.client.post(
                    f'/pi/{pi_id}/edit',
                    data=edit_data(
                        self.client, protected_qty=3, unlock=True,
                        reason='客户确认增加采购数量', note='admin updated',
                    ),
                )
            self.assertEqual(updated.status_code, 302)
            with application.app.app_context():
                item = db.session.get(PIItem, protected_item_id)
                self.assertEqual(item.quantity, 3)
                self.assertEqual(item.pi_id, pi_id)
                self.assertIsNotNone(db.session.get(PIItem, extra_item_id))
                current = db.session.get(PI, pi_id)
                self.assertEqual(current.received_amount, 25)
                self.assertFalse(current.paid)
                audit = AuditLog.query.filter_by(
                    entity_type='pi', entity_id=pi_id, action='update',
                ).order_by(AuditLog.id.desc()).first()
                self.assertIn('下游解锁原因', audit.summary)
                self.assertIn('客户确认增加采购数量', audit.after_json)

            blocked_remove = self.client.post(
                f'/pi/{pi_id}/edit',
                data=edit_data(
                    self.client, include_protected=False, unlock=True,
                    reason='尝试移除已有下游记录产品',
                ),
            )
            self.assertEqual(blocked_remove.status_code, 302)
            with application.app.app_context():
                self.assertIsNotNone(db.session.get(PIItem, protected_item_id))
                self.assertEqual(Procurement.query.filter_by(
                    pi_item_id=protected_item_id
                ).count(), 1)
                self.assertEqual(PackingItem.query.filter_by(
                    pi_item_id=protected_item_id
                ).count(), 1)
        finally:
            with application.app.app_context():
                pi = db.session.get(PI, pi_id)
                if pi:
                    db.session.delete(pi)
                    db.session.flush()
                supplier = db.session.get(Supplier, supplier_id)
                if supplier:
                    db.session.delete(supplier)
                product = db.session.get(Product, extra_product_id)
                if product:
                    db.session.delete(product)
                customer = db.session.get(Customer, self.alice_customer)
                application._recalculate_customer_deal(customer)
                db.session.commit()

    def test_direct_pi_pdf_download_uses_current_default_export_renderer(self):
        self.login('alice')
        calls = []

        def fake_render(export_pi, template, output_format, work_dir, preview=False):
            path = os.path.join(work_dir, 'current-template.pdf')
            with open(path, 'wb') as output:
                output.write(b'%PDF-current-default-template')
            calls.append((
                export_pi.pi_number,
                template.code,
                output_format,
                preview,
            ))
            return path, 'application/pdf'

        with patch.object(application, '_render_pi_export', side_effect=fake_render):
            response = self.client.get(f'/pi/{self.alice_pi}/download')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.data, b'%PDF-current-default-template')
            self.assertEqual(response.mimetype, 'application/pdf')
            self.assertIn('no-store', response.headers.get('Cache-Control', ''))
            response.close()

        self.assertEqual(calls, [
            ('PI-TEST-001', 'system-default', 'pdf', False),
        ])

    def test_salesperson_cannot_read_another_customer(self):
        self.login()
        self.assertEqual(self.client.get(f'/customers/{self.alice_customer}').status_code, 200)
        self.assertEqual(self.client.get(f'/customers/{self.bob_customer}').status_code, 403)

    def test_customer_detail_summarizes_purchased_products_without_cross_customer_data(self):
        self.login('alice')
        with application.app.app_context():
            product = Product.query.filter_by(name='Original Product').one()
            extra_pi = PI(
                pi_number='PI-CUSTOMER-SUMMARY', customer_id=self.alice_customer,
                salesperson='Alice', currency='RMB', exchange_rate=7.0,
                total_amount=37.5, issue_date=date(2030, 1, 2),
            )
            db.session.add(extra_pi)
            db.session.flush()
            db.session.add(PIItem(
                pi_id=extra_pi.id, product_id=product.id, quantity=3,
                unit_price=12.5, amount=37.5, name_override='Latest Product Name',
                code_override='LATEST-CODE', spec_override='Latest spec',
            ))
            db.session.commit()
            extra_pi_id = extra_pi.id

        html = self.client.get(f'/customers/{self.alice_customer}').get_data(as_text=True)
        self.assertIn('成交产品汇总（1）', html)
        self.assertIn('Latest Product Name', html)
        self.assertIn('LATEST-CODE', html)
        self.assertIn('Latest spec', html)
        self.assertIn('2030-01-02', html)
        self.assertIn('PI-CUSTOMER-SUMMARY', html)
        self.assertIn('¥12.500', html)
        self.assertRegex(html, r'<td class="text-end fw-bold">\s*4\s*</td>')
        self.assertNotIn('Bob Customer', html)

        with application.app.app_context():
            PIItem.query.filter_by(pi_id=extra_pi_id).delete()
            db.session.delete(db.session.get(PI, extra_pi_id))
            db.session.commit()

    def test_customer_follow_up_defaults_updates_dashboard_and_is_scoped(self):
        self.login('alice')
        detail_url = f'/customers/{self.alice_customer}'
        detail = self.client.get(detail_url)
        self.assertEqual(detail.status_code, 200)
        detail_html = detail.get_data(as_text=True)
        self.assertIn('id="follow-ups"', detail_html)
        self.assertIn('新增跟进', detail_html)
        token = self.token(detail_url)
        response = self.client.post(
            f'/customers/{self.alice_customer}/follow-ups',
            data={
                'csrf_token': token,
                'content': '已发送新报价，等待客户确认。',
                'status': 'waiting_reply',
                'contacted_at': date.today().isoformat(),
                'next_follow_up_date': '',
            },
        )
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            customer = db.session.get(Customer, self.alice_customer)
            follow_up = CustomerFollowUp.query.filter_by(customer_id=customer.id).order_by(CustomerFollowUp.id.desc()).first()
            self.assertIsNotNone(follow_up)
            self.assertEqual(follow_up.content, '已发送新报价，等待客户确认。')
            expected_days = application._follow_up_default_days(customer, 'waiting_reply')
            self.assertEqual(customer.next_follow_up_date, date.today() + timedelta(days=expected_days))
            follow_up_id = follow_up.id
        listing = self.client.get('/customers?follow=waiting_reply').get_data(as_text=True)
        self.assertIn('Alice Customer', listing)
        dashboard = self.client.get('/').get_data(as_text=True)
        self.assertIn('今天需要跟进', dashboard)
        self.client.get('/logout')
        self.login('bob')
        forbidden = self.client.post(
            f'/customers/{self.alice_customer}/follow-ups',
            data={'csrf_token': self.token('/customers'), 'content': '越权', 'status': 'needs_followup'},
        )
        self.assertEqual(forbidden.status_code, 403)
        with application.app.app_context():
            follow_up = db.session.get(CustomerFollowUp, follow_up_id)
            if follow_up:
                db.session.delete(follow_up)
            customer = db.session.get(Customer, self.alice_customer)
            customer.follow_up_status = 'needs_followup'
            customer.next_follow_up_date = None
            customer.last_follow_up_at = None
            db.session.commit()

    def test_salesperson_cannot_open_admin_pages(self):
        self.login()
        self.assertEqual(self.client.get('/users').status_code, 403)

    def test_customs_record_is_scoped_audited_and_does_not_change_order_state(self):
        self.login('alice')
        list_page = self.client.get('/pi/list')
        self.assertEqual(list_page.status_code, 200)
        list_html = list_page.get_data(as_text=True)
        self.assertIn('报关', list_html)
        self.assertIn('customsBtn-', list_html)

        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            original_state = (
                pi.received_amount, pi.procurement_confirmed,
                pi.shipping_completed, pi.procurement_status,
            )

        token = self.token('/pi/list')
        invalid = self.client.post(
            f'/api/pi/{self.alice_pi}/customs-declaration',
            json={'customs_required': None, 'note': ''},
            headers={'X-CSRFToken': token},
        )
        self.assertEqual(invalid.status_code, 400)

        saved = self.client.post(
            f'/api/pi/{self.alice_pi}/customs-declaration',
            json={'customs_required': True, 'note': '一般贸易报关'},
            headers={'X-CSRFToken': token},
        )
        self.assertEqual(saved.status_code, 200)
        self.assertTrue(saved.get_json()['customs_required'])

        detail = self.client.get(
            f'/api/pi/{self.alice_pi}/customs-declaration'
        )
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.get_json()['note'], '一般贸易报关')
        self.assertEqual(detail.get_json()['recorded_by'], 'alice')

        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            self.assertIs(pi.customs_required, True)
            self.assertEqual(pi.customs_note, '一般贸易报关')
            self.assertEqual(
                (pi.received_amount, pi.procurement_confirmed,
                 pi.shipping_completed, pi.procurement_status),
                original_state,
            )
            self.assertIsNotNone(AuditLog.query.filter_by(
                entity_type='pi', entity_id=pi.id, action='update'
            ).filter(AuditLog.summary.contains('报关记录')).first())

        self.client.get('/logout')
        self.login('bob')
        self.assertEqual(self.client.get(
            f'/api/pi/{self.alice_pi}/customs-declaration'
        ).status_code, 403)
        self.assertEqual(self.client.post(
            f'/api/pi/{self.alice_pi}/customs-declaration',
            json={'customs_required': False},
            headers={'X-CSRFToken': self.token('/pi/list')},
        ).status_code, 403)

        self.client.get('/logout')
        self.login('admin-test')
        changed = self.client.post(
            f'/api/pi/{self.alice_pi}/customs-declaration',
            json={'customs_required': False, 'note': '管理员复核'},
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(changed.status_code, 200)
        self.assertIs(changed.get_json()['customs_required'], False)

        # Keep this shared fixture neutral for tests that run afterward.
        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            pi.customs_required = None
            pi.customs_note = ''
            pi.customs_recorded_at = None
            pi.customs_recorded_by = ''
            db.session.commit()

    def test_product_translation_is_cached_and_normalized(self):
        self.login('alice')
        token = self.token('/products/add')
        source_key = application._translation_source_key('Unique Connector 987')
        with application.app.app_context():
            TranslationCache.query.filter_by(source_key=source_key).delete()
            db.session.commit()
        try:
            with patch.object(
                application,
                '_translate_product_name',
                return_value=('唯一连接器', 'test-provider'),
            ) as upstream:
                first = self.client.post(
                    '/api/translate', data={'text': 'Unique Connector 987'},
                    headers={'X-CSRFToken': token},
                )
                second = self.client.post(
                    '/api/translate', data={'text': '  UNIQUE   connector 987  '},
                    headers={'X-CSRFToken': token},
                )
            self.assertEqual(first.status_code, 200)
            self.assertFalse(first.get_json()['cached'])
            self.assertEqual(second.status_code, 200)
            self.assertTrue(second.get_json()['cached'])
            self.assertEqual(second.get_json()['translated'], '唯一连接器')
            upstream.assert_called_once_with('Unique Connector 987')
            with application.app.app_context():
                cache = TranslationCache.query.filter_by(source_key=source_key).one()
                self.assertEqual(cache.provider, 'test-provider')
        finally:
            with application.app.app_context():
                TranslationCache.query.filter_by(source_key=source_key).delete()
                db.session.commit()

    def test_product_translation_failure_is_friendly_and_retryable(self):
        self.login('alice')
        token = self.token('/products/add')
        with patch.object(
            application,
            '_translate_product_name',
            side_effect=application.TranslationUnavailable(),
        ):
            response = self.client.post(
                '/api/translate', data={'text': 'Never Cached Translation 654'},
                headers={'X-CSRFToken': token},
            )
        self.assertEqual(response.status_code, 503)
        payload = response.get_json()
        self.assertTrue(payload['retryable'])
        self.assertIn('稍后重试', payload['error'])
        self.assertNotIn('HTTP Error', payload['error'])
        too_long = self.client.post(
            '/api/translate', data={'text': 'a' * 301},
            headers={'X-CSRFToken': token},
        )
        self.assertEqual(too_long.status_code, 400)

    def test_every_new_product_entry_requires_a_chinese_name(self):
        self.login('admin-test')
        add_page = self.client.get('/products/add').get_data(as_text=True)
        self.assertRegex(add_page, r'id="chinese_name"[^>]*required')

        rejected = self.client.post('/products/add', data={
            'name': 'Regular Missing Chinese',
            'product_code': 'REG-MISSING-CN',
            'unit_price': '1',
            'csrf_token': self.token('/products/add'),
        })
        self.assertEqual(rejected.status_code, 200)
        self.assertIn('中文名称不能为空', rejected.get_data(as_text=True))
        with application.app.app_context():
            self.assertIsNone(Product.query.filter_by(product_code='REG-MISSING-CN').first())

        def workbook_upload(chinese_name):
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(['序号', '产品英文名称', '中文名称', '编码', '规格', '单价', '图片'])
            sheet.append([1, 'Imported Required Chinese', chinese_name, 'IMPORT-CN-REQ', 'S', 2.5, ''])
            output = BytesIO()
            workbook.save(output)
            workbook.close()
            output.seek(0)
            return output, 'products.xlsx'

        missing_import = self.client.post('/products/import', data={
            'csrf_token': self.token('/products/import'),
            'excel_file': workbook_upload(''),
        }, content_type='multipart/form-data', follow_redirects=True)
        self.assertIn('中文名称不能为空', missing_import.get_data(as_text=True))
        with application.app.app_context():
            self.assertIsNone(Product.query.filter_by(product_code='IMPORT-CN-REQ').first())

        imported = self.client.post('/products/import', data={
            'csrf_token': self.token('/products/import'),
            'excel_file': workbook_upload('导入中文名称'),
        }, content_type='multipart/form-data', follow_redirects=True)
        self.assertIn('已导入 1 个产品', imported.get_data(as_text=True))
        with application.app.app_context():
            product = Product.query.filter_by(product_code='IMPORT-CN-REQ').one()
            self.assertEqual(product.chinese_name, '导入中文名称')
            db.session.delete(product)
            db.session.commit()

    def test_salesperson_can_manage_individual_products_but_not_admin_bulk_tools(self):
        self.login()
        product_page = self.client.get('/products')
        self.assertEqual(product_page.status_code, 200)
        product_html = product_page.get_data(as_text=True)
        self.assertIn('添加产品', product_html)
        self.assertIn('title="复制"', product_html)
        self.assertIn('title="完整编辑（含图片）"', product_html)
        self.assertIn('title="删除或停用"', product_html)
        self.assertNotIn('仅查看', product_html)
        self.assertEqual(self.client.get('/products/add').status_code, 200)

        response = self.client.post('/products/add', data={
            'name': 'Sales Added Product',
            'chinese_name': '业务员新增产品',
            'product_code': 'SALES-ADD',
            'unit_price': '12.5',
            'unit_price_rmb': '87.5',
            'csrf_token': self.token('/products/add'),
        })
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            product = Product.query.filter_by(product_code='SALES-ADD').one()
            audit = AuditLog.query.filter_by(
                entity_type='product', entity_id=product.id, action='create'
            ).one()
            self.assertEqual(audit.username, 'alice')
            product_id = product.id

        create_pi_html = self.client.get('/pi/create').get_data(as_text=True)
        self.assertIn('data-bs-target="#quickAddProductModal"', create_pi_html)
        self.assertIn('id="qa_chinese_name"', create_pi_html)
        self.assertIn('id="qa_translate"', create_pi_html)
        self.assertIn('id="qa_chinese_name" required', create_pi_html)
        self.assertIn("fd.append('chinese_name'", create_pi_html)
        self.assertIn("fetch('/api/translate'", create_pi_html)
        missing_chinese = self.client.post('/api/products/add', data={
            'name': 'Missing Chinese Product',
            'product_code': 'MISSING-CN',
            'csrf_token': self.token('/pi/create'),
        })
        self.assertEqual(missing_chinese.status_code, 400)
        self.assertIn('中文名称不能为空', missing_chinese.get_json()['error'])
        api_response = self.client.post('/api/products/add', data={
            'name': 'Sales Quick Product',
            'chinese_name': '业务员快速产品',
            'product_code': 'SALES-QUICK',
            'unit_price': '5',
            'csrf_token': self.token('/pi/create'),
        })
        self.assertEqual(api_response.status_code, 200)
        self.assertTrue(api_response.get_json()['success'])
        self.assertEqual(api_response.get_json()['product']['chinese_name'], '业务员快速产品')

        self.assertEqual(self.client.get(f'/products/{product_id}/edit').status_code, 200)
        edit_pi_html = self.client.get(f'/pi/{self.alice_pi}/edit').get_data(as_text=True)
        self.assertIn('id="qa_prod_chinese_name"', edit_pi_html)
        self.assertIn('id="qa_prod_translate"', edit_pi_html)
        self.assertIn("fd.append('chinese_name'", edit_pi_html)
        edited = self.client.post(f'/products/{product_id}/edit', data={
            'name': 'Sales Edited Product',
            'product_code': 'SALES-EDITED',
            'unit_price': '13.5',
            'unit_price_rmb': '94.5',
            'csrf_token': self.token(f'/products/{product_id}/edit'),
        })
        self.assertEqual(edited.status_code, 302)
        inline_updated = self.client.post(
            f'/api/products/{product_id}/update',
            json={'specification': 'Sales inline edit'},
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(inline_updated.status_code, 200)

        copied = self.client.post(
            f'/api/products/{product_id}/copy',
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(copied.status_code, 200)
        copied_id = copied.get_json()['product']['id']
        deleted_copy = self.client.post(
            f'/api/products/{copied_id}/delete',
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(deleted_copy.status_code, 200)
        self.assertEqual(deleted_copy.get_json()['action'], 'deleted')

        self.assertEqual(self.client.get('/products/import').status_code, 403)
        denied_batch = self.client.post('/products/batch-delete', data={
            'ids': str(product_id),
            'csrf_token': self.token('/products'),
        })
        self.assertEqual(denied_batch.status_code, 403)
        with application.app.app_context():
            product = db.session.get(Product, product_id)
            self.assertEqual(product.name, 'Sales Edited Product')
            self.assertEqual(product.specification, 'Sales inline edit')
            self.assertIsNotNone(AuditLog.query.filter_by(
                entity_type='product', entity_id=product_id,
                action='update', username='alice',
            ).first())
            db.session.add(PIItem(
                pi_id=self.alice_pi,
                product_id=product_id,
                quantity=1,
                unit_price=13.5,
                amount=13.5,
            ))
            db.session.commit()

        retired = self.client.post(
            f'/api/products/{product_id}/delete',
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(retired.status_code, 200)
        self.assertEqual(retired.get_json()['action'], 'disabled')
        with application.app.app_context():
            product = db.session.get(Product, product_id)
            self.assertFalse(product.active)
            PIItem.query.filter_by(
                pi_id=self.alice_pi, product_id=product_id
            ).delete()
            db.session.delete(product)
            db.session.commit()

    def test_pi_keyword_search_and_order_descriptions(self):
        self.login('admin-test')
        with application.app.app_context():
            product = Product(name='36KD Gas Nozzle', product_code='FUZZY-NOZZLE',
                              specification='Copper', unit_price=5)
            exact = Product(name='36KD nozzle', product_code='FUZZY-EXACT', unit_price=5)
            unrelated = Product(name='36KD Torch', product_code='FUZZY-TORCH', unit_price=5)
            db.session.add_all([product, exact, unrelated])
            db.session.commit()
            product_id, exact_id = product.id, exact.id
        for url in ('/api/products/search?q=36kd%20NOZZLE',
                    '/api/products/search?picker=1&q=36kd%20NOZZLE'):
            payload = self.client.get(url).get_json()
            results = payload['items'] if isinstance(payload, dict) else payload
            self.assertEqual(results[0]['id'], exact_id)
            self.assertIn(product_id, [row['id'] for row in results])
            self.assertNotIn('FUZZY-TORCH', [row['code'] for row in results])
        cross_field = self.client.get('/api/products/search?q=36KD%20Copper').get_json()
        self.assertIn(product_id, [row['id'] for row in cross_field])
        data = {
            'customer_id': str(self.alice_customer), 'salesperson': 'Alice',
            'currency': 'USD', 'exchange_rate': '7', 'company': 'klista',
            'issue_date': '2026-09-17', 'notes': 'Description override test',
            'account_id': self.approved_account,
            f'selected_{product_id}': 'on', f'qty_{product_id}': '2',
            f'unit_price_{product_id}': '5', 'product_order': str(product_id),
            f'item_name_{product_id}': 'Customer nozzle', f'item_spec_{product_id}': '',
            f'item_code_{product_id}': 'CLIENT-001',
            'csrf_token': self.token('/pi/create'),
        }
        with patch.object(application, '_generate_default_pi_documents',
                          return_value=('test.pdf', 'test.xlsx')):
            response = self.client.post('/pi/create', data=data)
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            pi = PI.query.filter_by(notes='Description override test').one()
            item = pi.items[0]
            pi_id, item_id = pi.id, item.id
            self.assertEqual(item.display_code, 'CLIENT-001')
            self.assertEqual(item.product.product_code, 'FUZZY-NOZZLE')
            self.assertEqual(item.to_dict()['product_code'], 'CLIENT-001')
            self.assertEqual(item.display_name, 'Customer nozzle')
            self.assertEqual(item.display_specification, '')
            self.assertEqual(item.product.name, '36KD Gas Nozzle')
            self.assertEqual(item.product.specification, 'Copper')
            from document_export import _item_values, render_excel_template
            self.assertEqual(_item_values(item, 1)['item.name'], 'Customer nozzle')
            self.assertEqual(_item_values(item, 1)['item.specification'], '')
            export_copy = application._pi_export_copy(pi)
            self.assertEqual(export_copy.items[0].product.name, 'Customer nozzle')
            template = Workbook()
            template.active.append(['{{pi_number}}'])
            template.active.append(['{{item.name}}', '{{item.specification}}', '{{item.code}}'])
            template_path = os.path.join(TEST_ROOT.name, 'description-template.xlsx')
            output_path = os.path.join(TEST_ROOT.name, 'description-output.xlsx')
            template.save(template_path)
            render_excel_template(template_path, pi, output_path)
            exported = load_workbook(output_path).active
            self.assertEqual(exported['A2'].value, 'Customer nozzle')
            self.assertIsNone(exported['B2'].value)
            self.assertEqual(exported['C2'].value, 'CLIENT-001')
            self.assertEqual(export_copy.items[0].product.product_code, 'CLIENT-001')
            preload = application._preload_products(pi)[0]
            self.assertEqual(preload['originalName'], '36KD Gas Nozzle')
            self.assertEqual(preload['name'], 'Customer nozzle')
            version = pi.version
        page = self.client.get(f'/pi/{pi_id}/edit').get_data(as_text=True)
        self.assertIn('sel-name pi-structural-control', page)
        data.update({'version': str(version), 'csrf_token': self.token(f'/pi/{pi_id}/edit'),
                     f'item_name_{product_id}': 'Edited nozzle', f'item_spec_{product_id}': 'Custom spec',
                     f'item_code_{product_id}': ''})
        with patch.object(application, '_generate_default_pi_documents',
                          return_value=('test.pdf', 'test.xlsx')):
            response = self.client.post(f'/pi/{pi_id}/edit', data=data)
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            pi = db.session.get(PI, pi_id)
            self.assertEqual(pi.items[0].id, item_id)
            self.assertEqual(pi.items[0].display_code, '')
            self.assertEqual(pi.items[0].product.product_code, 'FUZZY-NOZZLE')
            self.assertEqual(pi.items[0].display_name, 'Edited nozzle')
            self.assertEqual(pi.items[0].display_specification, 'Custom spec')
            self.assertEqual(pi.items[0].product.name, '36KD Gas Nozzle')
            submitted, _ = application._submitted_pi_items(MultiDict(dict(data, **{
                f'item_name_{product_id}': 'Changed after purchase', f'item_code_{product_id}': 'CHANGED'})))
            changes = application._pi_structural_changes(
                pi, customer_id=pi.customer_id, salesperson=pi.salesperson,
                currency=pi.currency, exchange_rate=pi.exchange_rate, company=pi.company,
                issue_date=pi.issue_date, shipping_cost=pi.shipping_cost, selected_items=submitted)
            self.assertTrue(any('品名' in change for change in changes))
            self.assertTrue(any('编码' in change for change in changes))
        with application.app.app_context():
            source = db.session.get(PI, pi_id)
            source.issue_date = date(2020, 1, 1)
            db.session.commit()
        list_page = self.client.get('/pi/list').get_data(as_text=True)
        self.assertIn('复制 PI', list_page)
        self.assertIn(f'action="/pi/{pi_id}/copy"', list_page)
        token = self.token('/pi/list')
        with patch.object(application, '_generate_default_pi_documents',
                          return_value=('test.pdf', 'test.xlsx')):
            self.client.post(f'/pi/{pi_id}/copy', data={'csrf_token': token})
        with application.app.app_context():
            copied = PI.query.filter(PI.notes == 'Description override test', PI.id != pi_id).one()
            self.assertEqual(copied.issue_date, date.today())
            self.assertTrue(copied.pi_number.startswith('PI-' + date.today().strftime('%Y%m%d') + '-'))
            self.assertEqual(copied.items[0].display_code, '')
            self.assertEqual(copied.items[0].display_name, 'Edited nozzle')
            self.assertEqual(copied.items[0].display_specification, 'Custom spec')
        data[f'item_name_{product_id}'] = '   '
        data['csrf_token'] = self.token('/pi/create')
        self.assertEqual(self.client.post('/pi/create', data=data).status_code, 400)

    def test_customer_history_is_admin_only_and_separate_from_system_turnover(self):
        with application.app.app_context():
            customer = Customer(name='History single test', salesperson='Alice', total_deal_usd=25)
            db.session.add(customer)
            db.session.commit()
            customer_id, version = customer.id, customer.version
        self.login()
        self.assertEqual(self.client.get(f'/customers/{customer_id}/history').status_code, 403)
        self.assertEqual(self.client.get('/customers/history/import').status_code, 403)
        self.login('admin-test')
        token = self.token(f'/customers/{customer_id}/history')
        data = {'csrf_token': token, 'version': str(version), 'historical_deal_usd': '1234.56',
                'historical_deal_cutoff': '2026-09-17', 'historical_deal_note': 'Old USD business',
                'exclude_system_orders': '1'}
        response = self.client.post(f'/customers/{customer_id}/history', data=data)
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            customer = db.session.get(Customer, customer_id)
            self.assertEqual(customer.total_deal_usd, 25)
            self.assertEqual(customer.cumulative_deal_usd, 1259.56)
            self.assertEqual(customer.historical_deal_cutoff, date(2026, 9, 17))
            self.assertTrue(AuditLog.query.filter_by(entity_type='customer', entity_id=customer_id).count())
            application._recalculate_customer_deal(customer)
            db.session.commit()
            self.assertEqual(customer.total_deal_usd, 0)
            self.assertEqual(customer.historical_deal_usd, 1234.56)
            self.assertEqual(customer.cumulative_deal_usd, 1234.56)
        data['historical_deal_usd'] = '999'
        self.client.post(f'/customers/{customer_id}/history', data=data)
        with application.app.app_context():
            self.assertEqual(db.session.get(Customer, customer_id).historical_deal_usd, 1234.56)
        detail = self.client.get(f'/customers/{customer_id}').get_data(as_text=True)
        self.assertIn('历史成交额（USD）', detail)
        listing = self.client.get('/customers?search=History+single+test&deal_min=1000').get_data(as_text=True)
        self.assertIn('History single test', listing)
        self.assertIn('$1234.56', listing)

    def test_customer_history_import_preview_confirmation_and_atomic_conflicts(self):
        from customer_history import HEADERS, preview_workbook, history_values
        self.login('admin-test')
        with application.app.app_context():
            customers = [Customer(name='History import one', salesperson='Alice', historical_deal_usd=10),
                         Customer(name='History import two', salesperson='Alice'),
                         Customer(name='History ambiguous', salesperson='Alice'),
                         Customer(name='History ambiguous', salesperson='Bob')]
            db.session.add_all(customers)
            db.session.commit()
            first_id, second_id = customers[0].id, customers[1].id
        def workbook_upload(rows):
            workbook = Workbook()
            workbook.active.title = '历史成交额'
            workbook.active.append(list(HEADERS))
            for row in rows:
                workbook.active.append(row)
            output = BytesIO()
            workbook.save(output)
            output.seek(0)
            return output, 'history.xlsx'
        def preview(rows):
            response = self.client.post('/customers/history/import', data={
                'csrf_token': self.token('/customers/history/import'), 'file': workbook_upload(rows)})
            self.assertEqual(response.status_code, 200)
            page = response.get_data(as_text=True)
            match = re.search(r'name="preview_token" value="([a-f0-9]+)"', page)
            return page, match.group(1) if match else None
        rows = [[first_id, 'History import one', 100, '2026-09-17', 'First'],
                [None, 'History import two', 200, None, 'Second']]
        page, token = preview(rows)
        self.assertTrue(token)
        with application.app.app_context():
            self.assertEqual(db.session.get(Customer, first_id).historical_deal_usd, 10)
        response = self.client.post('/customers/history/import', data={
            'csrf_token': self.token('/customers/history/import'), 'action': 'confirm',
            'preview_token': token, 'exclude_system_orders': '1'})
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            self.assertEqual(db.session.get(Customer, first_id).historical_deal_usd, 100)
            self.assertEqual(db.session.get(Customer, second_id).historical_deal_usd, 200)
        page, token = preview(rows)
        self.client.post('/customers/history/import', data={
            'csrf_token': self.token('/customers/history/import'), 'action': 'confirm',
            'preview_token': token, 'exclude_system_orders': '1'})
        with application.app.app_context():
            self.assertEqual(db.session.get(Customer, first_id).historical_deal_usd, 100)
        rows[0][2], rows[1][2] = 500, 600
        page, token = preview(rows)
        with application.app.app_context():
            customer = db.session.get(Customer, second_id)
            customer.notes = 'Changed after preview'
            db.session.commit()
        response = self.client.post('/customers/history/import', data={
            'csrf_token': self.token('/customers/history/import'), 'action': 'confirm',
            'preview_token': token, 'exclude_system_orders': '1'})
        self.assertIn('未保存任何记录', response.get_data(as_text=True))
        with application.app.app_context():
            self.assertEqual(db.session.get(Customer, first_id).historical_deal_usd, 100)
            self.assertEqual(db.session.get(Customer, second_id).historical_deal_usd, 200)
        for bad in ([first_id, 'Wrong name', 10, None, None],
                    [None, 'History ambiguous', 10, None, None],
                    [first_id, None, -1, None, None],
                    [first_id, None, '=1+1', None, None],
                    [first_id, None, None, None, None]):
            page, token = preview([bad])
            self.assertIsNone(token)
            self.assertIn('未保存任何记录', page)
        page, token = preview([rows[0], rows[0]])
        self.assertIsNone(token)
        self.assertIn('重复出现', page)
        for amount in ('NaN', 'Infinity', '-1', '1.001', '', '1e100'):
            with self.assertRaises(ValueError):
                history_values(amount, '2026-09-17', '')
        template_response = self.client.get('/customers/history/template.xlsx')
        self.assertEqual(template_response.status_code, 200)
        loaded = load_workbook(BytesIO(template_response.data))
        self.assertEqual(tuple(cell.value for cell in loaded['历史成交额'][1]), HEADERS)
        loaded.close()
        csv_response = self.client.get('/customers/history/customer-ids.csv')
        self.assertEqual(csv_response.status_code, 200)
        self.assertIn('客户编号', csv_response.data.decode('utf-8-sig'))

    def test_supplier_forms_block_duplicate_names_without_changing_existing_data(self):
        self.login('admin-test')
        with application.app.app_context():
            first = Supplier(name='Duplicate Supplier Alpha', phone='Original phone')
            second = Supplier(name='Duplicate Supplier Beta', phone='Second phone')
            db.session.add_all([first, second])
            db.session.commit()
            first_id, second_id = first.id, second.id
            count = Supplier.query.count()
        response = self.client.post('/suppliers/add', data={
            'name': '  DUPLICATE SUPPLIER ALPHA  ', 'phone': 'Keep my input',
            'csrf_token': self.token('/suppliers/add')})
        self.assertEqual(response.status_code, 200)
        self.assertIn('该供应商名称已经存在。', response.get_data(as_text=True))
        self.assertIn('Keep my input', response.get_data(as_text=True))
        response = self.client.post(f'/suppliers/{second_id}/edit', data={
            'name': 'Duplicate Supplier Alpha', 'phone': 'Rejected phone',
            'csrf_token': self.token(f'/suppliers/{second_id}/edit')})
        self.assertIn('该供应商名称已经存在。', response.get_data(as_text=True))
        with application.app.app_context():
            self.assertEqual(Supplier.query.count(), count)
            second = db.session.get(Supplier, second_id)
            self.assertEqual(second.name, 'Duplicate Supplier Beta')
            self.assertEqual(second.phone, 'Second phone')
        response = self.client.post(f'/suppliers/{first_id}/edit', data={
            'name': 'Duplicate Supplier Alpha', 'phone': 'Updated phone',
            'csrf_token': self.token(f'/suppliers/{first_id}/edit')})
        self.assertEqual(response.status_code, 302)
        api_response = self.client.post('/api/suppliers', json={'name': 'duplicate supplier alpha'},
            headers={'X-CSRFToken': self.token('/suppliers')})
        self.assertEqual(api_response.status_code, 409)
        with application.app.app_context():
            self.assertEqual(db.session.get(Supplier, first_id).phone, 'Updated phone')

    def test_duplicate_pi_rows_remain_independent_and_protect_references(self):
        self.login('admin-test')
        with application.app.app_context():
            product = Product(name='Duplicate source', product_code='DUP-SOURCE', unit_price=3)
            pi = PI(pi_number='PI-DUP-ROWS', customer_id=self.alice_customer,
                    salesperson='Alice', currency='USD', exchange_rate=7, total_amount=6)
            db.session.add_all([product, pi])
            db.session.flush()
            original = PIItem(pi_id=pi.id, product_id=product.id, quantity=2, unit_price=3, amount=6)
            db.session.add(original)
            db.session.commit()
            pi_id, product_id, original_id = pi.id, product.id, original.id
            form = MultiDict({'product_order': '-1,-2', 'selected_-1': 'on', 'selected_-2': 'on',
                              'row_product_-1': str(product_id), 'row_product_-2': str(product_id),
                              'row_item_-1': str(original_id), 'qty_-1': '2', 'qty_-2': '4',
                              'unit_price_-1': '3', 'unit_price_-2': '5',
                              'item_spec_-1': 'Original', 'item_spec_-2': 'Copied spec',
                              'item_code_-2': 'CUSTOM-COPY'})
            rows, total = application._submitted_pi_items(form)
            self.assertEqual(total, 26)
            application._reconcile_pi_items(pi, rows)
            db.session.commit()
            db.session.expire_all()
            pi = db.session.get(PI, pi_id)
            self.assertEqual(len(pi.items), 2)
            self.assertEqual(pi.items[0].id, original_id)
            copy_id = pi.items[1].id
            self.assertEqual(pi.items[1].display_code, 'CUSTOM-COPY')
            self.assertEqual(pi.items[1].display_specification, 'Copied spec')
            self.assertEqual(product.product_code, 'DUP-SOURCE')
            supplier = Supplier(name='Duplicate test supplier')
            db.session.add(supplier)
            db.session.flush()
            db.session.add(Procurement(pi_id=pi_id, pi_item_id=copy_id, supplier_id=supplier.id,
                                       quantity=4, unit_price=1))
            db.session.commit()
            form.pop('selected_-2')
            removed, _ = application._submitted_pi_items(form)
            self.assertTrue(application._validate_pi_item_reconciliation(pi, removed))
            form['selected_-2'] = 'on'
            form['row_item_-2'] = str(copy_id)
            form['qty_-2'] = '1'
            reduced, _ = application._submitted_pi_items(form)
            self.assertTrue(application._validate_pi_item_reconciliation(pi, reduced))
            form['qty_-2'] = '4'
            form['product_order'] = '-2,-1'
            reordered, _ = application._submitted_pi_items(form)
            application._reconcile_pi_items(pi, reordered)
            db.session.commit()
            db.session.expire_all()
            self.assertEqual([item.id for item in db.session.get(PI, pi_id).items], [copy_id, original_id])
            form['row_item_-2'] = str(original_id)
            with self.assertRaises(Exception) as caught:
                application._match_pi_rows(pi, application._submitted_pi_items(form)[0])
            self.assertEqual(caught.exception.code, 400)

    def test_product_list_shares_keyword_search_and_preserves_filters(self):
        self.login()
        with application.app.app_context():
            products = [
                Product(name='36KD Gas Nozzle', product_code='LIST-FUZZY', specification='Copper'),
                Product(name='36KD nozzle', product_code='LIST-EXACT', image='test.png'),
                Product(name='36KD Torch', product_code='LIST-UNRELATED'),
                Product(name='36KD Inactive Nozzle', product_code='LIST-INACTIVE', active=False),
            ]
            db.session.add_all(products)
            db.session.commit()
        html = self.client.get('/products?search=36kd%20NOZZLE').get_data(as_text=True)
        self.assertIn('LIST-FUZZY', html)
        self.assertIn('LIST-EXACT', html)
        self.assertNotIn('LIST-UNRELATED', html)
        self.assertNotIn('LIST-INACTIVE', html)
        self.assertLess(html.index('LIST-EXACT'), html.index('LIST-FUZZY'))
        html = self.client.get('/products?search=36KD%20Copper').get_data(as_text=True)
        self.assertIn('LIST-FUZZY', html)
        self.assertNotIn('LIST-EXACT', html)
        html = self.client.get('/products?search=36KD%20nozzle&filter=no_image').get_data(as_text=True)
        self.assertIn('LIST-FUZZY', html)
        self.assertNotIn('LIST-EXACT', html)
        html = self.client.get('/products?search=36KD%20nozzle&filter=inactive').get_data(as_text=True)
        self.assertIn('LIST-INACTIVE', html)
        self.assertNotIn('LIST-FUZZY', html)
        html = self.client.get('/products?search=36KD%20nozzle&filter=all&sort=name_asc').get_data(as_text=True)
        self.assertIn('LIST-INACTIVE', html)
        self.assertLess(html.index('LIST-FUZZY'), html.index('LIST-EXACT'))
        html = self.client.get('/products?search=%25').get_data(as_text=True)
        self.assertNotIn('LIST-FUZZY', html)
        with application.app.app_context():
            Product.query.filter(Product.product_code.in_([
                'LIST-FUZZY', 'LIST-EXACT', 'LIST-UNRELATED', 'LIST-INACTIVE'
            ])).delete(synchronize_session=False)
            db.session.commit()

    def test_product_picker_supports_pagination_search_and_batch_ui(self):
        self.login()
        create_pi = self.client.get('/pi/create')
        self.assertEqual(create_pi.status_code, 200)
        html = create_pi.get_data(as_text=True)
        self.assertIn('id="productPickerModal"', html)
        self.assertIn('id="picker_add_selected"', html)
        self.assertIn('跨页多选', html)
        self.assertIn('id="product_order"', html)
        self.assertIn('var pickerSelectionOrder = [];', html)
        self.assertIn('var ids = pickerSelectionOrder.slice();', html)
        self.assertNotIn('Object.keys(pickerSelections)', html)

        edit_pi = self.client.get(f'/pi/{self.alice_pi}/edit')
        self.assertEqual(edit_pi.status_code, 200)
        edit_html = edit_pi.get_data(as_text=True)
        self.assertIn('id="productPickerModal"', edit_html)
        self.assertIn('id="picker_add_selected"', edit_html)
        self.assertIn('可搜索、翻页并跨页多选', edit_html)
        self.assertIn('if (structuralEditLocked()) return;', edit_html)
        self.assertIn('var ids = pickerSelectionOrder.slice();', edit_html)
        self.assertNotIn('id="product_dropdown"', edit_html)

        response = self.client.get('/api/products/search?picker=1&page=1&per_page=24')
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertIn('items', payload)
        self.assertEqual(payload['page'], 1)
        self.assertLessEqual(len(payload['items']), 24)
        self.assertGreaterEqual(payload['total'], 1)

        matched = self.client.get('/api/products/search?picker=1&q=ORIG&page=1&per_page=24').get_json()
        product = next(item for item in matched['items'] if item['name'] == 'Original Product')
        self.assertEqual(product['code'], 'ORIG')
        self.assertEqual(product['product_code'], 'ORIG')

        legacy = self.client.get('/api/products/search?q=ORIG').get_json()
        self.assertIsInstance(legacy, list)
        legacy_product = next(item for item in legacy if item['name'] == 'Original Product')
        self.assertEqual(legacy_product['code'], 'ORIG')

    def test_referenced_product_is_disabled_and_can_be_enabled(self):
        self.login('admin-test')
        image_name = 'referenced-product.png'
        image_path = os.path.join(application.app.config['UPLOAD_DIR'], image_name)
        with open(image_path, 'wb') as image_file:
            image_file.write(b'product-image')

        with application.app.app_context():
            customer = Customer(name='Product Retire Customer', salesperson='Alice')
            product = Product(
                name='Referenced Product To Retire',
                product_code='REF-RETIRE',
                unit_price=3,
                image=image_name,
            )
            pi = PI(
                pi_number='PI-PRODUCT-RETIRE', customer=customer,
                salesperson='Alice', currency='USD', exchange_rate=7,
                total_amount=3,
            )
            db.session.add_all([customer, product, pi])
            db.session.flush()
            db.session.add(PIItem(
                pi_id=pi.id, product_id=product.id,
                quantity=1, unit_price=3, amount=3,
            ))
            db.session.commit()
            product_id = product.id
            pi_id = pi.id

        response = self.client.post(
            f'/api/products/{product_id}/delete',
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()['action'], 'disabled')
        self.assertEqual(response.get_json()['reference_count'], 1)
        self.assertTrue(os.path.isfile(image_path))

        with application.app.app_context():
            product = db.session.get(Product, product_id)
            self.assertIsNotNone(product)
            self.assertFalse(product.active)
            self.assertEqual(PIItem.query.filter_by(product_id=product_id).count(), 1)

        search = self.client.get('/api/products/search?q=REF-RETIRE').get_json()
        self.assertEqual(search, [])

        enabled = self.client.post(
            f'/api/products/{product_id}/enable',
            headers={'X-CSRFToken': self.token('/products?filter=inactive')},
        )
        self.assertEqual(enabled.status_code, 200)
        self.assertTrue(enabled.get_json()['success'])
        self.assertEqual(
            [item['product_code'] for item in self.client.get(
                '/api/products/search?q=REF-RETIRE'
            ).get_json()],
            ['REF-RETIRE'],
        )

        with application.app.app_context():
            PIItem.query.filter_by(product_id=product_id).delete()
            db.session.delete(db.session.get(PI, pi_id))
            db.session.delete(db.session.get(Product, product_id))
            db.session.commit()
        os.remove(image_path)

    def test_unreferenced_shared_product_image_is_removed_only_after_last_product(self):
        self.login('admin-test')
        image_name = 'legacy-shared-product.png'
        image_path = os.path.join(application.app.config['UPLOAD_DIR'], image_name)
        with open(image_path, 'wb') as image_file:
            image_file.write(b'legacy-shared-image')

        with application.app.app_context():
            first = Product(name='Legacy Shared First', product_code='SHARED-1', image=image_name)
            second = Product(name='Legacy Shared Second', product_code='SHARED-2', image=image_name)
            db.session.add_all([first, second])
            db.session.commit()
            first_id, second_id = first.id, second.id

        first_delete = self.client.post(
            f'/api/products/{first_id}/delete',
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(first_delete.get_json()['action'], 'deleted')
        self.assertTrue(os.path.isfile(image_path))

        second_delete = self.client.post(
            f'/api/products/{second_id}/delete',
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(second_delete.get_json()['action'], 'deleted')
        self.assertFalse(os.path.exists(image_path))

    def test_product_copy_owns_an_independent_image_file(self):
        self.login('admin-test')
        image_name = 'copy-source-product.png'
        image_path = os.path.join(application.app.config['UPLOAD_DIR'], image_name)
        with open(image_path, 'wb') as image_file:
            image_file.write(b'copy-source-image')

        with application.app.app_context():
            source = Product(name='Copy Image Source', product_code='COPY-IMAGE', image=image_name)
            db.session.add(source)
            db.session.commit()
            source_id = source.id

        copied = self.client.post(
            f'/api/products/{source_id}/copy',
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(copied.status_code, 200)
        copied_product = copied.get_json()['product']
        self.assertNotEqual(copied_product['image'], image_name)
        copied_path = os.path.join(application.app.config['UPLOAD_DIR'], copied_product['image'])
        self.assertTrue(os.path.isfile(copied_path))
        self.assertTrue(os.path.isfile(image_path))

        deleted_copy = self.client.post(
            f"/api/products/{copied_product['id']}/delete",
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertEqual(deleted_copy.get_json()['action'], 'deleted')
        self.assertFalse(os.path.exists(copied_path))
        self.assertTrue(os.path.isfile(image_path))

        self.client.post(
            f'/api/products/{source_id}/delete',
            headers={'X-CSRFToken': self.token('/products')},
        )
        self.assertFalse(os.path.exists(image_path))

    def test_pi_rate_defaults_from_settings_but_existing_pi_keeps_its_rate(self):
        with application.app.app_context():
            original_settings = application._load_settings().copy()
            settings = original_settings.copy()
            settings['exchange_rate'] = '8.25'
            application._save_settings(settings)
            pi = db.session.get(PI, self.alice_pi)
            original_rate = pi.exchange_rate
            pi.exchange_rate = 6.75
            db.session.commit()

        try:
            self.login('alice')
            create_html = self.client.get('/pi/create').get_data(as_text=True)
            self.assertIn('name="exchange_rate"', create_html)
            self.assertIn('value="8.25"', create_html)

            edit_html = self.client.get(f'/pi/{self.alice_pi}/edit').get_data(as_text=True)
            self.assertIn('本单业务汇率', edit_html)
            self.assertIn('value="6.75"', edit_html)

            with application.app.app_context():
                self.assertEqual(db.session.get(PI, self.alice_pi).exchange_rate, 6.75)
        finally:
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                pi.exchange_rate = original_rate
                db.session.commit()
                application._save_settings(original_settings)

    def test_pi_list_combines_payment_order_and_customs_filters(self):
        self.login('admin-test')
        with application.app.app_context():
            product = Product.query.filter_by(product_code='ORIG').one()
            supplier = Supplier(name='PI List Filter Supplier')
            pi = PI(
                pi_number='PI-FILTER-STATUS-001',
                customer_id=self.alice_customer,
                salesperson='Alice',
                currency='USD',
                total_amount=10,
                received_amount=5,
                paid=False,
                customs_required=True,
                procurement_confirmed=False,
                shipping_completed=False,
            )
            db.session.add_all([supplier, pi])
            db.session.flush()
            item = PIItem(
                pi_id=pi.id, product_id=product.id,
                quantity=1, unit_price=10, amount=10,
            )
            db.session.add(item)
            db.session.commit()
            pi_id = pi.id
            supplier_id = supplier.id
            item_id = item.id

        try:
            combined = self.client.get(
                '/pi/list?payment_status=partial&order_status=pending&customs_status=required'
            )
            self.assertEqual(combined.status_code, 200)
            combined_html = combined.get_data(as_text=True)
            self.assertIn('PI-FILTER-STATUS-001', combined_html)
            self.assertIn('name="payment_status"', combined_html)
            self.assertIn('name="order_status"', combined_html)
            self.assertIn('name="customs_status"', combined_html)
            self.assertIn('回款状态：<strong>部分回款</strong>', combined_html)
            self.assertIn('订单进度：<strong>待采购</strong>', combined_html)
            self.assertIn('报关登记：<strong>需要报关</strong>', combined_html)

            paid_only = self.client.get('/pi/list?payment_status=paid')
            self.assertNotIn('PI-FILTER-STATUS-001', paid_only.get_data(as_text=True))
            unregistered = self.client.get('/pi/list?customs_status=unregistered')
            self.assertNotIn('PI-FILTER-STATUS-001', unregistered.get_data(as_text=True))

            with application.app.app_context():
                db.session.add(Procurement(
                    pi_id=pi_id,
                    pi_item_id=item_id,
                    supplier_id=supplier_id,
                    unit_price=8,
                    quantity=1,
                    total=8,
                ))
                db.session.commit()
            partial = self.client.get('/pi/list?order_status=partial')
            self.assertIn('PI-FILTER-STATUS-001', partial.get_data(as_text=True))

            with application.app.app_context():
                pi = db.session.get(PI, pi_id)
                pi.procurement_confirmed = True
                pi.procurement_status = '采购完成'
                db.session.commit()
            purchased = self.client.get('/pi/list?order_status=purchased')
            self.assertIn('PI-FILTER-STATUS-001', purchased.get_data(as_text=True))

            with application.app.app_context():
                pi = db.session.get(PI, pi_id)
                pi.shipping_completed = True
                pi.procurement_status = '发货完成'
                pi.customs_required = False
                db.session.commit()
            shipped = self.client.get(
                '/pi/list?order_status=shipped&customs_status=not_required'
            )
            self.assertIn('PI-FILTER-STATUS-001', shipped.get_data(as_text=True))
        finally:
            with application.app.app_context():
                pi = db.session.get(PI, pi_id)
                if pi:
                    db.session.delete(pi)
                    db.session.flush()
                supplier = db.session.get(Supplier, supplier_id)
                if supplier:
                    db.session.delete(supplier)
                db.session.commit()

    def test_customer_notes_are_visible_on_pi_pages_and_available_to_templates(self):
        with application.app.app_context():
            customer = db.session.get(Customer, self.alice_customer)
            original_notes = customer.notes
            customer.notes = 'External customer instruction'
            db.session.commit()
            pi = db.session.get(PI, self.alice_pi)
            self.assertEqual(_fixed_values(pi, None)['customer_notes'], 'External customer instruction')
        try:
            self.login('alice')
            create_html = self.client.get(
                f'/pi/create?customer_id={self.alice_customer}'
            ).get_data(as_text=True)
            edit_html = self.client.get(f'/pi/{self.alice_pi}/edit').get_data(as_text=True)
            preview_html = self.client.get(f'/pi/{self.alice_pi}/preview').get_data(as_text=True)
            self.assertIn('客户备注', create_html)
            self.assertIn('External customer instruction', create_html)
            self.assertIn('External customer instruction', edit_html)
            self.assertIn('CUSTOMER NOTES', preview_html)

            with application.app.app_context():
                template_id = DocumentTemplate.query.filter_by(code='system-default').one().id
                item = db.session.get(PI, self.alice_pi).items[0]
                item_id = item.id
                item_quantity = item.quantity
                item_price = item.unit_price
            exported = self.client.post(f'/pi/{self.alice_pi}/export', data={
                'template_id': template_id,
                'price_terms': 'FOB Shanghai',
                'output_format': 'xlsx',
                'mode': 'download',
                f'qty_{item_id}': str(item_quantity),
                f'price_{item_id}': str(item_price),
                'csrf_token': self.token(f'/pi/{self.alice_pi}/export'),
            })
            self.assertEqual(exported.status_code, 200)
            rendered = load_workbook(BytesIO(exported.data), data_only=False)
            rendered_values = [
                str(cell.value) for row in rendered.active.iter_rows() for cell in row
                if cell.value is not None
            ]
            rendered.close()
            exported.close()
            self.assertNotIn('External customer instruction', rendered_values)
            self.assertIn('Trade Terms:', rendered_values)
            self.assertIn('FOB Shanghai', rendered_values)

            template_path = os.path.join(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                'assets', 'system_default_pi_template.xlsx',
            )
            workbook = load_workbook(template_path, data_only=False)
            values = [
                str(cell.value) for row in workbook.active.iter_rows() for cell in row
                if cell.value is not None
            ]
            workbook.close()
            self.assertIn('{{price_terms}}', values)
            self.assertNotIn('{{customer_notes}}', values)
        finally:
            with application.app.app_context():
                customer = db.session.get(Customer, self.alice_customer)
                customer.notes = original_notes
                db.session.commit()

    def test_dingtalk_switches_and_shu_kei_group_routing(self):
        original_settings = application._load_settings().copy()
        try:
            enabled = {
                'dingtalk_webhook': 'https://example.invalid/default',
                'dingtalk_shu_kei_webhook': 'https://example.invalid/shu-kei',
                'dingtalk_report_enabled': '1',
                'dingtalk_task_enabled': '0',
            }
            self.assertEqual(
                application._dingtalk_report_webhook('Shu Kei', enabled),
                enabled['dingtalk_shu_kei_webhook'],
            )
            self.assertEqual(
                application._dingtalk_report_webhook('Alice', enabled),
                enabled['dingtalk_webhook'],
            )
            self.assertFalse(application._setting_enabled(enabled, 'dingtalk_task_enabled'))
            disabled = dict(enabled, dingtalk_report_enabled='0')
            self.assertEqual(application._dingtalk_report_webhook('Shu Kei', disabled), '')

            self.login('admin-test')
            response = self.client.post('/settings', data={
                'dingtalk_report_enabled': 'on',
                'exchange_rate': '7.2',
                'csrf_token': self.token('/settings'),
            })
            self.assertEqual(response.status_code, 302)
            saved = application._load_settings()
            self.assertEqual(saved['dingtalk_report_enabled'], '1')
            self.assertEqual(saved['dingtalk_task_enabled'], '0')
            html = self.client.get('/settings').get_data(as_text=True)
            self.assertIn('id="dingtalk_report_enabled"', html)
            self.assertIn('id="dingtalk_task_enabled"', html)
            self.assertIn('Shu Kei 喜报群 Webhook', html)
        finally:
            application._save_settings(original_settings)

    def test_sales_performance_uses_fixed_seven_not_pi_or_default_rate(self):
        with application.app.app_context():
            original_settings = application._load_settings().copy()
            settings = original_settings.copy()
            settings['exchange_rate'] = '8.88'
            application._save_settings(settings)
            pi = db.session.get(PI, self.alice_pi)
            original = {
                'currency': pi.currency,
                'exchange_rate': pi.exchange_rate,
                'received_amount': pi.received_amount,
                'paid': pi.paid,
            }
            pi.currency = 'RMB'
            pi.exchange_rate = 9.5
            pi.received_amount = 70
            pi.paid = True
            db.session.commit()
        try:
            self.login('alice')
            html = self.client.get('/sales-stats').get_data(as_text=True)
            self.assertIn('1 美元 = 7.00 人民币', html)
            self.assertIn('$10.00', html)
        finally:
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                for key, value in original.items():
                    setattr(pi, key, value)
                db.session.commit()
                application._save_settings(original_settings)

    def test_procurement_status_is_record_driven_and_shipping_is_scoped_to_pi_owner(self):
        with application.app.app_context():
            Procurement.query.filter_by(pi_id=self.alice_pi).delete()
            Payment.query.filter_by(pi_id=self.alice_pi).delete()
            pi = db.session.get(PI, self.alice_pi)
            pi.received_amount = 0
            pi.paid = False
            pi.procurement_confirmed = False
            pi.shipping_completed = False
            pi.procurement_status = '未回款'
            db.session.commit()

        self.login('alice')
        salesperson_page = self.client.get('/pi/list').get_data(as_text=True)
        self.assertIn('订单进度', salesperson_page)
        self.assertIn('未回款', salesperson_page)
        self.assertNotIn('onchange="updateProcurementStatus(', salesperson_page)
        unavailable_record = self.client.get(
            f'/api/pi/{self.alice_pi}/shipping-record',
        )
        self.assertEqual(unavailable_record.status_code, 200)
        self.assertFalse(unavailable_record.get_json()['can_ship'])
        self.assertIn('尚未回款', unavailable_record.get_json()['unavailable_reason'])
        denied = self.client.post(
            f'/api/pi/{self.alice_pi}/shipping-complete',
            json={'completed': True},
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(denied.status_code, 403)

        self.client.get('/logout')
        self.login('admin-test')
        no_payment = self.client.post(
            f'/api/pi/{self.alice_pi}/shipping-complete',
            json={'completed': True},
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(no_payment.status_code, 400)
        self.assertIn('尚未回款', no_payment.get_json()['error'])

        with application.app.app_context():
            payment = Payment(
                pi_id=self.alice_pi,
                amount=5,
                fee=0,
                order_no='PROC-STATUS-001',
                order_no_normalized='proc-status-001',
                idempotency_key='proc-status-test-001',
                receiving_account_id=self.approved_account,
                receiving_account_name='Approved',
                receiving_account_currency='USD',
            )
            db.session.add(payment)
            pi = db.session.get(PI, self.alice_pi)
            application._recalculate_pi_payments(pi)
            db.session.commit()
            self.assertEqual(pi.effective_procurement_status, '待采购')
            item_id = PIItem.query.filter_by(pi_id=self.alice_pi).one().id

        new_supplier = self.client.post(
            '/api/suppliers',
            json={
                'name': 'Procurement Test Supplier',
                'contact_person': 'Quick Add Contact',
                'phone': '123456',
            },
            headers={'X-CSRFToken': self.token(f'/procurement/{self.alice_pi}')},
        )
        self.assertEqual(new_supplier.status_code, 201)
        self.assertTrue(new_supplier.get_json()['success'])
        supplier_id = new_supplier.get_json()['supplier']['id']

        admin_page = self.client.get('/pi/list').get_data(as_text=True)
        self.assertNotIn('updateProcurementStatus(', admin_page)
        self.assertIn('待采购', admin_page)
        procurement_page = self.client.get(f'/procurement/{self.alice_pi}')
        self.assertEqual(procurement_page.status_code, 200)
        procurement_html = procurement_page.get_data(as_text=True)
        self.assertIn('预览采购汇总', procurement_html)
        self.assertIn('id="procurementPreviewModal"', procurement_html)
        self.assertIn('preview-product-image', procurement_html)
        self.assertIn("row.querySelector('.proc-image-cell img')", procurement_html)
        self.assertIn('<th>照片</th><th>产品</th><th>编码</th><th>规格</th><th class="text-end">数量</th><th class="text-end">采购单价</th>', procurement_html)
        self.assertIn('供应商采购单预览', procurement_html)
        self.assertIn('已按供应商拆分为', procurement_html)
        self.assertIn('导出该供应商采购单', procurement_html)
        self.assertIn('supplier-orders/export', procurement_html)
        self.assertIn('PI 单价（USD / RMB）', procurement_html)
        self.assertIn('$10.00', procurement_html)
        self.assertIn('data-pi-quantity="1"', procurement_html)
        self.assertRegex(procurement_html, r'class="[^"]*proc-qty[^"]*"[^>]*value="1"')
        self.assertNotIn('订单单价（RMB）', procurement_html)
        self.assertIn('供应商采购运费成本', procurement_html)
        self.assertIn('id="newSupplierModal"', procurement_html)
        self.assertIn('新增供应商', procurement_html)
        self.assertIn('<i class="bi bi-save"></i> 暂存', procurement_html)
        self.assertIn('id="procurementReadiness"', procurement_html)
        self.assertIn('只能暂存', procurement_html)
        self.assertRegex(procurement_html, r'id="confirmBtn"[^>]*disabled')
        self.assertIn("priceValue !== ''", procurement_html)
        self.assertIn('允许填 0', procurement_html)
        self.assertIn('id="procurementRiskSummary"', procurement_html)
        self.assertIn('缺少采购信息', procurement_html)
        self.assertIn('采购价高于 PI', procurement_html)
        self.assertIn('>序号</th>', procurement_html)
        self.assertIn('class="text-center proc-sequence">1</td>', procurement_html)
        self.assertIn('proc-missing-field', procurement_html)
        self.assertIn('proc-high-price-field', procurement_html)
        self.assertIn('highPriceCount', procurement_html)
        self.assertIn('.proc-price::-webkit-inner-spin-button', procurement_html)
        self.assertIn('-moz-appearance: textfield', procurement_html)
        self.assertIn('利润 = 订单产品总金额 − 采购产品总金额 − 供应商采购运费', procurement_html)
        self.assertNotIn('实际运费成本（内部）', procurement_html)
        self.assertNotIn('客户费用 / 折扣（对外）', procurement_html)
        incomplete = self.client.post(
            f'/api/procurement/{self.alice_pi}/confirm',
            headers={'X-CSRFToken': self.token(f'/procurement/{self.alice_pi}')},
        )
        self.assertEqual(incomplete.status_code, 400)
        self.assertIn('尚未完整', incomplete.get_json()['error'])

        draft_only = self.client.post(
            '/api/procurement/save',
            json={
                'pi_id': self.alice_pi,
                'supplier_freight_cost': 1.5,
                'exchange_rate': 8.25,
                'procurement_date': '2026-09-02',
                'draft': True,
                'items': [],
            },
            headers={'X-CSRFToken': self.token(f'/procurement/{self.alice_pi}')},
        )
        self.assertEqual(draft_only.status_code, 200)
        self.assertTrue(draft_only.get_json()['draft'])

        missing_price = self.client.post(
            '/api/procurement/save',
            json={
                'pi_id': self.alice_pi,
                'supplier_freight_cost': 1.5,
                'exchange_rate': 8.25,
                'procurement_date': '2026-09-02',
                'draft': True,
                'items': [{
                    'pi_item_id': item_id,
                    'supplier_id': supplier_id,
                    'quantity': 1,
                }],
            },
            headers={'X-CSRFToken': self.token(f'/procurement/{self.alice_pi}')},
        )
        self.assertEqual(missing_price.status_code, 400)
        self.assertIn('缺少采购单价', missing_price.get_json()['error'])

        saved = self.client.post(
            '/api/procurement/save',
            json={
                'pi_id': self.alice_pi,
                'supplier_freight_cost': 1.5,
                'exchange_rate': 8.25,
                'procurement_date': '2026-09-02',
                'items': [{
                    'pi_item_id': item_id,
                    'supplier_id': supplier_id,
                    'unit_price': 0,
                    'quantity': 1,
                    'note': 'test',
                }],
            },
            headers={'X-CSRFToken': self.token(f'/procurement/{self.alice_pi}')},
        )
        self.assertEqual(saved.status_code, 200)
        self.assertTrue(saved.get_json()['complete'])
        with application.app.app_context():
            saved_procurement = Procurement.query.filter_by(pi_id=self.alice_pi).one()
            self.assertEqual(saved_procurement.unit_price, 0)
            self.assertEqual(db.session.get(PI, self.alice_pi).effective_procurement_status, '部分采购')
            self.assertIsNotNone(AuditLog.query.filter_by(
                entity_type='supplier', entity_id=supplier_id, action='create'
            ).first())
        procurement_list = self.client.get('/procurement').get_data(as_text=True)
        self.assertIn('部分采购', procurement_list)
        self.assertIn('继续采购', procurement_list)

        export_payload = {
            'procurement_date': '2026-09-02',
            'items': [{
                'pi_item_id': item_id,
                'supplier_id': supplier_id,
                'unit_price': 7,
                'quantity': 1,
                'note': 'test',
            }],
        }
        supplier_export = self.client.post(
            f'/procurement/{self.alice_pi}/supplier-orders/export',
            json=dict(export_payload, supplier_id=supplier_id),
            headers={'X-CSRFToken': self.token(f'/procurement/{self.alice_pi}')},
        )
        self.assertEqual(supplier_export.status_code, 200)
        self.assertEqual(
            supplier_export.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        workbook = load_workbook(BytesIO(supplier_export.data))
        workbook_text = '\n'.join(
            str(cell.value)
            for row in workbook.active.iter_rows()
            for cell in row
            if cell.value is not None
        )
        self.assertIn('采购订单', workbook_text)
        self.assertIn('Procurement Test Supplier', workbook_text)
        self.assertIn('Original Product', workbook_text)
        self.assertIn('ORIG', workbook_text)
        self.assertNotIn('Alice Customer', workbook_text)
        self.assertNotIn('利润', workbook_text)

        bundle_export = self.client.post(
            f'/procurement/{self.alice_pi}/supplier-orders/export',
            json=export_payload,
            headers={'X-CSRFToken': self.token(f'/procurement/{self.alice_pi}')},
        )
        self.assertEqual(bundle_export.status_code, 200)
        self.assertEqual(bundle_export.mimetype, 'application/zip')
        with zipfile.ZipFile(BytesIO(bundle_export.data)) as bundle:
            self.assertEqual(len(bundle.namelist()), 1)
            self.assertTrue(bundle.namelist()[0].endswith('.xlsx'))
        with application.app.app_context():
            saved_pi = db.session.get(PI, self.alice_pi)
            self.assertEqual(saved_pi.exchange_rate, 8.25)
            self.assertEqual(saved_pi.supplier_freight_cost, 1.5)

        confirmed = self.client.post(
            f'/api/procurement/{self.alice_pi}/confirm',
            headers={'X-CSRFToken': self.token(f'/procurement/{self.alice_pi}')},
        )
        self.assertEqual(confirmed.status_code, 200)
        self.assertEqual(confirmed.get_json()['procurement_status'], '采购完成')

        confirmed_procurement_html = self.client.get(
            f'/procurement/{self.alice_pi}'
        ).get_data(as_text=True)
        self.assertIn('发货由业务员在 PI 列表登记', confirmed_procurement_html)
        self.assertNotIn('setShippingCompleted(', confirmed_procurement_html)

        self.client.get('/logout')
        self.login('bob')
        bob_read = self.client.get(f'/api/pi/{self.alice_pi}/shipping-record')
        self.assertEqual(bob_read.status_code, 403)
        bob_write = self.client.post(
            f'/api/pi/{self.alice_pi}/shipping-record',
            json={'shipping_date': '2026-09-07'},
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(bob_write.status_code, 403)

        self.client.get('/logout')
        self.login('alice')
        available_record = self.client.get(f'/api/pi/{self.alice_pi}/shipping-record')
        self.assertEqual(available_record.status_code, 200)
        self.assertTrue(available_record.get_json()['can_ship'])
        missing_date = self.client.post(
            f'/api/pi/{self.alice_pi}/shipping-record',
            json={'tracking_no': 'SF123456'},
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(missing_date.status_code, 400)

        shipped = self.client.post(
            f'/api/pi/{self.alice_pi}/shipping-record',
            json={
                'shipping_date': '2026-09-07',
                'tracking_no': 'SF123456',
                'note': '已交付物流',
            },
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(shipped.status_code, 200)
        self.assertEqual(shipped.get_json()['procurement_status'], '发货完成')
        self.assertEqual(shipped.get_json()['shipping_date'], '2026-09-07')
        self.assertEqual(shipped.get_json()['tracking_no'], 'SF123456')
        self.assertEqual(shipped.get_json()['recorded_by'], 'alice')
        self.assertIn('发货完成', self.client.get('/pi/list').get_data(as_text=True))

        duplicate = self.client.post(
            f'/api/pi/{self.alice_pi}/shipping-record',
            json={'shipping_date': '2026-09-08'},
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(duplicate.status_code, 409)

        self.client.get('/logout')
        self.login('admin-test')
        admin_record = self.client.get(f'/api/pi/{self.alice_pi}/shipping-record')
        self.assertEqual(admin_record.status_code, 200)
        self.assertTrue(admin_record.get_json()['shipping_completed'])
        confirmation_required = self.client.post(
            f'/api/pi/{self.alice_pi}/shipping-complete',
            json={'completed': False},
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(confirmation_required.status_code, 409)
        self.assertTrue(confirmation_required.get_json()['confirmation_required'])
        with application.app.app_context():
            self.assertTrue(db.session.get(PI, self.alice_pi).shipping_completed)
        reverted = self.client.post(
            f'/api/pi/{self.alice_pi}/shipping-complete',
            json={'completed': False, 'confirm_downstream_reset': True},
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(reverted.status_code, 200)
        self.assertEqual(reverted.get_json()['procurement_status'], '采购完成')

        with application.app.app_context():
            Procurement.query.filter_by(pi_id=self.alice_pi).delete()
            Payment.query.filter_by(pi_id=self.alice_pi).delete()
            pi = db.session.get(PI, self.alice_pi)
            pi.exchange_rate = 7
            pi.supplier_freight_cost = 0
            pi.shipping_date = None
            pi.shipping_tracking_no = ''
            pi.shipping_record_note = ''
            pi.shipping_recorded_at = None
            pi.shipping_recorded_by = ''
            application._recalculate_pi_payments(pi)
            Supplier.query.filter_by(id=supplier_id).delete()
            db.session.commit()
            self.assertEqual(pi.effective_procurement_status, '待采购')

    def test_csrf_is_required_and_customer_owner_is_forced(self):
        self.login()
        salesperson_form = self.client.get('/pi/create').get_data(as_text=True)
        self.assertIn('id="qa_cust_salesperson" value="Alice"', salesperson_form)
        self.assertIn('value="Alice" readonly aria-label="当前业务员"', salesperson_form)
        self.assertEqual(self.client.post('/api/customers/add', data={'name': 'Rejected'}).status_code, 400)
        response = self.client.post('/api/customers/add', data={
            'name': 'Owned Customer',
            'salesperson': 'Bob',
            'csrf_token': self.token('/customers'),
        })
        self.assertEqual(response.status_code, 200)
        with application.app.app_context():
            self.assertEqual(Customer.query.filter_by(name='Owned Customer').one().salesperson, 'Alice')
        self.client.get('/logout')
        self.login('admin-test')
        admin_form = self.client.get('/pi/create').get_data(as_text=True)
        self.assertIn('id="qa_cust_salesperson" required', admin_form)
        missing_salesperson = self.client.post('/api/customers/add', data={
            'name': 'No Owner',
            'csrf_token': self.token('/pi/create'),
        })
        self.assertEqual(missing_salesperson.status_code, 400)
        self.assertIn('必须指定业务员', missing_salesperson.get_json()['error'])

    def test_profit_report_is_admin_only_and_sales_can_add_own_actual_cost(self):
        with application.app.app_context():
            original_grand_total = db.session.get(PI, self.alice_pi).grand_total
            product = Product.query.filter_by(product_code='ORIG').one()
            bob_pi = PI(
                pi_number='PI-TEST-BOB-COST', customer_id=self.bob_customer,
                salesperson='Bob', currency='USD', total_amount=10,
            )
            db.session.add(bob_pi)
            db.session.flush()
            db.session.add(PIItem(
                pi_id=bob_pi.id, product_id=product.id,
                quantity=1, unit_price=10, amount=10,
            ))
            db.session.commit()
            bob_pi_id = bob_pi.id

        self.login('admin-test')
        page = self.client.get('/pi/list').get_data(as_text=True)
        self.assertIn(f'openExpense({self.alice_pi},', page)
        self.assertIn('新增真实成本', page)
        self.assertIn('利润表', page)
        self.assertLess(page.index('采购'), page.index('利润表'))
        self.assertEqual(self.client.get('/procurement').status_code, 200)
        profit_page = self.client.get('/profit-report')
        self.assertEqual(profit_page.status_code, 200)
        profit_html = profit_page.get_data(as_text=True)
        self.assertIn('管理员专属', profit_html)
        self.assertIn('每单利润明细', profit_html)
        self.assertIn('默认显示关键金额', profit_html)
        self.assertIn('查看计算口径', profit_html)
        self.assertIn('净利润 = 净收入 − 订单真实成本合计', profit_html)
        self.assertIn('订单真实成本结构', profit_html)
        self.assertIn('采购产品成本', profit_html)
        invalid_amount = self.client.post('/api/expenses', data={
            'pi_id': self.alice_pi,
            'category': '实际运费',
            'amount': 'nan',
            'currency': 'RMB',
            'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(invalid_amount.status_code, 400)
        invalid_currency = self.client.post('/api/expenses', data={
            'pi_id': self.alice_pi,
            'category': '实际运费',
            'amount': '1',
            'currency': 'EUR',
            'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(invalid_currency.status_code, 400)
        response = self.client.post('/api/expenses', data={
            'pi_id': self.alice_pi,
            'category': '实际运费',
            'amount': '25.50',
            'currency': 'RMB',
            'note': 'PI 列表快捷添加',
            'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(response.status_code, 200)
        expense_id = response.get_json()['expense']['id']
        with application.app.app_context():
            expense = db.session.get(Expense, expense_id)
            self.assertEqual(expense.pi_id, self.alice_pi)
            self.assertEqual(expense.amount, 25.5)
            self.assertEqual(db.session.get(PI, self.alice_pi).grand_total, original_grand_total)
            self.assertIsNotNone(AuditLog.query.filter_by(
                entity_type='expense', entity_id=expense_id, action='create').first())
        self.client.get('/logout')
        self.login('alice')
        salesperson_page = self.client.get('/pi/list').get_data(as_text=True)
        self.assertIn('已登记（1笔）', salesperson_page)
        self.assertIn('已登记记录', salesperson_page)
        history_response = self.client.get(f'/api/expenses?pi_id={self.alice_pi}')
        self.assertEqual(history_response.status_code, 200)
        self.assertEqual(len(history_response.get_json()), 1)
        self.assertNotIn('利润表', salesperson_page)
        self.assertNotIn('href="/procurement"', salesperson_page)
        self.assertEqual(self.client.get('/profit-report').status_code, 403)
        self.assertEqual(self.client.get('/fees').status_code, 403)
        self.assertEqual(self.client.get('/procurement').status_code, 403)
        self.assertEqual(self.client.get(f'/procurement/{self.alice_pi}').status_code, 403)
        self.assertEqual(self.client.get(f'/api/procurement/{self.alice_pi}').status_code, 403)
        from PIL import Image
        receipt = BytesIO()
        Image.new('RGB', (4, 4), '#ffffff').save(receipt, format='PNG')
        receipt.seek(0)
        own_response = self.client.post('/api/expenses', data={
            'pi_id': self.alice_pi,
            'category': '报关费',
            'amount': '1',
            'currency': 'RMB',
            'attachment': (receipt, 'customs-receipt.png'),
            'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(own_response.status_code, 200)
        own_expense = own_response.get_json()['expense']
        own_expense_id = own_expense['id']
        self.assertTrue(own_expense['has_attachment'])
        history_response = self.client.get(f'/api/expenses?pi_id={self.alice_pi}')
        self.assertEqual(len(history_response.get_json()), 2)
        receipt_response = self.client.get(f'/expenses/{own_expense_id}/attachment')
        self.assertEqual(receipt_response.status_code, 200)
        self.assertEqual(receipt_response.mimetype, 'image/png')
        self.assertIn('no-store', receipt_response.headers.get('Cache-Control', ''))
        detail_html = self.client.get(f'/pi/{self.alice_pi}').get_data(as_text=True)
        self.assertIn('订单真实成本（内部）', detail_html)
        self.assertIn('查看图片', detail_html)
        denied = self.client.post('/api/expenses', data={
            'pi_id': bob_pi_id,
            'category': '报关费',
            'amount': '1',
            'currency': 'RMB',
            'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(denied.status_code, 403)
        invalid_image = self.client.post('/api/expenses', data={
            'pi_id': self.alice_pi,
            'category': '报关费',
            'amount': '1',
            'currency': 'RMB',
            'attachment': (BytesIO(b'not-an-image'), 'fake.png'),
            'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(invalid_image.status_code, 400)
        self.client.get('/logout')
        self.login('bob')
        self.assertEqual(self.client.get(f'/expenses/{own_expense_id}/attachment').status_code, 403)
        self.client.get('/logout')
        self.login('admin-test')
        with application.app.app_context():
            receipt_name = db.session.get(Expense, own_expense_id).attachment
            receipt_path = os.path.join(application.app.config['UPLOAD_DIR'], receipt_name)
            self.assertTrue(os.path.isfile(receipt_path))
        delete_response = self.client.delete(
            f'/api/expenses/{own_expense_id}',
            headers={'X-CSRFToken': self.token('/pi/list')},
        )
        self.assertEqual(delete_response.status_code, 200)
        self.assertFalse(os.path.exists(receipt_path))
        with application.app.app_context():
            db.session.delete(db.session.get(Expense, expense_id))
            db.session.delete(db.session.get(PI, bob_pi_id))
            db.session.commit()

    def test_procurement_and_profit_reports_share_salesperson_and_date_filters(self):
        with application.app.app_context():
            product = Product.query.filter_by(product_code='ORIG').one()
            filtered_pi = PI(
                pi_number='PI-TEST-REPORT-FILTER',
                customer_id=self.bob_customer,
                salesperson='Bob',
                issue_date=application.date(2024, 4, 15),
                currency='USD',
                total_amount=18,
                received_amount=18,
                paid=True,
            )
            db.session.add(filtered_pi)
            db.session.flush()
            db.session.add(PIItem(
                pi_id=filtered_pi.id,
                product_id=product.id,
                quantity=1,
                unit_price=18,
                amount=18,
            ))
            db.session.commit()
            filtered_pi_id = filtered_pi.id

        self.login('admin-test')
        query_string = '?salesperson=Bob&date_from=2024-04-01&date_to=2024-04-30'
        for path in ('/procurement', '/profit-report'):
            response = self.client.get(path + query_string)
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            self.assertIn('PI-TEST-REPORT-FILTER', html)
            self.assertIn('name="salesperson"', html)
            self.assertIn('name="date_from" value="2024-04-01"', html)
            self.assertIn('name="date_to" value="2024-04-30"', html)
            self.assertIn('value="this_month"', html)
            self.assertIn('value="last_month"', html)
            self.assertIn('value="this_year"', html)
            self.assertIn('当前结果：<strong>1</strong> 份 PI', html)

            excluded = self.client.get(
                path + '?salesperson=Bob&date_from=2024-05-01&date_to=2024-05-31'
            ).get_data(as_text=True)
            self.assertNotIn('PI-TEST-REPORT-FILTER', excluded)
            self.assertIn('当前结果：<strong>0</strong> 份 PI', excluded)

        procurement_html = self.client.get(
            '/procurement' + query_string
        ).get_data(as_text=True)
        self.assertIn('procurement-list-table', procurement_html)
        self.assertIn('proc-products-col', procurement_html)
        self.assertIn('width: 320px', procurement_html)
        self.assertIn('white-space: nowrap', procurement_html)
        self.assertIn('text-overflow: ellipsis', procurement_html)

        with application.app.app_context():
            db.session.delete(db.session.get(PI, filtered_pi_id))
            db.session.commit()

    def test_profit_report_calculates_complete_per_order_structure_in_rmb(self):
        with application.app.app_context():
            original_settings = application._load_settings().copy()
            settings = original_settings.copy()
            # The editable default is deliberately different from the PI's
            # saved rate: historical profit must continue to use the PI rate.
            settings['exchange_rate'] = '9'
            application._save_settings(settings)

            Procurement.query.filter_by(pi_id=self.alice_pi).delete()
            Payment.query.filter_by(pi_id=self.alice_pi).delete()
            Expense.query.filter_by(pi_id=self.alice_pi).delete()
            pi = db.session.get(PI, self.alice_pi)
            pi.currency = 'USD'
            pi.exchange_rate = 7
            pi.total_amount = 10
            pi.shipping_cost = 2
            pi.actual_shipping_cost = 3
            pi.procurement_confirmed = True
            item = PIItem.query.filter_by(pi_id=self.alice_pi).one()
            supplier = Supplier(name='Profit Test Supplier')
            db.session.add(supplier)
            db.session.flush()
            db.session.add_all([
                Procurement(
                    pi_id=pi.id, pi_item_id=item.id, supplier_id=supplier.id,
                    unit_price=60, quantity=1, total=60,
                ),
                Payment(
                    pi_id=pi.id, amount=10, fee=1,
                    order_no='PROFIT-TEST-001', order_no_normalized='profit-test-001',
                    idempotency_key='profit-test-idempotency',
                    receiving_account_id=self.approved_account,
                    receiving_account_name='Approved', receiving_account_currency='USD',
                ),
                Expense(pi_id=pi.id, category='报关费', amount=7, currency='RMB'),
            ])
            application._recalculate_pi_payments(pi)
            db.session.commit()
            supplier_id = supplier.id

        self.login('admin-test')
        response = self.client.get('/profit-report')
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        # Order total: (USD 10 + USD 2) × 7 = RMB 84.
        # Net income: 84 - payment fee 7 = RMB 77.
        # Cost: procurement 60 + freight 3 + expense 7 = RMB 70.
        # Profit: RMB 7.
        self.assertIn('¥84.00', html)
        self.assertIn('¥60.00', html)
        self.assertIn('¥77.00', html)
        self.assertIn('¥70.00', html)
        self.assertIn('¥7.00', html)
        self.assertNotIn('正式利润', html)
        self.assertIn('收款手续费', html)
        self.assertIn('供应商采购运费', html)
        self.assertIn('业务员登记的真实成本', html)
        self.assertIn('采购产品成本', html)
        self.assertIn('订单真实成本合计', html)
        self.assertIn('订单总金额', html)
        self.assertIn('$12.00', html)
        self.assertIn('净收入', html)
        self.assertIn('净利润', html)
        self.assertIn('1 笔登记成本', html)
        self.assertIn('查看明细', html)
        self.assertNotIn('成本记录（1）', html)
        self.assertIn(f'id="profitDetail{self.alice_pi}"', html)
        self.assertNotIn('<th class="text-end">采购成本</th>', html)
        self.assertIn('本单汇率：1 USD = ¥7.0000', html)
        self.assertIn('新单默认汇率：1 USD = ¥9.0000', html)

        with application.app.app_context():
            Procurement.query.filter_by(pi_id=self.alice_pi).delete()
            Payment.query.filter_by(pi_id=self.alice_pi).delete()
            Expense.query.filter_by(pi_id=self.alice_pi).delete()
            pi = db.session.get(PI, self.alice_pi)
            pi.currency = 'USD'
            pi.total_amount = 10
            pi.shipping_cost = 0
            pi.actual_shipping_cost = 0
            pi.procurement_confirmed = False
            pi.shipping_completed = False
            application._recalculate_pi_payments(pi)
            Supplier.query.filter_by(id=supplier_id).delete()
            db.session.commit()
            application._save_settings(original_settings)

    def test_sqlite_safety_pragmas(self):
        with application.app.app_context():
            self.assertEqual(db.session.execute(db.text('PRAGMA journal_mode')).scalar(), 'wal')
            self.assertEqual(db.session.execute(db.text('PRAGMA foreign_keys')).scalar(), 1)

    def test_disabled_user_existing_session_is_rejected(self):
        self.login()
        with application.app.app_context():
            user = User.query.filter_by(username='alice').one()
            user.active = False
            user.auth_version += 1
            db.session.commit()
        response = self.client.get('/customers')
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.endswith('/login'))
        with application.app.app_context():
            user = User.query.filter_by(username='alice').one()
            user.active = True
            user.auth_version += 1
            db.session.commit()

    def test_authenticated_responses_are_not_cached(self):
        self.login()
        response = self.client.get('/customers')
        self.assertIn('no-store', response.headers.get('Cache-Control', ''))

    def test_authenticated_static_assets_are_cacheable(self):
        self.login()
        response = self.client.get('/static/style.css')
        self.assertEqual(response.status_code, 200)
        cache_control = response.headers.get('Cache-Control', '')
        self.assertIn('public', cache_control)
        self.assertIn('max-age=3600', cache_control)
        self.assertNotIn('no-store', cache_control)

    def test_product_thumbnail_is_small_and_privately_cacheable(self):
        from PIL import Image

        filename = 'thumbnail-test.png'
        upload_dir = application.app.config['UPLOAD_DIR']
        source_path = os.path.join(upload_dir, filename)
        Image.new('RGB', (1200, 800), '#225588').save(source_path, format='PNG')
        with application.app.app_context():
            product = Product.query.filter_by(product_code='ORIG').one()
            previous_image = product.image
            product.image = filename
            db.session.commit()
        try:
            self.login()
            response = self.client.get(f'/uploads/thumb/{filename}')
            self.assertEqual(response.status_code, 200)
            self.assertIn('private', response.headers.get('Cache-Control', ''))
            self.assertIn('max-age=604800', response.headers.get('Cache-Control', ''))
            with Image.open(BytesIO(response.data)) as thumbnail:
                self.assertLessEqual(thumbnail.width, 320)
                self.assertLessEqual(thumbnail.height, 320)
        finally:
            with application.app.app_context():
                product = Product.query.filter_by(product_code='ORIG').one()
                product.image = previous_image
                db.session.commit()
            if os.path.exists(source_path):
                os.remove(source_path)
            thumb_path = os.path.join(upload_dir, '.thumbs', 'thumbnail-test.webp')
            if os.path.exists(thumb_path):
                os.remove(thumb_path)

    def test_salesperson_live_edit_cannot_change_protected_fields(self):
        with application.app.test_request_context('/pi/test'):
            user = User.query.filter_by(username='alice').one()
            application.session['user_id'] = user.id
            application.session['auth_version'] = user.auth_version
            pi = db.session.get(PI, self.alice_pi)
            item = pi.items[0]
            form = MultiDict({
                'pi_number': 'FORGED-001',
                'salesperson': 'Bob',
                'bank_info': 'ATTACKER BANK',
                'currency': 'USD',
                'issue_date': pi.issue_date.strftime('%Y-%m-%d'),
                'shipping_cost': '0',
                'cust_name': 'Forged Customer',
                'cust_country': 'Forged Country',
                f'qty_{item.id}': '2',
                f'price_{item.id}': '12.50',
                f'prod_name_{item.id}': 'Forged Product',
                f'prod_code_{item.id}': 'FORGED',
                f'prod_spec_{item.id}': 'Forged spec',
            })
            application._apply_form_to_pi(form, pi)
            self.assertEqual(pi.pi_number, 'PI-TEST-001')
            self.assertEqual(pi.salesperson, 'Alice')
            self.assertNotEqual(pi.bank_info, 'ATTACKER BANK')
            self.assertEqual(pi.customer.name, 'Alice Customer')
            self.assertEqual(item.product.name, 'Original Product')
            self.assertEqual(item.quantity, 2)
            db.session.rollback()

    def test_payment_history_does_not_build_user_input_as_html(self):
        template_path = os.path.join(os.path.dirname(application.__file__), 'templates', 'pi_list.html')
        with open(template_path, encoding='utf-8') as handle:
            source = handle.read()
        self.assertNotIn("histHtml +=", source)
        self.assertIn("cell.textContent = value", source)

    def test_order_total_includes_charges_and_discounts(self):
        self.login()
        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            pi.total_amount = 100
            pi.shipping_cost = -10
            db.session.commit()
            self.assertEqual(pi.product_subtotal, 100)
            self.assertEqual(pi.other_charges, -10)
            self.assertEqual(pi.grand_total, 90)
        dashboard = self.client.get('/').get_data(as_text=True)
        self.assertIn('$90.00', dashboard)
        self.assertIn('lang="zh-CN"', dashboard)
        self.assertIn('<meta name="google" content="notranslate">', dashboard)
        self.assertIn('<span translate="no" class="notranslate">$90.00</span>', dashboard)
        preview = self.client.get(f'/pi/{self.alice_pi}/preview').get_data(as_text=True)
        self.assertIn('$100.00', preview)
        self.assertIn('$-10.00', preview)
        self.assertIn('$90.00', preview)
        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            pi.received_amount = 40
            pi.paid = False
            db.session.commit()
        stats = self.client.get('/sales-stats').get_data(as_text=True)
        self.assertIn('订单总金额', stats)
        self.assertIn('实际回款金额', stats)
        self.assertIn('$90.00', stats)
        self.assertIn('$40.00', stats)
        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            pi.total_amount = 10
            pi.shipping_cost = 0
            pi.received_amount = 0
            pi.paid = False
            db.session.commit()

    def test_pi_exports_label_adjustments_as_other_charges(self):
        workbook = load_workbook(
            application.SYSTEM_DEFAULT_TEMPLATE_SOURCE,
            data_only=False,
            read_only=True,
        )
        try:
            self.assertEqual(
                workbook['PI Template']['E19'].value,
                'Other Charges ({{currency}}):',
            )
        finally:
            workbook.close()

        self.login()
        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            previous = (
                pi.shipping_cost, pi.shipping_note, pi.shipping_note_en,
            )
            pi.shipping_cost = 5
            pi.shipping_note = ''
            pi.shipping_note_en = ''
            db.session.commit()
        try:
            preview = self.client.get(
                f'/pi/{self.alice_pi}/preview'
            ).get_data(as_text=True)
            self.assertIn('Other Charges (USD):', preview)
            self.assertNotIn('Other Charges / Discount', preview)
        finally:
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                (
                    pi.shipping_cost, pi.shipping_note, pi.shipping_note_en,
                ) = previous
                db.session.commit()

    def test_migration_normalizes_legacy_empty_deleted_dates(self):
        with application.app.app_context():
            db.session.execute(db.text(
                "UPDATE customers SET deleted_at = '' WHERE id = :customer_id"
            ), {'customer_id': self.alice_customer})
            db.session.commit()
            application._migrate_db()
            raw_value = db.session.execute(db.text(
                "SELECT deleted_at FROM customers WHERE id = :customer_id"
            ), {'customer_id': self.alice_customer}).scalar()
            self.assertIsNone(raw_value)

    def test_customer_delete_is_recoverable_and_audited(self):
        self.login()
        customer_page = self.client.get('/customers').get_data(as_text=True)
        self.assertIn(
            "onsubmit='return typedConfirm(this, \"Alice Customer\",",
            customer_page,
        )
        with application.app.app_context():
            customer = Customer(name='Recoverable Customer', salesperson='Alice')
            db.session.add(customer)
            db.session.commit()
            customer_id = customer.id
        response = self.client.post(f'/customers/{customer_id}/delete', data={
            'confirm_value': 'Recoverable Customer',
            'csrf_token': self.token('/customers'),
        })
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            self.assertIsNotNone(db.session.get(Customer, customer_id).deleted_at)
            self.assertIsNotNone(AuditLog.query.filter_by(entity_type='customer', entity_id=customer_id, action='soft_delete').first())
        self.client.get('/logout')
        self.login('admin-test')
        self.client.post(f'/recycle-bin/customer/{customer_id}/restore', data={
            'csrf_token': self.token('/recycle-bin'),
        })
        with application.app.app_context():
            self.assertIsNone(db.session.get(Customer, customer_id).deleted_at)

    def test_stale_customer_form_cannot_overwrite_newer_data(self):
        self.login()
        with application.app.app_context():
            customer = Customer(name='Concurrent Customer', salesperson='Alice')
            db.session.add(customer)
            db.session.commit()
            customer_id = customer.id
            stale_version = customer.version
            customer.name = 'Newer Server Value'
            db.session.commit()
        response = self.client.post(f'/customers/{customer_id}/edit', data={
            'name': 'Stale Browser Value', 'salesperson': 'Alice',
            'version': stale_version, 'csrf_token': self.token(f'/customers/{customer_id}/edit'),
        })
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            self.assertEqual(db.session.get(Customer, customer_id).name, 'Newer Server Value')

    def test_payment_reference_is_globally_unique_and_idempotent(self):
        self.login()
        with application.app.app_context():
            audit_count_before = AuditLog.query.filter_by(entity_type='payment', action='create').count()
        first_token = 'a' * 32
        response = self.client.post(f'/pi/{self.alice_pi}/toggle-paid', data={
            'received_amount': '2', 'fee': '0', 'order_no': ' BANK-REF-001 ',
            'receiving_account_id': self.approved_account,
            'attachment': self.image_upload('bank-ref.png'),
            'idempotency_key': first_token, 'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(response.status_code, 302)
        response = self.client.post(f'/pi/{self.alice_pi}/toggle-paid', data={
            'received_amount': '2', 'fee': '0', 'order_no': 'bank-ref-001',
            'receiving_account_id': self.approved_account,
            'idempotency_key': 'b' * 32, 'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            self.assertEqual(Payment.query.filter_by(order_no_normalized='bank-ref-001').count(), 1)
            self.assertEqual(
                AuditLog.query.filter_by(entity_type='payment', action='create').count(),
                audit_count_before + 1,
            )

    def test_payment_requires_matching_account_and_allows_same_currency_switch(self):
        self.login()
        with application.app.app_context():
            usd_alternate = Account(name='USD Alternate', currency='USD')
            rmb_account = Account(name='RMB Account', currency='RMB')
            pi = PI(
                pi_number='PI-PAYMENT-ACCOUNT-001',
                customer_id=self.alice_customer,
                salesperson='Alice',
                bank_info='SAFE BANK\nA/C: 123\nSWIFT: SAFE',
                currency='USD',
                total_amount=10,
            )
            db.session.add_all([usd_alternate, rmb_account, pi])
            db.session.commit()
            pi_id = pi.id
            usd_alternate_id = usd_alternate.id
            rmb_account_id = rmb_account.id

        missing = self.client.post(f'/pi/{pi_id}/toggle-paid', data={
            'received_amount': '1', 'fee': '0', 'order_no': 'NO-ACCOUNT',
            'idempotency_key': 'd' * 32, 'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(missing.status_code, 302)
        wrong_currency = self.client.post(f'/pi/{pi_id}/toggle-paid', data={
            'received_amount': '1', 'fee': '0', 'order_no': 'WRONG-CURRENCY',
            'receiving_account_id': rmb_account_id,
            'idempotency_key': 'e' * 32, 'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(wrong_currency.status_code, 302)
        first = self.client.post(f'/pi/{pi_id}/toggle-paid', data={
            'received_amount': '4', 'fee': '0', 'order_no': 'USD-PRIMARY',
            'receiving_account_id': self.approved_account,
            'attachment': self.image_upload('primary.png'),
            'idempotency_key': 'f' * 32, 'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(first.status_code, 302)
        switched = self.client.post(f'/pi/{pi_id}/toggle-paid', data={
            'received_amount': '1', 'fee': '0', 'order_no': 'USD-ALTERNATE',
            'receiving_account_id': usd_alternate_id,
            'attachment': self.image_upload('alternate.png'),
            'idempotency_key': 'g' * 32, 'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(switched.status_code, 302)

        with application.app.app_context():
            payments = Payment.query.filter_by(pi_id=pi_id).order_by(Payment.id).all()
            self.assertEqual(len(payments), 2)
            self.assertEqual(payments[0].receiving_account_name, 'Approved')
            self.assertEqual(payments[1].receiving_account_name, 'USD Alternate')
            self.assertEqual(payments[1].receiving_account_currency, 'USD')

        payment_data = self.client.get(f'/api/pi/{pi_id}/payments').get_json()
        self.assertEqual(payment_data['default_account_id'], usd_alternate_id)
        self.assertEqual(
            [payment['receiving_account_name'] for payment in payment_data['payments']],
            ['Approved', 'USD Alternate'],
        )
        list_html = self.client.get('/pi/list').get_data(as_text=True)
        self.assertIn('name="receiving_account_id"', list_html)
        self.assertIn('name="attachment" id="payAttachment" required', list_html)
        self.assertIn('onsubmit="return preparePaymentSubmit(this);"', list_html)
        self.assertNotIn("Processing...';return true;\">", list_html)

    def test_payment_receipt_is_required_validated_and_access_controlled(self):
        self.login()
        with application.app.app_context():
            pi = PI(
                pi_number='PI-PAYMENT-RECEIPT-001',
                customer_id=self.alice_customer,
                salesperson='Alice',
                currency='USD',
                total_amount=10,
            )
            db.session.add(pi)
            db.session.commit()
            pi_id = pi.id

        missing = self.client.post(f'/pi/{pi_id}/toggle-paid', data={
            'received_amount': '1', 'fee': '0', 'order_no': 'RECEIPT-MISSING',
            'receiving_account_id': self.approved_account,
            'idempotency_key': 'h' * 32, 'csrf_token': self.token('/pi/list'),
        }, follow_redirects=True)
        self.assertIn('请上传回款凭证图片', missing.get_data(as_text=True))

        invalid = self.client.post(f'/pi/{pi_id}/toggle-paid', data={
            'received_amount': '1', 'fee': '0', 'order_no': 'RECEIPT-INVALID',
            'receiving_account_id': self.approved_account,
            'attachment': (BytesIO(b'not-an-image'), 'fake.png'),
            'idempotency_key': 'i' * 32, 'csrf_token': self.token('/pi/list'),
        }, follow_redirects=True)
        self.assertIn('必须是有效的 JPG', invalid.get_data(as_text=True))

        saved = self.client.post(f'/pi/{pi_id}/toggle-paid', data={
            'received_amount': '1', 'fee': '0', 'order_no': 'RECEIPT-VALID',
            'receiving_account_id': self.approved_account,
            'attachment': self.image_upload('receipt.png'),
            'idempotency_key': 'j' * 32, 'csrf_token': self.token('/pi/list'),
        })
        self.assertEqual(saved.status_code, 302)
        with application.app.app_context():
            payment = Payment.query.filter_by(order_no='RECEIPT-VALID').one()
            payment_id = payment.id
            receipt_path = os.path.join(application.app.config['UPLOAD_DIR'], payment.attachment)
            self.assertTrue(os.path.isfile(receipt_path))

        payment_data = self.client.get(f'/api/pi/{pi_id}/payments').get_json()
        self.assertTrue(payment_data['payments'][0]['has_attachment'])
        receipt_response = self.client.get(f'/payments/{payment_id}/attachment')
        self.assertEqual(receipt_response.status_code, 200)
        self.assertEqual(receipt_response.mimetype, 'image/png')
        self.assertIn('no-store', receipt_response.headers.get('Cache-Control', ''))

        self.client.get('/logout')
        self.login('bob')
        self.assertEqual(self.client.get(f'/payments/{payment_id}/attachment').status_code, 403)

    def test_payment_delete_requires_typed_reference(self):
        self.login()
        with application.app.app_context():
            payment = Payment(pi_id=self.alice_pi, amount=1, order_no='DELETE-ME',
                              order_no_normalized='delete-me', idempotency_key='c' * 32)
            db.session.add(payment)
            pi = db.session.get(PI, self.alice_pi)
            pi.procurement_confirmed = True
            pi.procurement_status = '发货完成'
            pi.shipping_completed = True
            pi.shipping_date = date(2026, 9, 9)
            pi.shipping_tracking_no = 'KEEP-TRACKING'
            pi.shipping_record_note = 'KEEP-SHIPPING-NOTE'
            pi.shipping_recorded_by = 'alice'
            pi.customs_required = True
            pi.customs_note = 'KEEP-CUSTOMS-NOTE'
            db.session.commit()
            payment_id = payment.id
        self.client.post(f'/pi/{self.alice_pi}/payment/{payment_id}/delete', data={
            'confirm_value': 'wrong', 'csrf_token': self.token('/pi/list'),
        })
        with application.app.app_context():
            self.assertIsNone(db.session.get(Payment, payment_id).deleted_at)
        self.client.post(f'/pi/{self.alice_pi}/payment/{payment_id}/delete', data={
            'confirm_value': 'DELETE-ME', 'csrf_token': self.token('/pi/list'),
        })
        with application.app.app_context():
            self.assertIsNotNone(db.session.get(Payment, payment_id).deleted_at)
            application._migrate_db()
            pi = db.session.get(PI, self.alice_pi)
            self.assertEqual(pi.received_amount, 0)
            self.assertFalse(pi.paid)
            self.assertTrue(pi.procurement_confirmed)
            self.assertTrue(pi.shipping_completed)
            self.assertEqual(pi.shipping_date, date(2026, 9, 9))
            self.assertEqual(pi.shipping_tracking_no, 'KEEP-TRACKING')
            self.assertEqual(pi.shipping_record_note, 'KEEP-SHIPPING-NOTE')
            self.assertEqual(pi.shipping_recorded_by, 'alice')
            self.assertIs(pi.customs_required, True)
            self.assertEqual(pi.customs_note, 'KEEP-CUSTOMS-NOTE')
            self.assertEqual(pi.effective_procurement_status, '发货完成')

    def test_pi_create_requires_valid_account_for_admin_and_sales_but_not_drafts(self):
        for user in ['admin-test', 'alice']:
            self.login(user)
            token = self.token('/pi/create')
            with application.app.app_context():
                count = PI.query.count()
                product_id = Product.query.first().id
            data = {'csrf_token': token, 'customer_id': self.alice_customer,
                    'salesperson': 'Alice', 'notes': 'account required',
                    'currency': 'USD', 'exchange_rate': '7',
                    'selected_' + str(product_id): 'on', 'qty_' + str(product_id): '1',
                    'unit_price_' + str(product_id): '1.235'}
            for account_id in ['', 'invalid', '999999999']:
                with patch.object(application, '_generate_default_pi_documents') as renderer:
                    response = self.client.post('/pi/create', data=dict(data, account_id=account_id))
                    self.assertEqual(response.status_code, 400)
                    renderer.assert_not_called()
                with application.app.app_context():
                    self.assertEqual(PI.query.count(), count)
            without_currency = dict(data, account_id=str(self.approved_account))
            without_currency.pop('currency')
            with patch.object(application, '_generate_default_pi_documents') as renderer:
                response = self.client.post('/pi/create', data=without_currency)
                self.assertEqual(response.status_code, 400)
                self.assertIn('请选择币种', response.get_data(as_text=True))
                renderer.assert_not_called()
            with application.app.app_context():
                self.assertEqual(PI.query.count(), count)
            response = self.client.post('/pi/drafts/save', data={'csrf_token': token, 'draft_rows': '[]'})
            self.assertEqual(response.status_code, 200)
            html = self.client.get('/pi/create').get_data(as_text=True)
            self.assertIn('id="account_select" required', html)
            self.assertIn('id="currency_select" required', html)
            self.assertIn('— 请选择币种 —', html)
            self.assertIn("alert('请选择币种。')", html)

    def test_export_keeps_saved_account_despite_browser_newlines(self):
        self.login('admin-test')
        with application.app.app_context():
            account = db.session.get(Account, self.approved_account)
            pi = PI(pi_number='BANK-NEWLINE-TEST', customer_id=self.alice_customer,
                    salesperson='Alice', company='klista', currency='USD', notes='bank transport')
            application._apply_bank_snapshot(pi, account)
            db.session.add(pi)
            db.session.flush()
            db.session.add(PIItem(pi_id=pi.id, product_id=Product.query.first().id,
                                  quantity=1, unit_price=1.235, amount=1.24))
            db.session.commit()
            pi_id, original, version = pi.id, pi.bank_info, pi.version
        url = '/pi/%s/export' % pi_id
        html = self.client.get(url).get_data(as_text=True)
        self.assertIn('原账户与抬头匹配时可直接导出', html)
        for newline in ['\r\n', '\r', '\n']:
            submitted = original.replace('\n', newline)
            response = self.client.post(url, data={
                'csrf_token': self.token(url), 'company_header': 'klista',
                'account_id': '', 'bank_info': submitted, 'shipping_cost': '0',
                'mode': 'download', 'output_format': 'xlsx',
            })
            self.assertEqual(response.status_code, 200)
            response.close()
        with application.app.app_context():
            pi = db.session.get(PI, pi_id)
            self.assertEqual((pi.bank_info, pi.bank_receiving_account_id, pi.version),
                             (original, self.approved_account, version))
            copied = application._pi_export_copy(pi)
            copied.bank_info = original.replace('\n', '\r\n')
            application._validate_export_account_brand(copied)
            copied.bank_receiving_account_id = None
            self.assertEqual(application._matching_account_for_pi(copied).id, self.approved_account)
        with patch.object(application, '_render_pi_export') as renderer:
            response = self.client.post(url, data={
                'csrf_token': self.token(url), 'account_id': '',
                'bank_info': original.replace('123', 'changed-account'),
                'shipping_cost': '0', 'output_format': 'xlsx', 'mode': 'download',
            })
            self.assertEqual(response.status_code, 400)
            renderer.assert_not_called()

    def test_account_brand_multiselect_validation_and_legacy_preservation(self):
        self.login('admin-test')
        with application.app.app_context():
            a = Account(name='Brand Test', brand='qisuo', currency='USD')
            db.session.add(a)
            db.session.commit()
            account_id = a.id
        url = f'/accounts/{account_id}/edit'
        data = {'name': 'Brand Test', 'currency': 'USD',
                'current_password': 'AdminPass123!', 'brand': ['klista', 'qisuo']}
        data['csrf_token'] = self.token(url)
        self.assertEqual(self.client.post(url, data=data).status_code, 302)
        with application.app.app_context():
            a = db.session.get(Account, account_id)
            self.assertEqual(a.brands, ['klista', 'qisuo'])
            self.assertEqual(a.brand, 'klista,qisuo')
        for values in ([], ['klista', 'invalid']):
            data['brand'] = values
            data['name'] = 'Must not change'
            data['csrf_token'] = self.token(url)
            self.assertEqual(self.client.post(url, data=data).status_code, 400)
            with application.app.app_context():
                a = db.session.get(Account, account_id)
                self.assertEqual(a.name, 'Brand Test')
                self.assertEqual(a.brands, ['klista', 'qisuo'])
        self.assertIn('type="checkbox" name="brand"', self.client.get(url).get_data(as_text=True))

    def test_receiving_account_association_and_ambiguous_legacy_snapshot(self):
        with application.app.test_request_context('/pi/test'):
            pi = db.session.get(PI, self.alice_pi)
            copy = application._pi_export_copy(pi)
            a = db.session.get(Account, self.approved_account)
            application._apply_bank_snapshot(copy, a)
            self.assertEqual(copy.bank_receiving_account_id, a.id)
            self.assertEqual(application._matching_account_for_pi(copy).id, a.id)
            application._validate_export_account_brand(copy)
            workbook = Workbook()
            workbook.active['A1'] = 'CHANGZHOU QISUO CO., LTD'
            source = os.path.join(TEST_ROOT.name, 'fixed-company-test.xlsx')
            workbook.save(source)
            workbook.close()
            with patch.object(application, '_template_file_path', return_value=source):
                with self.assertRaises(ValueError):
                    application._validate_export_account_brand(copy, SimpleNamespace(template_type='xlsx'))
            clone = Account(name='Ambiguous account', brand='qisuo',
                            bank_name=a.bank_name, account_no=a.account_no,
                            swift_code=a.swift_code, currency=a.currency)
            db.session.add(clone)
            db.session.flush()
            try:
                copy.bank_receiving_account_id = None
                self.assertIsNone(application._matching_account_for_pi(copy))
                with self.assertRaises(ValueError):
                    application._validate_export_account_brand(copy)
                copy.bank_receiving_account_id = a.id
                application._validate_export_account_brand(copy)
                user = User.query.filter_by(username='admin-test').one()
                application.session['user_id'] = user.id
                application.session['auth_version'] = user.auth_version
                with self.assertRaises(ValueError):
                    application._apply_export_form(MultiDict({
                        'bank_info': 'FAKE RECIPIENT', 'shipping_cost': '0',
                    }), copy)
                saved_bank_info = copy.bank_info
                copy.bank_info = 'FAKE BANK TEXT'
                with self.assertRaises(ValueError):
                    application._validate_export_account_brand(copy)
                copy.bank_info = saved_bank_info
                copy._company_name_override = 'FAKE COMPANY'
                with self.assertRaises(ValueError):
                    application._validate_export_account_brand(copy)
            finally:
                db.session.rollback()

    def test_export_brand_mismatch_blocks_preview_download_and_direct_routes(self):
        self.login('admin-test')
        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            old = (pi.company, pi.bank_info, pi.bank_receiving_account_id)
            a = Account(name='Export Brand Guard', bank_name='GUARD BANK',
                        account_no='brand-guard-123', brand='klista')
            db.session.add(a)
            db.session.flush()
            account_id = a.id
            pi.company = 'qisuo'
            pi.bank_info = application._legacy_account_bank_info(a)
            # Legacy PI deliberately has no account ID: uniquely match its snapshot.
            pi.bank_receiving_account_id = None
            db.session.commit()
            template_id = DocumentTemplate.query.filter_by(code='system-default').one().id
        url = f'/pi/{self.alice_pi}/export'
        try:
            for mode, fmt in [('preview', 'pdf'), ('download', 'pdf'), ('download', 'xlsx')]:
                data = {'company_header': 'qisuo', 'template_id': template_id,
                        'mode': mode, 'output_format': fmt, 'shipping_cost': '0',
                        'csrf_token': self.token(url)}
                with patch.object(application, '_render_pi_export') as renderer:
                    response = self.client.post(url, data=data)
                    self.assertEqual(response.status_code, 400)
                    self.assertIn('不一致', response.get_json()['error'])
                    renderer.assert_not_called()
            with patch.object(application, '_render_pi_export') as renderer:
                self.assertEqual(self.client.get(f'/pi/{self.alice_pi}/download').status_code, 302)
                renderer.assert_not_called()
            self.assertEqual(self.client.get(f'/pi/{self.alice_pi}/excel').status_code, 302)
            # Mismatch fixed by selecting a matching company; real xlsx generation succeeds.
            data.update(company_header='klista', mode='download', output_format='xlsx')
            response = self.client.post(url, data=data)
            self.assertEqual(response.status_code, 200)
            response.close()
            with application.app.app_context():
                a = db.session.get(Account, account_id)
                a.brand = 'klista,qisuo'
                db.session.commit()
            data['company_header'] = 'qisuo'
            response = self.client.post(url, data=data)
            self.assertEqual(response.status_code, 200)
            response.close()
            with application.app.app_context():
                a = db.session.get(Account, account_id)
                a.brand = ''
                db.session.commit()
            self.assertEqual(self.client.post(url, data=data).status_code, 400)
            data['account_id'] = '99999999'
            self.assertEqual(self.client.post(url, data=data).status_code, 400)
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                self.assertEqual(pi.company, 'qisuo')
                self.assertEqual(pi.bank_receiving_account_id, None)
                self.assertEqual(pi.bank_info, 'GUARD BANK\nA/C: brand-guard-123')
        finally:
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                pi.company, pi.bank_info, pi.bank_receiving_account_id = old
                db.session.commit()

    def test_bank_account_change_requires_current_admin_password(self):
        self.login('admin-test')
        token = self.token('/accounts/add')
        denied = self.client.post('/accounts/add', data={
            'name': 'Protected Account', 'current_password': 'wrong', 'csrf_token': token,
        })
        self.assertEqual(denied.status_code, 403)
        with application.app.app_context():
            self.assertIsNone(Account.query.filter_by(name='Protected Account').first())
        allowed = self.client.post('/accounts/add', data={
            'name': 'Protected Account', 'brand': 'klista', 'current_password': 'AdminPass123!',
            'csrf_token': self.token('/accounts/add'),
        })
        self.assertEqual(allowed.status_code, 302)
        with application.app.app_context():
            account = Account.query.filter_by(name='Protected Account').one()
            self.assertIsNotNone(AuditLog.query.filter_by(entity_type='account', entity_id=account.id, action='create').first())

    def test_field_management_controls_pi_note_choices(self):
        self.login('admin-test')
        page = self.client.get('/field-management')
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn('字段管理', html)
        self.assertIn('客户费用 / 折扣类型（对外）', html)
        self.assertIn('订单真实成本类别（内部）', html)
        self.assertIn('英文名称（对外单据）', html)
        self.assertIn('收款账户', html)
        self.assertIn('<th>币种</th>', html)
        self.assertIn('账户币种', self.client.get('/accounts/add').get_data(as_text=True))

        response = self.client.post('/field-options/add', data={
            'field_key': 'shipping_note',
            'value': '测试附加费',
            'english_value': 'Test Surcharge',
            'csrf_token': self.token('/field-management'),
        })
        self.assertEqual(response.status_code, 302)
        with application.app.app_context():
            option = FieldOption.query.filter_by(
                field_key='shipping_note', value='测试附加费'
            ).one()
            self.assertEqual(option.english_value, 'Test Surcharge')
            self.assertIsNotNone(AuditLog.query.filter_by(
                entity_type='field_option', entity_id=option.id, action='create'
            ).first())

        create_html = self.client.get('/pi/create').get_data(as_text=True)
        self.assertIn('name="shipping_note" id="shipping_note"', create_html)
        self.assertIn('测试附加费', create_html)
        self.assertIn('Test Surcharge', create_html)
        with application.app.app_context():
            self.assertEqual(
                application._managed_field_value('shipping_note', '测试附加费'),
                '测试附加费',
            )
            self.assertEqual(
                application._managed_field_english('shipping_note', '测试附加费'),
                'Test Surcharge',
            )
            pi = db.session.get(PI, self.alice_pi)
            old_note, old_note_en, old_cost = pi.shipping_note, pi.shipping_note_en, pi.shipping_cost
            pi.shipping_note = '测试附加费'
            pi.shipping_note_en = 'Test Surcharge'
            pi.shipping_cost = 2
            self.assertEqual(_fixed_values(pi, {})['shipping_note'], 'Test Surcharge')
            self.assertEqual(_fixed_values(pi, {})['shipping_note_zh'], '测试附加费')
            pi.shipping_note, pi.shipping_note_en, pi.shipping_cost = old_note, old_note_en, old_cost
            with self.assertRaises(ValueError):
                application._managed_field_value('shipping_note', '未配置备注')

        cost_response = self.client.post('/field-options/add', data={
            'field_key': 'expense_category',
            'value': '测试真实成本',
            'csrf_token': self.token('/field-management'),
        })
        self.assertEqual(cost_response.status_code, 302)
        self.assertIn('测试真实成本', self.client.get('/pi/list').get_data(as_text=True))

        self.client.get('/logout')
        self.login('alice')
        self.assertEqual(self.client.get('/field-management').status_code, 403)
        denied = self.client.post('/field-options/add', data={
            'field_key': 'shipping_note',
            'value': '越权选项',
            'english_value': 'Unauthorized Option',
            'csrf_token': self.token('/pi/create'),
        })
        self.assertEqual(denied.status_code, 403)
        with application.app.app_context():
            self.assertIsNone(FieldOption.query.filter_by(value='越权选项').first())

    def test_export_workbench_uses_copy_and_keeps_original_pi_unchanged(self):
        self.login('alice')
        page = self.client.get(f'/pi/{self.alice_pi}/export')
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn('PI 导出工作台', html)
        self.assertIn('原始 PI 不会被修改', html)
        self.assertNotIn('保存并重新生成 PDF', html)
        self.assertIn('系统默认 PI 模板', html)
        list_html = self.client.get('/pi/list').get_data(as_text=True)
        self.assertIn('> 预览\n', list_html)
        self.assertIn('bi-box-arrow-up-right text-primary"></i>导出', list_html)
        self.assertNotIn('在线编辑', list_html)
        self.assertIn('class="pi-list-workflow"', list_html)
        self.assertNotIn('>业务处理</button>', list_html)
        self.assertIn('更多', list_html)
        self.assertNotIn('<th class="text-center">装箱单</th>', list_html)
        self.assertNotIn('<th class="text-center">报关</th>', list_html)

        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            item = pi.items[0]
            template_id = DocumentTemplate.query.filter_by(code='system-default').one().id
            original = {
                'company': pi.company,
                'bank_info': pi.bank_info,
                'issue_date': pi.issue_date,
                'payment_terms': pi.payment_terms,
                'shipping_cost': pi.shipping_cost,
                'total_amount': pi.total_amount,
                'contact': pi.customer.contact_person,
                'quantity': item.quantity,
                'price': item.unit_price,
            }
            item_id = item.id

        response = self.client.post(f'/pi/{self.alice_pi}/export', data={
            'company_header': 'qisuo',
            'company_name': 'FORGED COMPANY',
            'bank_info': 'FORGED BANK',
            'template_id': template_id,
            'output_format': 'xlsx',
            'mode': 'download',
            'pi_number': 'FORGED-EXPORT-PI',
            'salesperson': 'Bob',
            'currency': 'RMB',
            'issue_date': '2026-12-31',
            'payment_terms': '仅用于本次导出',
            'shipping_address': 'Temporary export address',
            'shipping_note': '运费',
            'shipping_cost': '5.50',
            'notes': 'Temporary export note',
            'cust_contact': 'Temporary Contact',
            'cust_email': 'temporary@example.com',
            'cust_phone': '+86 10000',
            'cust_address': 'Temporary customer address',
            f'qty_{item_id}': '2',
            f'price_{item_id}': '12.50',
            'csrf_token': self.token(f'/pi/{self.alice_pi}/export'),
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        payload = response.data
        response.close()
        workbook = load_workbook(BytesIO(payload), data_only=False)
        rendered_values = [
            cell.value for row in workbook.active.iter_rows() for cell in row
            if cell.value is not None
        ]
        self.assertTrue(any('仅用于本次导出' in str(value) for value in rendered_values))
        self.assertTrue(any('Freight' in str(value) for value in rendered_values))
        self.assertTrue(any('Changzhou Qisuo' in str(value) for value in rendered_values))
        self.assertTrue(any('Jingchuang Road' in str(value) for value in rendered_values))
        self.assertFalse(any('FORGED COMPANY' in str(value) or 'FORGED BANK' in str(value) for value in rendered_values))
        self.assertIn(25, rendered_values)
        self.assertTrue(any('PI-TEST-001' in str(value) for value in rendered_values))
        self.assertFalse(any('FORGED-EXPORT-PI' in str(value) for value in rendered_values))
        workbook.close()

        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            item = db.session.get(PIItem, item_id)
            self.assertEqual(pi.company, original['company'])
            self.assertEqual(pi.bank_info, original['bank_info'])
            self.assertEqual(pi.issue_date, original['issue_date'])
            self.assertEqual(pi.payment_terms, original['payment_terms'])
            self.assertEqual(pi.shipping_cost, original['shipping_cost'])
            self.assertEqual(pi.total_amount, original['total_amount'])
            self.assertEqual(pi.customer.contact_person, original['contact'])
            self.assertEqual(item.quantity, original['quantity'])
            self.assertEqual(item.unit_price, original['price'])
            audit = AuditLog.query.filter_by(
                entity_type='pi', entity_id=self.alice_pi, action='export'
            ).order_by(AuditLog.id.desc()).first()
            self.assertIsNotNone(audit)
            self.assertIn('original_pi_unchanged', audit.after_json)

    def test_multiple_custom_excel_templates_can_be_selected(self):
        self.login('admin-test')
        self.assertEqual(self.client.get('/document-templates').status_code, 200)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'PI'
        sheet['A1'] = 'PI {{pi_number}}'
        sheet['F1'] = '{{grand_total}}'
        sheet['A3'] = 'No.'
        sheet['B3'] = 'Product'
        sheet['C3'] = 'Qty'
        sheet['D3'] = 'Amount'
        sheet['A4'] = '{{item.no}}'
        sheet['B4'] = '{{item.name}}'
        sheet['C4'] = '{{item.quantity}}'
        sheet['D4'] = '{{item.amount}}'
        sheet['A6'] = 'Notes:'
        sheet['B6'] = '{{notes}}'
        sheet.merge_cells('B6:D6')
        sheet.print_area = 'A1:F6'
        sheet['B4'].font = Font(bold=True, color='FFFFFF')
        sheet['B4'].fill = PatternFill('solid', fgColor='1A3A5C')
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        buffer.seek(0)

        upload = self.client.post('/document-templates/add', data={
            'name': '测试客户专用模板',
            'notes': '多模板切换测试',
            'current_password': 'AdminPass123!',
            'template_file': (buffer, 'custom-template.xlsx'),
            'csrf_token': self.token('/document-templates'),
        }, content_type='multipart/form-data')
        self.assertEqual(upload.status_code, 302)

        with application.app.app_context():
            custom_template = DocumentTemplate.query.filter_by(name='测试客户专用模板').one()
            template_id = custom_template.id
            second_product = Product(name='Second Product', product_code='SECOND', unit_price=3)
            db.session.add(second_product)
            db.session.flush()
            second_item = PIItem(
                pi_id=self.alice_pi, product_id=second_product.id,
                quantity=3, unit_price=3, amount=9,
            )
            db.session.add(second_item)
            db.session.commit()
            second_item_id = second_item.id

        self.client.get('/logout')
        self.login('alice')
        export_html = self.client.get(f'/pi/{self.alice_pi}/export').get_data(as_text=True)
        self.assertIn('测试客户专用模板', export_html)
        self.assertEqual(self.client.get('/document-templates').status_code, 403)

        with application.app.app_context():
            pi = db.session.get(PI, self.alice_pi)
            item_ids = [item.id for item in pi.items]
        data = {
            'template_id': template_id,
            'output_format': 'xlsx',
            'mode': 'download',
            'issue_date': '2026-01-01',
            'payment_terms': 'TT',
            'shipping_address': '',
            'shipping_note': '',
            'shipping_cost': '0',
            'notes': '',
            'cust_contact': '',
            'cust_email': '',
            'cust_phone': '',
            'cust_address': '',
            'csrf_token': self.token(f'/pi/{self.alice_pi}/export'),
        }
        for item_id in item_ids:
            with application.app.app_context():
                item = db.session.get(PIItem, item_id)
                data[f'qty_{item_id}'] = str(item.quantity)
                data[f'price_{item_id}'] = str(item.unit_price)
        response = self.client.post(f'/pi/{self.alice_pi}/export', data=data)
        self.assertEqual(response.status_code, 200)
        payload = response.data
        response.close()
        rendered = load_workbook(BytesIO(payload), data_only=False)
        sheet = rendered['PI']
        self.assertEqual(sheet['A1'].value, 'PI PI-TEST-001')
        self.assertEqual(sheet['A4'].value, 1)
        self.assertEqual(sheet['A5'].value, 2)
        self.assertEqual(sheet['B4'].value, 'Original Product')
        self.assertEqual(sheet['B5'].value, 'Second Product')
        self.assertEqual(sheet['C5'].value, 3)
        self.assertEqual(sheet['B5'].font.bold, True)
        self.assertEqual(sheet['B5'].fill.fgColor.rgb, sheet['B4'].fill.fgColor.rgb)
        self.assertEqual(sheet['A7'].value, 'Notes:')
        self.assertIn('B7:D7', {str(cell_range) for cell_range in sheet.merged_cells.ranges})
        self.assertTrue(str(sheet.print_area).endswith('$A$1:$F$7'))
        rendered.close()

        with application.app.app_context():
            db.session.delete(db.session.get(PIItem, second_item_id))
            db.session.commit()

    def test_system_default_excel_source_can_be_downloaded_and_replaced(self):
        self.login('admin-test')
        page = self.client.get('/document-templates')
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn('系统默认 Excel', html)
        self.assertIn('替换 Excel 源文件（可选）', html)
        self.assertIn('id="packingDocumentTemplates"', html)
        self.assertIn('标准完整装箱单', html)
        self.assertIn('精简 100×150 装箱单', html)
        self.assertIn('100×150 mm', html)
        self.assertIn('白底黑字', html)
        self.assertIn('/packing-lists', html)

        with application.app.app_context():
            packing_templates = DocumentTemplate.query.filter(
                DocumentTemplate.template_type.in_(['packing_a4', 'packing_compact'])
            ).all()
            self.assertEqual(len(packing_templates), 2)
            compact_template = next(
                item for item in packing_templates
                if item.template_type == 'packing_compact'
            )
            compact_template_id = compact_template.id
        packing_source = self.client.get(
            f'/document-templates/{compact_template_id}/source'
        )
        self.assertEqual(packing_source.status_code, 200)
        source_workbook = load_workbook(BytesIO(packing_source.data), data_only=False)
        self.assertEqual(str(source_workbook.active.page_setup.paperWidth), '100mm')
        self.assertEqual(str(source_workbook.active.page_setup.paperHeight), '150mm')
        source_workbook.close()
        packing_source.close()

        with application.app.app_context():
            template = DocumentTemplate.query.filter_by(code='system-default').one()
            self.assertEqual(template.template_type, 'xlsx')
            original_filename = template.filename
            original_name = template.name
            original_notes = template.notes
            template_id = template.id
            self.assertTrue(os.path.isfile(os.path.join(
                application.app.config['DOCUMENT_TEMPLATE_DIR'], original_filename,
            )))

        source = self.client.get(f'/document-templates/{template_id}/source')
        self.assertEqual(source.status_code, 200)
        downloaded = load_workbook(BytesIO(source.data), data_only=False)
        self.assertEqual(downloaded['PI Template']['A1'].value, '{{company_name}}')
        downloaded.close()
        source.close()

        replacement = Workbook()
        sheet = replacement.active
        sheet.title = 'PI Template'
        sheet['A1'] = 'Replacement {{pi_number}}'
        sheet['A3'] = '{{item.no}}'
        sheet['B3'] = '{{item.name}}'
        sheet['C3'] = '{{item.quantity}}'
        replacement_buffer = BytesIO()
        replacement.save(replacement_buffer)
        replacement.close()
        replacement_buffer.seek(0)

        response = self.client.post(
            f'/document-templates/{template_id}/edit',
            data={
                'name': original_name,
                'notes': '管理员替换测试',
                'current_password': 'AdminPass123!',
                'template_file': (replacement_buffer, 'replacement.xlsx'),
                'csrf_token': self.token('/document-templates'),
            },
            content_type='multipart/form-data',
        )
        self.assertEqual(response.status_code, 302)

        with application.app.app_context():
            template = db.session.get(DocumentTemplate, template_id)
            self.assertNotEqual(template.filename, original_filename)
            self.assertTrue(template.active)
            self.assertEqual(template.notes, '管理员替换测试')
            self.assertIsNotNone(AuditLog.query.filter_by(
                entity_type='document_template', entity_id=template_id, action='update'
            ).first())
            template.filename = original_filename
            template.name = original_name
            template.notes = original_notes
            db.session.commit()

    def test_packing_list_permissions_validation_and_excel_export(self):
        second_product_id = None
        second_item_id = None
        second_admin_id = None
        packing_list_id = None
        original_paid = None
        original_received_amount = None

        def client_token(client, path):
            html = client.get(path).get_data(as_text=True)
            match = re.search(r'<meta name="csrf-token" content="([^"]+)"', html)
            self.assertIsNotNone(match)
            return match.group(1)

        def login_client(account, password):
            client = application.app.test_client()
            response = client.post('/login', data={
                'account': account,
                'password': password,
                'csrf_token': client_token(client, '/login'),
            })
            self.assertEqual(response.status_code, 302)
            return client

        try:
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                original_paid = bool(pi.paid)
                original_received_amount = float(pi.received_amount or 0)
                pi.paid = False
                pi.received_amount = 0
                second_product = Product(
                    name='Packing Product', product_code='PACK',
                    specification='Packing spec', unit_price=4,
                )
                second_admin = User(
                    account='admin-two-login', username='admin-two',
                    password_hash=generate_password_hash('AdminTwoPass123!'),
                    role='admin', salesperson_name='', must_change_password=False,
                )
                db.session.add_all([second_product, second_admin])
                db.session.flush()
                second_item = PIItem(
                    pi_id=self.alice_pi, product_id=second_product.id,
                    quantity=2, unit_price=4, amount=8,
                )
                db.session.add(second_item)
                db.session.commit()
                first_item_id = PIItem.query.filter_by(
                    pi_id=self.alice_pi,
                ).filter(PIItem.id != second_item.id).one().id
                second_product_id = second_product.id
                second_item_id = second_item.id
                second_admin_id = second_admin.id

            self.login('admin-test')
            index = self.client.get('/packing-lists')
            self.assertEqual(index.status_code, 200)
            self.assertNotIn('PI-TEST-001', index.get_data(as_text=True))
            self.assertEqual(
                self.client.get(f'/packing-list/{self.alice_pi}').status_code, 302,
            )

            overpacked = {
                'action': 'draft',
                'version': 0,
                'packing_date': '2026-09-07',
                'boxes': [
                    {
                        'box_no': '1', 'net_weight': 1, 'gross_weight': 2,
                        'length_cm': 10, 'width_cm': 10, 'height_cm': 10,
                        'shipping_mark': 'MARK-A', 'note': '',
                        'items': [
                            {'pi_item_id': second_item_id, 'quantity': 2, 'note': ''},
                        ],
                    },
                    {
                        'box_no': '2', 'net_weight': 1, 'gross_weight': 2,
                        'length_cm': 10, 'width_cm': 10, 'height_cm': 10,
                        'shipping_mark': 'MARK-B', 'note': '',
                        'items': [
                            {'pi_item_id': second_item_id, 'quantity': 1, 'note': ''},
                        ],
                    },
                ],
            }
            response = self.client.post(
                f'/api/packing-list/{self.alice_pi}', json=overpacked,
                headers={'X-CSRFToken': self.token('/packing-lists')},
            )
            self.assertEqual(response.status_code, 409)
            self.assertIn('尚未回款', response.get_json()['error'])

            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                pi.paid = False
                pi.received_amount = 1
                db.session.commit()

            index = self.client.get('/packing-lists')
            self.assertIn('PI-TEST-001', index.get_data(as_text=True))
            editor = self.client.get(f'/packing-list/{self.alice_pi}')
            self.assertEqual(editor.status_code, 200)
            editor_html = editor.get_data(as_text=True)
            self.assertIn('id="packingProductSearch"', editor_html)
            self.assertIn('id="packingProductFilters"', editor_html)
            self.assertIn('data-packing-filter="remaining"', editor_html)
            self.assertIn("productFilter === 'remaining' && product.remaining <= 0", editor_html)
            self.assertIn('所有产品都已装箱', editor_html)
            self.assertIn('PI 产品清单', editor_html)
            self.assertIn('完整 PI，不显示价格', editor_html)
            self.assertIn('全部产品始终保留', editor_html)
            self.assertNotIn('id="showAllPackingProducts"', editor_html)
            self.assertIn('全选当前结果', editor_html)
            self.assertIn('id="packingSelectionCount"', editor_html)
            self.assertIn('id="clearSelectedProducts"', editor_html)
            self.assertIn('const selectedProductIds = new Set()', editor_html)
            self.assertIn('for (const id of selectedProductIds)', editor_html)
            self.assertIn('将全部剩余数量放入当前箱', editor_html)
            self.assertIn('id="packingTargetLabel"', editor_html)
            self.assertIn('packing-inline-quantity', editor_html)
            self.assertNotIn('产品备注（可选）', editor_html)
            self.assertIn('清空当前箱产品', editor_html)
            self.assertIn('完成箱子', editor_html)
            self.assertIn('finishBoxTop', editor_html)
            self.assertIn('function finishBoxAndCreateNext(boxIndex)', editor_html)
            self.assertIn("title='上移箱子'", editor_html)
            self.assertIn("title='下移箱子'", editor_html)
            self.assertIn('function moveBox(boxIndex, direction)', editor_html)
            self.assertNotIn("field('箱号 *'", editor_html)
            self.assertIn("body.className='card-body packing-box-body p-2'", editor_html)
            self.assertIn("measureFields.className='row g-2 mb-2 packing-measure-fields'", editor_html)
            self.assertIn("body.append(measureFields,noteFields,cbm,itemList,boxActions)", editor_html)
            self.assertGreaterEqual(editor_html.count("{className:'col'}"), 5)
            self.assertEqual(
                editor_html.count('class="col-xl packing-side-column"'), 2,
            )
            self.assertIn('.packing-side-column { flex:1 1 0 !important;', editor_html)
            self.assertIn('class="col-xl-1"', editor_html)
            self.assertNotIn('class="col-xl-7"', editor_html)
            self.assertNotIn('class="col-xl-4"', editor_html)
            self.assertIn('箱子与已分配产品', editor_html)
            self.assertIn('packing-box-item', editor_html)
            self.assertNotIn('手动添加产品行', editor_html)
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                prefill = application._packing_list_payload(pi, prefill=True)
                self.assertEqual(len(prefill['boxes']), 1)
                self.assertEqual(prefill['boxes'][0]['items'], [])
                self.assertIn('image', prefill['pi_items'][0])

            response = self.client.post(
                f'/api/packing-list/{self.alice_pi}', json=overpacked,
                headers={'X-CSRFToken': self.token(f'/packing-list/{self.alice_pi}')},
            )
            self.assertEqual(response.status_code, 400)
            self.assertIn('超过 PI 数量', response.get_json()['error'])

            valid = {
                'action': 'complete',
                'version': 0,
                'packing_date': '2026-09-07',
                'boxes': [
                    {
                        'box_no': '9', 'net_weight': 1.25, 'gross_weight': 1.5,
                        'length_cm': 10, 'width_cm': 10, 'height_cm': 10,
                        'shipping_mark': 'MARK-A', 'note': 'Mixed carton',
                        'items': [
                            {'pi_item_id': first_item_id, 'quantity': 1, 'note': ''},
                            {'pi_item_id': second_item_id, 'quantity': 1, 'note': 'Part one'},
                        ],
                    },
                    {
                        'box_no': '3', 'net_weight': 2, 'gross_weight': 2.5,
                        'length_cm': 10, 'width_cm': 20, 'height_cm': 30,
                        'shipping_mark': 'MARK-B', 'note': 'Split carton',
                        'items': [
                            {'pi_item_id': second_item_id, 'quantity': 1, 'note': 'Part two'},
                        ],
                    },
                ],
            }
            saved = self.client.post(
                f'/api/packing-list/{self.alice_pi}', json=valid,
                headers={'X-CSRFToken': self.token(f'/packing-list/{self.alice_pi}')},
            )
            self.assertEqual(saved.status_code, 200)
            saved_json = saved.get_json()
            self.assertEqual(saved_json['data']['status'], 'completed')
            self.assertEqual(
                [box['box_no'] for box in saved_json['data']['boxes']], ['1', '2'],
            )
            first_version = saved_json['data']['version']

            # A completed packing list opens read-only. Administrators must
            # explicitly enter edit mode; the editor still supports saving a
            # draft or saving and completing again.
            completed_detail = self.client.get(
                f'/packing-list/{self.alice_pi}',
            ).get_data(as_text=True)
            self.assertIn('当前为只读查看', completed_detail)
            self.assertIn('编辑装箱单', completed_detail)
            self.assertNotIn('id="packingBoxes"', completed_detail)
            completed_editor = self.client.get(
                f'/packing-list/{self.alice_pi}?edit=1',
            ).get_data(as_text=True)
            self.assertIn('id="packingBoxes"', completed_editor)
            self.assertIn('保存草稿', completed_editor)
            self.assertIn('保存并完成', completed_editor)
            self.assertIn('精简 100×150', completed_detail)
            self.assertIn('compact-100x150.xlsx', completed_detail)
            self.assertIn('compact-100x150.pdf', completed_detail)
            self.assertIn('id="compactExportModal"', completed_detail)
            self.assertIn('选择要导出的箱子', completed_detail)
            self.assertIn('compact-export-box m-0 flex-shrink-0', completed_detail)
            self.assertIn('第 1 箱 · 箱号 1', completed_detail)
            self.assertIn('第 2 箱 · 箱号 2', completed_detail)

            packing_index = self.client.get('/packing-lists').get_data(as_text=True)
            self.assertIn('<i class="bi bi-eye"></i> 预览', packing_index)
            self.assertIn('选择 Excel 模板', packing_index)
            self.assertIn('选择 PDF 模板', packing_index)
            self.assertIn('完整 A4', packing_index)
            self.assertIn('精简 100×150（每箱一页）', packing_index)
            self.assertIn(
                f'/packing-list/{self.alice_pi}/compact-100x150.xlsx', packing_index,
            )
            self.assertIn(
                f'/packing-list/{self.alice_pi}/compact-100x150.pdf', packing_index,
            )

            with application.app.app_context():
                db.session.get(PI, self.alice_pi).shipping_completed = True
                db.session.commit()
            shipped_detail = self.client.get(
                f'/packing-list/{self.alice_pi}',
            ).get_data(as_text=True)
            self.assertIn('该订单已经登记发货，确认仍要编辑装箱单吗？', shipped_detail)
            with application.app.app_context():
                db.session.get(PI, self.alice_pi).shipping_completed = False
                db.session.commit()

            with application.app.app_context():
                packing_list = PackingList.query.filter_by(pi_id=self.alice_pi).one()
                packing_list_id = packing_list.id
                self.assertEqual(len(packing_list.boxes), 2)
                self.assertEqual([len(box.items) for box in packing_list.boxes], [2, 1])
                self.assertAlmostEqual(packing_list.boxes[0].volume_cbm, 0.001)
                self.assertAlmostEqual(packing_list.boxes[1].volume_cbm, 0.006)

            excel = self.client.get(
                f'/packing-list/{self.alice_pi}/export.xlsx',
            )
            self.assertEqual(excel.status_code, 200)
            workbook = load_workbook(BytesIO(excel.data), data_only=False)
            self.assertIn('装箱单', workbook.sheetnames)
            values = [
                cell.value for row in workbook['装箱单'].iter_rows() for cell in row
                if cell.value is not None
            ]
            self.assertTrue(any('PI-TEST-001' in str(value) for value in values))
            self.assertIn('Original Product', values)
            self.assertIn('Packing Product', values)
            self.assertTrue(str(workbook['装箱单'].print_area))
            workbook.close()
            excel.close()

            compact_excel = self.client.get(
                f'/packing-list/{self.alice_pi}/compact-100x150.xlsx',
            )
            self.assertEqual(compact_excel.status_code, 200)
            compact_workbook = load_workbook(BytesIO(compact_excel.data), data_only=False)
            self.assertEqual(len(compact_workbook.sheetnames), 2)
            for sheet in compact_workbook.worksheets:
                compact_values = [
                    cell.value for row in sheet.iter_rows() for cell in row
                    if cell.value is not None
                ]
                self.assertIn('PACKING LIST / 装箱单', compact_values)
                self.assertTrue(str(sheet.print_area))
                self.assertEqual(str(sheet.page_setup.paperWidth), '100mm')
                self.assertEqual(str(sheet.page_setup.paperHeight), '150mm')
                self.assertNotIn('Exporter / 出口商', compact_values)
                self.assertNotIn('Buyer / 客户', compact_values)
                self.assertIn('Sales / 业务员', compact_values)
                self.assertTrue(any(
                    'Packing spec' in str(value) for value in compact_values
                ))
                self.assertEqual(sheet['A1'].fill.fgColor.rgb, '00FFFFFF')
                self.assertEqual(sheet['A1'].font.color.rgb, '00000000')
                self.assertTrue(sheet['A1'].font.bold)
                self.assertGreaterEqual(sheet['A1'].font.sz, 17)
                self.assertEqual(sheet['A8'].fill.fgColor.rgb, '00FFFFFF')
                self.assertEqual(sheet['A8'].font.color.rgb, '00000000')
                self.assertTrue(sheet['A8'].font.bold)
                self.assertGreaterEqual(sheet['A8'].font.sz, 10)
            compact_workbook.close()
            compact_excel.close()

            selected_compact_excel = self.client.get(
                f'/packing-list/{self.alice_pi}/compact-100x150.xlsx?box=2',
            )
            self.assertEqual(selected_compact_excel.status_code, 200)
            self.assertIn(
                'Cartons-2',
                selected_compact_excel.headers.get('Content-Disposition', ''),
            )
            selected_workbook = load_workbook(
                BytesIO(selected_compact_excel.data), data_only=False,
            )
            self.assertEqual(len(selected_workbook.sheetnames), 1)
            self.assertTrue(selected_workbook.sheetnames[0].startswith('第2箱-'))
            self.assertIn('2/2', str(selected_workbook.active['C5'].value))
            selected_workbook.close()
            selected_compact_excel.close()
            self.assertEqual(
                self.client.get(
                    f'/packing-list/{self.alice_pi}/compact-100x150.xlsx?box=0',
                ).status_code,
                400,
            )
            self.assertEqual(
                self.client.get(
                    f'/packing-list/{self.alice_pi}/compact-100x150.xlsx?box=3',
                ).status_code,
                400,
            )

            with patch.object(
                application, 'convert_excel_to_pdf',
                side_effect=AssertionError(
                    'compact PDF must use the exact 100x150 renderer'
                ),
            ):
                compact_pdf = self.client.get(
                    f'/packing-list/{self.alice_pi}/compact-100x150.pdf',
                )
            self.assertEqual(compact_pdf.status_code, 200)
            self.assertEqual(compact_pdf.mimetype, 'application/pdf')
            self.assertTrue(compact_pdf.data.startswith(b'%PDF-'))
            self.assertEqual(len(re.findall(rb'/Type\s*/Page\b', compact_pdf.data)), 2)
            self.assertRegex(
                compact_pdf.data,
                rb'/MediaBox\s*\[\s*0\s+0\s+283\.[0-9]+\s+425\.[0-9]+\s*\]',
            )
            compact_pdf.close()

            selected_compact_pdf = self.client.get(
                f'/packing-list/{self.alice_pi}/compact-100x150.pdf?box=2',
            )
            self.assertEqual(selected_compact_pdf.status_code, 200)
            self.assertEqual(
                len(re.findall(rb'/Type\s*/Page\b', selected_compact_pdf.data)),
                1,
            )
            selected_compact_pdf.close()

            # Any administrator can edit a packing list, even when another
            # administrator created it.
            second_admin_client = login_client('admin-two-login', 'AdminTwoPass123!')
            second_admin_detail = second_admin_client.get(
                f'/packing-list/{self.alice_pi}',
            )
            self.assertEqual(second_admin_detail.status_code, 200)
            self.assertIn('编辑装箱单', second_admin_detail.get_data(as_text=True))
            second_admin_payload = dict(valid)
            second_admin_payload['version'] = first_version
            second_admin_payload['action'] = 'draft'
            second_admin_payload['boxes'] = [dict(box) for box in valid['boxes']]
            second_admin_payload['boxes'][0]['shipping_mark'] = 'EDITED-BY-SECOND-ADMIN'
            updated = second_admin_client.post(
                f'/api/packing-list/{self.alice_pi}', json=second_admin_payload,
                headers={
                    'X-CSRFToken': client_token(
                        second_admin_client, f'/packing-list/{self.alice_pi}',
                    ),
                },
            )
            self.assertEqual(updated.status_code, 200)
            updated_json = updated.get_json()
            self.assertEqual(updated_json['data']['status'], 'draft')
            self.assertGreater(updated_json['data']['version'], first_version)
            with application.app.app_context():
                packing_list = db.session.get(PackingList, packing_list_id)
                self.assertEqual(packing_list.updated_by, 'admin-two')
                self.assertIsNotNone(packing_list.updated_at)

            second_admin_payload['action'] = 'complete'
            second_admin_payload['version'] = updated_json['data']['version']
            recompleted = second_admin_client.post(
                f'/api/packing-list/{self.alice_pi}', json=second_admin_payload,
                headers={
                    'X-CSRFToken': client_token(
                        second_admin_client, f'/packing-list/{self.alice_pi}?edit=1',
                    ),
                },
            )
            self.assertEqual(recompleted.status_code, 200)
            self.assertEqual(recompleted.get_json()['data']['status'], 'completed')

            stale = self.client.post(
                f'/api/packing-list/{self.alice_pi}', json=valid,
                headers={'X-CSRFToken': self.token(f'/packing-list/{self.alice_pi}')},
            )
            self.assertEqual(stale.status_code, 409)

            # If all payments are later removed, retain its
            # historical packing list for read/export but lock all changes.
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                pi.paid = False
                pi.received_amount = 0
                db.session.commit()
            history_detail = self.client.get(f'/packing-list/{self.alice_pi}')
            self.assertEqual(history_detail.status_code, 200)
            self.assertIn('历史装箱单仅允许查看和导出', history_detail.get_data(as_text=True))
            self.assertIn(
                'PI-TEST-001', self.client.get('/packing-lists').get_data(as_text=True),
            )
            locked = self.client.post(
                f'/api/packing-list/{self.alice_pi}', json=second_admin_payload,
                headers={'X-CSRFToken': self.token(f'/packing-list/{self.alice_pi}')},
            )
            self.assertEqual(locked.status_code, 409)
            self.assertIn('尚未回款', locked.get_json()['error'])
            with application.app.app_context():
                pi = db.session.get(PI, self.alice_pi)
                pi.paid = True
                pi.received_amount = pi.grand_total
                db.session.commit()

            bob = login_client('bob-login', 'StrongPass123!')
            bob_index = bob.get('/packing-lists')
            self.assertEqual(bob_index.status_code, 200)
            self.assertNotIn('PI-TEST-001', bob_index.get_data(as_text=True))
            self.assertEqual(bob.get(f'/packing-list/{self.alice_pi}').status_code, 403)
            self.assertEqual(
                bob.get(f'/packing-list/{self.alice_pi}/export.xlsx').status_code,
                403,
            )
            self.assertEqual(
                bob.post(
                    f'/api/packing-list/{self.alice_pi}', json=valid,
                    headers={'X-CSRFToken': client_token(bob, '/packing-lists')},
                ).status_code,
                403,
            )

            alice = login_client('alice-login', 'StrongPass123!')
            alice_detail = alice.get(f'/packing-list/{self.alice_pi}')
            self.assertEqual(alice_detail.status_code, 200)
            self.assertIn('当前为只读查看', alice_detail.get_data(as_text=True))
            self.assertNotIn('编辑装箱单', alice_detail.get_data(as_text=True))
            self.assertNotIn(
                'id="packingBoxes"',
                alice.get(f'/packing-list/{self.alice_pi}?edit=1').get_data(as_text=True),
            )
            self.assertEqual(
                alice.get(f'/packing-list/{self.alice_pi}/export.xlsx').status_code,
                200,
            )
            self.assertEqual(
                alice.post(
                    f'/api/packing-list/{self.alice_pi}', json=valid,
                    headers={
                        'X-CSRFToken': client_token(
                            alice, f'/packing-list/{self.alice_pi}',
                        ),
                    },
                ).status_code,
                403,
            )
        finally:
            with application.app.app_context():
                db.session.rollback()
                if packing_list_id:
                    packing_list = db.session.get(PackingList, packing_list_id)
                    if packing_list:
                        db.session.delete(packing_list)
                        db.session.flush()
                if second_item_id:
                    second_item = db.session.get(PIItem, second_item_id)
                    if second_item:
                        db.session.delete(second_item)
                        db.session.flush()
                if second_product_id:
                    second_product = db.session.get(Product, second_product_id)
                    if second_product:
                        db.session.delete(second_product)
                if second_admin_id:
                    AuditLog.query.filter_by(user_id=second_admin_id).update(
                        {'user_id': None}, synchronize_session=False,
                    )
                    second_admin = db.session.get(User, second_admin_id)
                    if second_admin:
                        db.session.delete(second_admin)
                if original_paid is not None:
                    pi = db.session.get(PI, self.alice_pi)
                    if pi:
                        pi.paid = original_paid
                        pi.received_amount = original_received_amount
                db.session.commit()

    @unittest.skipUnless(find_soffice(), 'LibreOffice is only required on the production server')
    def test_server_can_convert_custom_excel_template_to_pdf(self):
        work_dir = tempfile.mkdtemp(prefix='pi-template-convert-test-')
        try:
            excel_path = os.path.join(work_dir, 'template.xlsx')
            pdf_path = os.path.join(work_dir, 'template.pdf')
            workbook = Workbook()
            sheet = workbook.active
            sheet['A1'] = 'PI template conversion test'
            workbook.save(excel_path)
            workbook.close()
            convert_excel_to_pdf(excel_path, pdf_path)
            with open(pdf_path, 'rb') as pdf_file:
                self.assertEqual(pdf_file.read(4), b'%PDF')
        finally:
            shutil.rmtree(work_dir, ignore_errors=True)


if __name__ == '__main__':
    unittest.main(verbosity=2)
