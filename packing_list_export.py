"""Generate packing-list workbooks used for both Excel and PDF exports."""

from io import BytesIO
from math import ceil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from reportlab.lib.colors import HexColor
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas

from pdf_generator import _FONT as COMPACT_PDF_FONT


NAVY = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F3F5F7"
GREEN = "198754"
COMPACT_BLACK = "000000"
COMPACT_WHITE = "FFFFFF"
COMPACT_PAGE_SIZE = (100 * mm, 150 * mm)


def _text(value):
    """Keep arbitrary business text from becoming an Excel formula."""
    value = str(value or "")
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _compact_number(value):
    return f"{float(value or 0):g}"


def _compact_sheet_title(index, box_no, used_titles):
    safe_box_no = "".join(
        char for char in str(box_no or "") if char not in "[]:*?/\\"
    ).strip()
    base = f"第{index}箱-{safe_box_no}"[:31] or f"第{index}箱"
    title = base
    suffix = 2
    while title in used_titles:
        tail = f"-{suffix}"
        title = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    used_titles.add(title)
    return title


def _compact_excel_sheet(sheet, pi, packing_list, box, box_index, box_count):
    """Lay out one 100 x 150 mm compact carton page on one worksheet."""
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.paperSize = None
    sheet.page_setup.paperWidth = "100mm"
    sheet.page_setup.paperHeight = "150mm"
    sheet.page_margins = PageMargins(
        left=0.15, right=0.15, top=0.18, bottom=0.18, header=0, footer=0
    )
    sheet.print_options.horizontalCentered = True
    sheet.print_options.verticalCentered = True

    for column, width in zip("ABCDEF", (5, 12, 12, 12, 12, 9)):
        sheet.column_dimensions[column].width = width

    thin = Side(style="thin", color=COMPACT_BLACK)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    normal = Font(name="Microsoft YaHei", size=9, color=COMPACT_BLACK)
    label = Font(name="Microsoft YaHei", size=8, bold=True, color=COMPACT_BLACK)

    sheet.merge_cells("A1:F2")
    sheet["A1"] = "PACKING LIST / 装箱单"
    sheet["A1"].font = Font(
        name="Microsoft YaHei", size=15, bold=True, color=COMPACT_BLACK
    )
    sheet["A1"].alignment = center
    sheet["A1"].fill = PatternFill("solid", fgColor=COMPACT_WHITE)
    for row in sheet["A1:F2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=COMPACT_WHITE)
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 14

    info_rows = [
        (3, "PI No. / PI编号", pi.pi_number),
        (4, "Sales / 业务员", pi.salesperson),
        (5, "Carton / 箱号", f"{box.box_no}  ({box_index}/{box_count})"),
    ]
    for row_no, field_label, value in info_rows:
        sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=2)
        sheet.merge_cells(start_row=row_no, start_column=3, end_row=row_no, end_column=6)
        sheet.cell(row_no, 1, field_label).font = label
        sheet.cell(row_no, 3, _text(value)).font = normal
        sheet.cell(row_no, 1).alignment = left
        sheet.cell(row_no, 3).alignment = left
        sheet.row_dimensions[row_no].height = 20

    header_row = 7
    sheet.merge_cells(start_row=header_row, start_column=2, end_row=header_row, end_column=5)
    for column, value in ((1, "#"), (2, "Product / 产品"), (6, "Qty / 数量")):
        cell = sheet.cell(header_row, column, value)
        cell.font = Font(
            name="Microsoft YaHei", size=8, bold=True, color=COMPACT_WHITE
        )
        cell.fill = PatternFill("solid", fgColor=COMPACT_BLACK)
        cell.alignment = center
    for column in range(1, 7):
        sheet.cell(header_row, column).fill = PatternFill(
            "solid", fgColor=COMPACT_BLACK
        )
        sheet.cell(header_row, column).border = border
    sheet.row_dimensions[header_row].height = 22

    item_row = header_row + 1
    items = list(box.items)
    if not items:
        items = [None]
    item_height = max(9, min(24, 150 / max(1, len(items))))
    item_font_size = max(6, min(9, item_height / 2.2))
    for item_index, item in enumerate(items, 1):
        sheet.merge_cells(start_row=item_row, start_column=2, end_row=item_row, end_column=5)
        values = (
            item_index if item else "",
            _text(item.product_name) if item else "No products / 暂无产品",
            int(item.quantity) if item else "",
        )
        for column, value in ((1, values[0]), (2, values[1]), (6, values[2])):
            cell = sheet.cell(item_row, column, value)
            cell.font = Font(
                name="Microsoft YaHei", size=item_font_size, color=COMPACT_BLACK
            )
            cell.alignment = center if column in (1, 6) else left
        for column in range(1, 7):
            sheet.cell(item_row, column).border = border
            sheet.cell(item_row, column).fill = PatternFill(
                "solid", fgColor=COMPACT_WHITE
            )
        sheet.row_dimensions[item_row].height = item_height
        item_row += 1

    metrics_row = item_row + 1
    metrics = [
        ("Size / 尺寸", f"{_compact_number(box.length_cm)} x {_compact_number(box.width_cm)} x {_compact_number(box.height_cm)} cm"),
        ("N.W. / 净重", f"{_compact_number(box.net_weight)} kg"),
        ("G.W. / 毛重", f"{_compact_number(box.gross_weight)} kg"),
    ]
    for offset, (field_label, value) in enumerate(metrics):
        row_no = metrics_row + offset
        sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=2)
        sheet.merge_cells(start_row=row_no, start_column=3, end_row=row_no, end_column=6)
        sheet.cell(row_no, 1, field_label).font = label
        sheet.cell(row_no, 3, value).font = Font(
            name="Microsoft YaHei", size=10, bold=True, color=COMPACT_BLACK
        )
        for column in range(1, 7):
            sheet.cell(row_no, column).border = border
            sheet.cell(row_no, column).fill = PatternFill(
                "solid", fgColor=COMPACT_WHITE
            )
        sheet.cell(row_no, 1).alignment = left
        sheet.cell(row_no, 3).alignment = left
        sheet.row_dimensions[row_no].height = 22

    last_row = metrics_row + len(metrics) - 1
    sheet.print_area = f"A1:F{last_row}"


def generate_compact_packing_list_workbook(pi, packing_list):
    """Generate the compact 100 x 150 mm workbook, one carton per sheet/page."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles = set()
    boxes = list(packing_list.boxes)
    for box_index, box in enumerate(boxes, 1):
        title = _compact_sheet_title(box_index, box.box_no, used_titles)
        sheet = workbook.create_sheet(title)
        _compact_excel_sheet(sheet, pi, packing_list, box, box_index, len(boxes))

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output.getvalue()


def _ensure_compact_pdf_font():
    pdfmetrics.getFont(COMPACT_PDF_FONT)
    return COMPACT_PDF_FONT


def _fit_pdf_text(value, font_name, font_size, max_width):
    text = str(value or "")
    if pdfmetrics.stringWidth(text, font_name, font_size) <= max_width:
        return text
    suffix = "..."
    while text and pdfmetrics.stringWidth(text + suffix, font_name, font_size) > max_width:
        text = text[:-1]
    return text + suffix


def generate_compact_packing_list_pdf(pi, packing_list):
    """Generate an exact 100 x 150 mm PDF with one carton on every page."""
    font_name = _ensure_compact_pdf_font()
    output = BytesIO()
    document = canvas.Canvas(output, pagesize=COMPACT_PAGE_SIZE, pageCompression=1)
    document.setTitle(f"Packing List {pi.pi_number} - Compact 100x150")
    page_width, page_height = COMPACT_PAGE_SIZE
    margin = 5 * mm
    content_width = page_width - 2 * margin
    boxes = list(packing_list.boxes)
    for box_index, box in enumerate(boxes, 1):
        document.setFillColor(HexColor(f"#{COMPACT_WHITE}"))
        document.setStrokeColor(HexColor(f"#{COMPACT_BLACK}"))
        document.setLineWidth(0.6)
        document.roundRect(
            margin, page_height - 24 * mm, content_width, 19 * mm,
            2 * mm, fill=1, stroke=1,
        )
        document.setFillColor(HexColor(f"#{COMPACT_BLACK}"))
        document.setFont(font_name, 13)
        document.drawString(margin + 4 * mm, page_height - 13 * mm, "PACKING LIST / 装箱单")
        document.setFont(font_name, 8)
        document.drawRightString(
            page_width - margin - 4 * mm,
            page_height - 20 * mm,
            f"Carton {box_index}/{len(boxes)}",
        )

        document.setFillColor(HexColor(f"#{COMPACT_BLACK}"))
        info_y = page_height - 30 * mm
        document.setFont(font_name, 8)
        document.drawString(margin, info_y, "PI No. / PI编号")
        document.setFont(font_name, 9)
        document.drawString(
            margin + 27 * mm, info_y,
            _fit_pdf_text(pi.pi_number, font_name, 9, content_width - 27 * mm),
        )
        info_y -= 7 * mm
        document.setFont(font_name, 8)
        document.drawString(margin, info_y, "Sales / 业务员")
        document.setFont(font_name, 9)
        document.drawString(
            margin + 27 * mm, info_y,
            _fit_pdf_text(pi.salesperson, font_name, 9, content_width - 27 * mm),
        )
        info_y -= 7 * mm
        document.setFont(font_name, 8)
        document.drawString(margin, info_y, "Carton / 箱号")
        document.setFont(font_name, 11)
        document.setFillColor(HexColor(f"#{COMPACT_BLACK}"))
        document.drawString(
            margin + 27 * mm, info_y,
            _fit_pdf_text(box.box_no, font_name, 11, content_width - 27 * mm),
        )

        table_top = page_height - 51 * mm
        table_bottom = 46 * mm
        table_height = table_top - table_bottom
        document.setFillColor(HexColor(f"#{COMPACT_BLACK}"))
        document.setFont(font_name, 8)
        document.drawString(margin, table_top + 3 * mm, "CONTENTS / 箱内产品")

        items = list(box.items)
        if not items:
            items = [None]
        column_count = min(3, max(1, ceil(len(items) / 12.0)))
        row_count = ceil(len(items) / float(column_count))
        cell_width = content_width / column_count
        row_height = table_height / max(1, row_count)
        font_size = max(4.5, min(8.5, row_height - 4))
        for item_index, item in enumerate(items):
            column = item_index // row_count
            row = item_index % row_count
            x = margin + column * cell_width
            y = table_top - (row + 1) * row_height
            document.setStrokeColor(HexColor(f"#{COMPACT_BLACK}"))
            document.setLineWidth(0.45)
            document.rect(x, y, cell_width, row_height, fill=0, stroke=1)
            if item:
                item_text = f"{item_index + 1}. {item.product_name or ''}"
                quantity_text = f"x {int(item.quantity)}"
            else:
                item_text = "No products / 暂无产品"
                quantity_text = ""
            document.setFillColor(HexColor(f"#{COMPACT_BLACK}"))
            document.setFont(font_name, font_size)
            quantity_width = (
                pdfmetrics.stringWidth(quantity_text, font_name, font_size) + 2 * mm
                if quantity_text else 0
            )
            document.drawString(
                x + 1.5 * mm,
                y + max(1.2 * mm, (row_height - font_size) / 2),
                _fit_pdf_text(
                    item_text, font_name, font_size,
                    cell_width - 3 * mm - quantity_width,
                ),
            )
            if quantity_text:
                document.drawRightString(
                    x + cell_width - 1.5 * mm,
                    y + max(1.2 * mm, (row_height - font_size) / 2),
                    quantity_text,
                )

        metric_top = 40 * mm
        metric_height = 9 * mm
        metrics = [
            ("Size / 尺寸", f"{_compact_number(box.length_cm)} x {_compact_number(box.width_cm)} x {_compact_number(box.height_cm)} cm"),
            ("N.W. / 净重", f"{_compact_number(box.net_weight)} kg"),
            ("G.W. / 毛重", f"{_compact_number(box.gross_weight)} kg"),
        ]
        for metric_index, (field_label, value) in enumerate(metrics):
            y = metric_top - (metric_index + 1) * metric_height
            document.setFillColor(HexColor(f"#{COMPACT_WHITE}"))
            document.setStrokeColor(HexColor(f"#{COMPACT_BLACK}"))
            document.rect(margin, y, content_width, metric_height, fill=1, stroke=1)
            document.setFillColor(HexColor(f"#{COMPACT_BLACK}"))
            document.setFont(font_name, 7.5)
            document.drawString(margin + 2 * mm, y + 3 * mm, field_label)
            document.setFillColor(HexColor(f"#{COMPACT_BLACK}"))
            document.setFont(font_name, 10)
            document.drawRightString(page_width - margin - 2 * mm, y + 2.7 * mm, value)

        document.showPage()

    document.save()
    output.seek(0)
    return output.getvalue()


def generate_packing_list_workbook(pi, packing_list, company_name, company_address):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "装箱单"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A11"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(
        left=0.22, right=0.22, top=0.35, bottom=0.35, header=0.15, footer=0.15
    )

    thin = Side(style="thin", color="B8C2CC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    normal = Font(name="Microsoft YaHei", size=9, color="1F2937")
    label = Font(name="Microsoft YaHei", size=9, bold=True, color="4A5568")

    widths = [7, 14, 25, 15, 18, 9, 11, 11, 9, 9, 9, 11, 22, 22]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width

    sheet.merge_cells("A1:N2")
    title = sheet["A1"]
    title.value = "PACKING LIST / 装箱单"
    title.font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    title.alignment = center
    for row in sheet["A1:N2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)

    customer = pi.customer
    details = [
        (4, "Exporter / 出口商", company_name, "PI No. / PI 编号", pi.pi_number),
        (5, "Address / 地址", company_address, "Packing Date / 装箱日期", packing_list.packing_date.isoformat() if packing_list.packing_date else ""),
        (6, "Buyer / 客户", customer.name if customer else "", "Salesperson / 业务员", pi.salesperson),
        (7, "Ship To / 收货地址", pi.shipping_address or (customer.address if customer else ""), "Status / 状态", "Completed / 已完成" if packing_list.is_completed else "Draft / 草稿"),
    ]
    for row_no, left_label, left_value, right_label, right_value in details:
        sheet.cell(row_no, 1, left_label).font = label
        sheet.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=8)
        sheet.cell(row_no, 2, _text(left_value)).font = normal
        sheet.cell(row_no, 9, right_label).font = label
        sheet.merge_cells(start_row=row_no, start_column=10, end_row=row_no, end_column=14)
        sheet.cell(row_no, 10, _text(right_value)).font = normal
        for col in range(1, 15):
            sheet.cell(row_no, col).alignment = left

    headers = [
        "箱号\nCarton No.", "产品编码\nCode", "产品\nProduct", "规格\nSpecification",
        "数量\nQty", "净重(kg)\nN.W.", "毛重(kg)\nG.W.", "长(cm)\nL",
        "宽(cm)\nW", "高(cm)\nH", "体积(m³)\nCBM", "唛头\nShipping Mark",
        "箱备注\nCarton Notes", "产品备注\nItem Notes",
    ]
    header_row = 10
    for col, value in enumerate(headers, 1):
        cell = sheet.cell(header_row, col, value)
        cell.font = Font(name="Microsoft YaHei", size=8, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = center
        cell.border = border
    sheet.row_dimensions[header_row].height = 36

    row_no = header_row + 1
    total_qty = total_net = total_gross = total_cbm = 0
    for box_index, box in enumerate(packing_list.boxes):
        box_items = list(box.items)
        for item_index, item in enumerate(box_items):
            values = [
                box.box_no if item_index == 0 else "",
                item.product_code,
                item.product_name,
                item.specification,
                int(item.quantity),
                box.net_weight if item_index == 0 else "",
                box.gross_weight if item_index == 0 else "",
                box.length_cm if item_index == 0 else "",
                box.width_cm if item_index == 0 else "",
                box.height_cm if item_index == 0 else "",
                box.volume_cbm if item_index == 0 else "",
                box.shipping_mark if item_index == 0 else "",
                box.note if item_index == 0 else "",
                item.note,
            ]
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row_no, col, _text(value) if isinstance(value, str) else value)
                cell.font = normal
                cell.border = border
                cell.alignment = center if col in (1, 2, 5, 6, 7, 8, 9, 10, 11) else left
                if box_index % 2:
                    cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
                if col in (6, 7, 8, 9, 10):
                    cell.number_format = '0.00'
                if col == 11:
                    cell.number_format = '0.0000'
            sheet.row_dimensions[row_no].height = 32
            total_qty += int(item.quantity)
            row_no += 1
        total_net += float(box.net_weight or 0)
        total_gross += float(box.gross_weight or 0)
        total_cbm += float(box.volume_cbm or 0)

    total_row = row_no
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    sheet.cell(total_row, 1, f"合计 / Total — {len(packing_list.boxes)} 箱 / cartons").alignment = right
    sheet.cell(total_row, 5, total_qty)
    sheet.cell(total_row, 6, round(total_net, 2))
    sheet.cell(total_row, 7, round(total_gross, 2))
    sheet.merge_cells(start_row=total_row, start_column=8, end_row=total_row, end_column=10)
    sheet.cell(total_row, 11, round(total_cbm, 4))
    sheet.merge_cells(start_row=total_row, start_column=12, end_row=total_row, end_column=14)
    for col in range(1, 15):
        cell = sheet.cell(total_row, col)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=GREEN)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.border = border
        cell.alignment = center if col != 1 else right
    sheet.row_dimensions[total_row].height = 28

    sheet.auto_filter.ref = f"A{header_row}:N{max(header_row, row_no - 1)}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.print_area = f"A1:N{total_row}"
    sheet.oddFooter.center.text = "Page &P / &N"
    sheet.oddFooter.center.size = 8

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()
